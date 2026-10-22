"""Column-sizing contract for pp_report.fixed().

Every case here pins behaviour that was silently removed once. Commit 593e7f7
(14.08.2026) rewrote `fixed()` while adding the brand assets and the glyph guard;
its message describes fonts, the weekly report and brand assets, and says nothing
about layout. The overflow compression and the word-boundary entry matching went
with it and nothing noticed for two weeks, because no test covered them.

These run without LibreOffice, Letta or a network — python-docx builds the table
in memory and the assertions read the resulting column widths back off the XML.
"""
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

docx = pytest.importorskip("docx", reason="python-docx is the engine's own dependency")
from docx import Document                                          # noqa: E402
import pp_report as R                                              # noqa: E402


def widths(rows, **kw):
    """Build a table from `rows` (first row = header), size it, return column cm."""
    d = Document()
    t = d.add_table(rows=len(rows), cols=len(rows[0]))
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            t.cell(i, j).text = v
    R.fixed(t, **kw)
    return [c.width.cm for c in t.columns]


WIDE = [
    ["Серија | Batch", "Параметар | Parameter",
     "Критериум за прифаќање | Acceptance criteria",
     "Добиена вредност | Obtained value", "Лабораторија | Laboratory"],
    ["GG1024_01", "Вкупен Δ9-THC | Total Δ9-THC", "18.0 – 22.0 %", "19.84 %",
     "Фармахем | Farmahem"],
    ["CJ062501/2", "TYMC", "≤ 10⁴ CFU/g", "4.2 × 10⁴ CFU/g", "ИЈЗ | IPH"],
]


# ------------------------------------------------------------------ the page
@pytest.mark.parametrize("rows", [WIDE, WIDE[:2], [["A", "B"], ["1", "2"]]])
def test_no_table_exceeds_the_page(rows):
    """The house rule at the top of pp_report.py: every table FITS the page.
    Fits, not fills — a narrow two-column table is left compact rather than
    stretched across the text block, which is why this is `<=` and not `==`."""
    assert sum(widths(rows)) <= R.PAGE_W + 0.05


def test_an_overflowing_table_is_widened_to_exactly_the_page():
    """Where the content genuinely does not fit, the sized table must occupy the
    full text block — anything less means content was discarded rather than
    compressed."""
    assert sum(widths(WIDE)) == pytest.approx(R.PAGE_W, abs=0.05)


def test_no_column_collapses():
    """Proportional sizing on an overflowing table can drive a narrow column to
    nothing. The compression path floors each column instead."""
    assert all(x > 0.5 for x in widths(WIDE))


# ------------------------------------------------- overflow compression path
def test_wide_table_is_compressed_not_scaled_proportionally():
    """The regression guard for 593e7f7.

    Plain proportional scaling shrinks every column by the same ratio, so the
    widest column keeps the largest share. Compression floors each column to its
    own data and header-word demand first, then gives the *slack* to the longest
    headers — so a column carrying short data but a long header ends up wider
    than its proportional share, and a column of long data keeps its floor.
    The two strategies therefore cannot produce the same widths."""
    w = widths(WIDE)
    ideal_sum_scaling = None  # what plain proportional sizing would have produced
    # Reproduce the proportional branch directly from the same measurements.
    CHCM, PAD, cap = 0.176, 0.42, 34
    ideal = [0.8] * len(WIDE[0])
    for row in WIDE:
        for j, cell in enumerate(row):
            parts = cell.split(" | ")
            L = (len(parts[0]) + 3 + 0.8 * len(parts[1])) if len(parts) == 2 else len(cell)
            ideal[j] = max(ideal[j], min(L, cap) * CHCM + PAD)
    total = sum(ideal)
    assert total > R.PAGE_W, "fixture must overflow, or this test proves nothing"
    ideal_sum_scaling = [x / total * R.PAGE_W for x in ideal]

    assert w != pytest.approx(ideal_sum_scaling, abs=0.01), (
        "widths match plain proportional scaling — the overflow compression "
        "branch is gone (this is exactly what 593e7f7 removed)")


def test_explicit_weights_bypass_compression():
    """A caller that passes weights has already decided; compression must not
    second-guess it. Several builders rely on this — cover_page() sizes its
    approval table with fixed(t, [5.0, 5.46, 3.5, 4.5])."""
    w = widths([["A", "B", "C", "D"], ["1", "2", "3", "4"]], weights=[5.0, 5.46, 3.5, 4.5])
    assert sum(w) == pytest.approx(R.PAGE_W, abs=0.05)
    assert w[1] > w[0] > w[3] > w[2], "the caller's proportions were not preserved"


# ------------------------------------------------ word-boundary entry sizing
def test_entry_column_matching_respects_word_boundaries():
    """`_ENTRY_TARGETS` purpose-sizes hand-written columns — signature, date,
    name. Matching on a bare substring makes 'име' hit inside 'примерок', so a
    sample-description column gets sized as a name field. The comment in the
    source calls this out by name; the substring form is what 593e7f7 restored
    by accident."""
    sample = widths([["Опис на примерок | Sample description", "Вредност | Value"],
                     ["x", "1"]])
    name = widths([["Име | Name", "Вредност | Value"], ["x", "1"]])
    assert sample[0] != pytest.approx(name[0], abs=0.01), (
        "'примерок' was sized as a name column — substring matching regressed")


@pytest.mark.parametrize("header,other", [
    ("Потпис | Signature", "Забелешка | Note"),
    ("Датум | Date", "Забелешка | Note"),
])
def test_real_entry_columns_still_match(header, other):
    """The counterpart: tightening to word boundaries must not stop the genuine
    signature and date columns from being purpose-sized."""
    tight = widths([[header, other], ["", ""]])
    loose = widths([["Опис | Description", other], ["", ""]])
    assert tight[0] != pytest.approx(loose[0], abs=0.01), f"{header} lost its entry sizing"


# ------------------------------------- content-aware sizing (Head of QC, 14.10.2026)
# The transport annexes of 13.10.2026 showed both the good and the bad case side by side:
# spare width was shared in proportion to each column's natural width, so an empty Yes/No
# column and a short reference column stretched while the long item text wrapped into tall
# rows, and a "Вкупно" label made the № column the widest on the page.

CHECKLIST = [
    ["№", "Ставка | Item", "Да | Yes", "Не | No", "N/A", "Референца | Reference"],
    ["A1", "Барањето до МВР поднесено пред поаѓањето, доказ за поднесување на досие | "
           "MoIA request submitted before departure, proof of submission on file", "", "", "", "A02"],
    ["A4", "Рутата е QUALIFIED за ова возило, конфигурација на палети и сезона | "
           "Lane QUALIFIED for this vehicle, pallet configuration and season", "", "", "", "QASOP_0XX_A06"],
]


def test_long_text_column_takes_the_spare_width():
    w = widths(CHECKLIST, mode="full")
    assert sum(w) == pytest.approx(R.PAGE_W, abs=0.05)
    assert w[1] == max(w), "the item text must be the widest column"
    assert w[1] > 0.55 * R.PAGE_W
    for j in (2, 3, 4):
        assert w[j] < 1.6, f"tick column {j} stretched to {w[j]:.2f} cm"
    assert w[5] < 3.5, "a short reference column must not take the width the text needs"
    assert w[0] < 1.5, "an ordinal column takes only what A1…A22 needs"


def test_a_total_row_does_not_widen_the_ordinal_column():
    rows = [["№", "Број на кеса | Bag No.", "Нето маса (g) | Net mass (g)",
             "Кесата е цела и затворена | Bag intact and sealed"]]
    rows += [[str(i), "", "", ""] for i in range(1, 13)]
    rows += [["Вкупно кеси / нето маса | Total bags / net mass", "", "", ""]]
    w = widths(rows, mode="full")
    assert w[0] == min(w), "the № column became wider than a write-in column"
    assert w[0] <= 2.5


def test_write_in_columns_share_the_spare_width_evenly():
    rows = [["Позиција | Position", "Кеси | Bags", "Нето маса (g) | Net mass (g)",
             "Бруто маса (g) | Gross mass (g)"],
            ["C01", "", "", ""], ["C02", "", "", ""]]
    w = widths(rows, mode="full")
    assert sum(w) == pytest.approx(R.PAGE_W, abs=0.05)
    assert max(w[1:]) - min(w[1:]) < 0.05, f"write-in columns are uneven: {w[1:]}"
    assert w[0] < w[1]


def test_signoff_role_column_takes_the_remainder():
    rows = [["Улога | Role", "Име | Name", "Датум/време | Date/time", "Потпис | Signature"],
            ["Независно проверил (кеси, серија, маса, етикета) | "
             "Independently verified (bags, batch, mass, label)", "", "", ""]]
    w = widths(rows, mode="full")
    assert w[1] == pytest.approx(5.2, abs=0.05)
    assert w[3] == pytest.approx(3.8, abs=0.05)
    assert w[2] == pytest.approx(2.6, abs=0.05)
    assert w[0] == pytest.approx(R.PAGE_W - 11.6, abs=0.05)


def test_placeholder_cells_are_still_write_in():
    """'[ ]' and '___' mark where a person writes; they are not printed content."""
    rows = [["Улога | Role", "Име | Name", "Датум | Date", "Потпис | Signature"],
            ["Одобрил и потпишал (QP) | Approved and signed (QP)", "[ ]", "[ ]", "[ ]"]]
    w = widths(rows, mode="full")
    assert w[1] == pytest.approx(5.2, abs=0.05) and w[3] == pytest.approx(3.8, abs=0.05)
