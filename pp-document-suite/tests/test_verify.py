"""pp_verify: what the gate reads and what it reports (review 19.10.2026)."""
import contextlib
import io
import os
import sys

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))
pytest.importorskip("docx")
import pp_verify  # noqa: E402
from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402


def _run(path, *args):
    buf = io.StringIO(); old = sys.argv; sys.argv = ["pp_verify", str(path), *args]
    try:
        with contextlib.redirect_stdout(buf), pytest.raises(SystemExit) as e:
            pp_verify.main()
    finally:
        sys.argv = old
    return e.value.code, buf.getvalue()


def test_nested_table_text_counts(tmp_path):
    """A table inside a cell (SOP §7 records) was invisible to the word count and the bilingual check."""
    d = Document(); d.add_paragraph("Only Latin here")
    inner = d.add_table(rows=1, cols=1).cell(0, 0).add_table(rows=1, cols=1)
    inner.cell(0, 0).text = "Само кирилица"
    p = tmp_path / "n.docx"; d.save(p)
    r = pp_verify.analyse(str(p))
    assert r["cyr"] and r["words"] == 5
    assert _run(p, "--require-bilingual", "true")[0] == 0


def test_an_inherited_font_size_is_a_warning_not_a_silent_pass(tmp_path):
    d = Document(); d.add_paragraph("Текст | Text"); p = tmp_path / "i.docx"; d.save(p)
    code, out = _run(p)
    assert code == 0 and "min font: inherited (styles not checked)  [WARN]" in out and "None pt" not in out
    d = Document(); d.add_paragraph().add_run("Текст | Text").font.size = Pt(11); p2 = tmp_path / "s.docx"; d.save(p2)
    code, out = _run(p2)
    assert code == 0 and "min font 11.0 pt  [OK]" in out
    d = Document(); d.add_paragraph().add_run("Текст | Text").font.size = Pt(5); p3 = tmp_path / "f.docx"; d.save(p3)
    assert _run(p3)[0] == 1
