"""pp_render: the Gotenberg request (updateIndexes, basic auth), the fallback to local
LibreOffice, and a clear error when neither renderer is available."""
import os
import sys
import threading
from http.server import BaseHTTPRequestHandler, HTTPServer

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))
import pp_render  # noqa: E402


class _Stub(BaseHTTPRequestHandler):
    seen = {}
    reply = b"%PDF-1.7 /BaseFont/AAAAAA+Calibri stub"

    def do_POST(self):
        n = int(self.headers["Content-Length"])
        _Stub.seen = {"path": self.path, "auth": self.headers.get("Authorization"),
                      "ctype": self.headers["Content-Type"], "body": self.rfile.read(n)}
        self.send_response(200); self.send_header("Content-Type", "application/pdf"); self.end_headers()
        self.wfile.write(_Stub.reply)

    def log_message(self, *a):
        pass


@pytest.fixture
def gotenberg(monkeypatch):
    srv = HTTPServer(("127.0.0.1", 0), _Stub)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    monkeypatch.setenv("GOTENBERG_URL", f"http://127.0.0.1:{srv.server_port}")
    yield _Stub
    srv.shutdown()


@pytest.fixture
def docx(tmp_path):
    from docx import Document          # a real document: the renderer now checks the fonts it declares
    d = Document(); r = d.add_paragraph().add_run("Текст | Text"); r.font.name = "Calibri"
    p = tmp_path / "in.docx"; d.save(p); return str(p)


def test_gotenberg_request(gotenberg, docx, tmp_path, monkeypatch):
    monkeypatch.setenv("GOTENBERG_USERNAME", "pp"); monkeypatch.setenv("GOTENBERG_PASSWORD", "s")
    out = tmp_path / "out.pdf"
    assert pp_render.render_pdf(docx, str(out)) == "gotenberg"
    s = gotenberg.seen
    assert s["path"] == "/forms/libreoffice/convert" and s["ctype"].startswith("multipart/form-data")
    assert b'name="updateIndexes"\r\n\r\ntrue' in s["body"] and b'filename="in.docx"' in s["body"]
    assert s["auth"] == "Basic cHA6cw=="          # pp:s
    assert out.read_bytes() == b"%PDF-1.7 /BaseFont/AAAAAA+Calibri stub"


def test_non_pdf_reply_is_an_error(gotenberg, docx, tmp_path, monkeypatch):
    monkeypatch.setattr(_Stub, "reply", b"<html>error</html>")
    with pytest.raises(pp_render.RenderError):
        pp_render.render_gotenberg(docx, str(tmp_path / "o.pdf"))


def test_auto_falls_back_to_local_when_gotenberg_unreachable(docx, tmp_path, monkeypatch):
    monkeypatch.setenv("GOTENBERG_URL", "http://127.0.0.1:9")
    def local(d, p):
        open(p, "wb").write(b"%PDF-1.7 /BaseFont/AAAAAA+Calibri-Bold")
        return "local"
    monkeypatch.setattr(pp_render, "render_local", local)
    assert pp_render.render_pdf(docx, str(tmp_path / "o.pdf")) == "local"


def test_no_renderer_names_both_reasons(docx, tmp_path, monkeypatch):
    monkeypatch.delenv("GOTENBERG_URL", raising=False)
    monkeypatch.setattr(pp_render.shutil, "which", lambda name: None)
    with pytest.raises(pp_render.RenderError, match="no local LibreOffice"):
        pp_render.render_pdf(docx, str(tmp_path / "o.pdf"))


def test_a_local_timeout_is_a_render_error_and_kills_the_group(docx, tmp_path, monkeypatch):
    """A timeout escaped as subprocess.TimeoutExpired (outside the RenderError contract) and killed only the
    soffice wrapper, leaving soffice.bin running on a deleted profile (review 14.10.2026)."""
    import shutil
    import subprocess
    fake = tmp_path / "bin"
    fake.mkdir()
    (fake / "soffice").write_text("#!/bin/sh\nsleep 30 &\nwait\n")
    (fake / "soffice").chmod(0o755)
    monkeypatch.setenv("PATH", f"{fake}{os.pathsep}{os.environ['PATH']}")
    assert shutil.which("soffice") == str(fake / "soffice")
    with pytest.raises(pp_render.RenderTimeout):
        pp_render.render_local(docx, str(tmp_path / "o.pdf"), timeout=1)
    left = subprocess.run(["pgrep", "-f", "sleep 30"], capture_output=True, text=True).stdout.split()
    assert not left, "the renderer's child process outlived the timeout"


# ---------------------------------------------------------------- font substitution (review 14.10.2026)
def test_a_pdf_printed_with_substitute_fonts_is_refused_and_removed(gotenberg, docx, tmp_path, monkeypatch):
    """Every PDF rendered where Calibri and Arial Narrow are not installed carried Carlito and Inter instead,
    against the 10.10.2026 ruling: never accept a substitute."""
    monkeypatch.setattr(_Stub, "reply", b"%PDF-1.7 /BaseFont/AAAAAA+Carlito-Regular /BaseFont/BAAAAA+DejaVuSans")
    out = tmp_path / "out.pdf"
    with pytest.raises(pp_render.RenderError, match="substitute fonts.*Calibri"):
        pp_render.render_pdf(docx, str(out))
    assert not out.exists(), "a refused PDF is never left behind"


def test_a_quote_in_the_docx_path_never_reaches_the_macro(tmp_path, monkeypatch):
    """The paths are spliced into a Basic call as "..." literals; a quote in the name ended the literal (review 19.10.2026)."""
    import re
    import shutil
    seen = {}

    class Popen:
        def __init__(self, args, **kw):
            seen["macro"] = args[-1]
            m = re.fullmatch(r'macro:///Standard\.Module1\.ToPdf\("([^"]*)","([^"]*)"\)', args[-1])
            assert m, args[-1]
            seen["in"] = m.group(1)
            open(m.group(2), "wb").write(b"%PDF-1.7 /BaseFont/AAAAAA+Calibri")
        def wait(self, timeout=None): return 0
    monkeypatch.setattr(pp_render.subprocess, "Popen", Popen)
    monkeypatch.setattr(pp_render.shutil, "which", lambda name: "/usr/bin/soffice")
    from docx import Document
    d = Document(); d.add_paragraph().add_run("Текст | Text").font.name = "Calibri"
    odd = tmp_path / 'say "hi".docx'; d.save(odd)
    assert pp_render.render_local(str(odd), str(tmp_path / "o.pdf")) == "local"
    assert '"' not in seen["in"] and seen["in"] != str(odd)          # a safe copy, not the quoted path
    assert (tmp_path / "o.pdf").read_bytes().startswith(b"%PDF")
