"""Font substitution: what a .docx declares against what its PDF embeds (review 14.10.2026).

Head of QC, 10.10.2026: fidelity comes from the real fonts; a missing font is installed, never substituted."""
import os
import sys
import zlib

import pytest

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "scripts"))
docx = pytest.importorskip("docx", reason="python-docx is the engine's own dependency")
from docx import Document  # noqa: E402
import pp_assets  # noqa: E402


def _doc(tmp_path, runs):
    d = Document()
    for text, face in runs:
        r = d.add_paragraph().add_run(text)
        if face:
            r.font.name = face
    p = tmp_path / "t.docx"; d.save(p); return str(p)


def test_declared_faces_follow_runs_then_defaults(tmp_path):
    p = _doc(tmp_path, [("Код | Code", "Arial Narrow"), ("Текст", "Calibri"), ("   ", "Times New Roman")])
    faces = pp_assets.declared_faces(p)
    assert {"Arial Narrow", "Calibri"} <= faces and "Times New Roman" not in faces   # a blank run asks for nothing


def test_pdf_fonts_read_plain_and_compressed_objects():
    plain = b"%PDF-1.7 /BaseFont/BAAAAA+Calibri-Bold /BaseFont /ArialNarrow"
    assert pp_assets.pdf_fonts(plain) == {"Calibri-Bold", "ArialNarrow"}
    hidden = zlib.compress(b"<< /Type /Font /BaseFont /CAAAAA+Calibri-Italic >>")
    objstm = b"%PDF-1.7 1 0 obj << /Type /ObjStm /Filter /FlateDecode >>\nstream\n" + hidden + b"\nendstream"
    assert "Calibri-Italic" in pp_assets.pdf_fonts(objstm)


def test_substitution_is_named_and_the_real_face_passes(tmp_path):
    p = _doc(tmp_path, [("Код | Code", "Arial Narrow"), ("Текст", "Calibri")])
    subs = pp_assets.font_substitutions(p, b"%PDF /BaseFont/AAAAAA+Carlito-Regular /BaseFont/BAAAAA+Inter-Regular")
    assert len(subs) == 2 and any(s.startswith("Arial Narrow") for s in subs)
    real = b"%PDF /BaseFont/AAAAAA+Calibri /BaseFont/BAAAAA+ArialNarrow-Bold /BaseFont/CAAAAA+DejaVuSans"
    assert pp_assets.font_substitutions(p, real) == []          # a symbol fallback alongside is allowed


def test_macedonian_in_a_latin_only_display_face_is_flagged(tmp_path, monkeypatch):
    monkeypatch.setattr(pp_assets, "_resolve", lambda face: "/fake.ttf")
    monkeypatch.setattr(pp_assets, "_cmap", lambda path: None)  # whatever the machine substitutes
    bad = pp_assets.audit_docx(_doc(tmp_path, [("Сорта Tangie", "Orbitron")]))
    assert bad and bad[0][0] == "Orbitron" and "С" in bad[0][1]
