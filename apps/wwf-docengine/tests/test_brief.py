# The Document Brief (Head of QC, 18.10.2026): one workflow, every answer chosen / suggested / locked, a scope change
# marks decided answers stale instead of rewriting them, every change is an event, build from any page.
import asyncio
import base64
import io
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app import brief as B  # noqa: E402
from app import db, main, pipeline, sources  # noqa: E402
from app.config import settings  # noqa: E402
from app.main import app  # noqa: E402

KEY = "k"


def fresh(**kw):
    return B.new_brief("sop", {"title_en": "Sampling", "code": "QCSOP011"}, **kw)


def test_a_change_needs_the_current_rev_and_counts_one_rev():
    b = fresh()
    b1, ev = B.apply_patch(b, {"rev": 0, "set": {"doc.purpose": "Sample finished product"}}, "bn")
    assert b1["rev"] == 1 and b1["fields"]["doc.purpose"]["state"] == "chosen" and ev[0]["op"] == "set"
    with pytest.raises(B.BriefConflict):
        B.apply_patch(b1, {"rev": 0, "set": {"purpose": "x"}})            # an old tab cannot overwrite
    same, ev = B.apply_patch(b1, {"rev": 1, "set": {"doc.purpose": "Sample finished product"}})
    assert same["rev"] == 1 and ev == []                                     # nothing changed, no new rev


def test_a_locked_answer_is_never_overwritten_until_unlocked():
    b, _ = B.apply_patch(fresh(), {"rev": 0, "set": {"sample_types": "Finished product"}, "lock": ["sample_types"]})
    with pytest.raises(B.BriefConflict):
        B.apply_patch(b, {"rev": 1, "set": {"sample_types": "Incoming"}})
    b2, ev = B.apply_patch(b, {"rev": 1, "suggest": {"sample_types": "Incoming"}})
    assert b2["fields"]["sample_types"]["value"] == "Finished product" and ev == []   # a suggestion never wins
    b3, ev = B.apply_patch(b, {"rev": 1, "unlock": ["sample_types"], "set": {"sample_types": "Incoming"}})
    assert [e["op"] for e in ev] == ["unlock", "set"] and b3["fields"]["sample_types"]["state"] == "chosen"


def test_a_scope_change_marks_decided_answers_stale_but_not_locked_ones():
    b, _ = B.apply_patch(fresh(), {"rev": 0, "set": {"doc.scope_in": "FP", "focus": "Potency", "retention": "Other (state in §7)"}, "lock": ["retention"]})
    b2, ev = B.apply_patch(b, {"rev": 1, "set": {"doc.scope_in": "FP and intermediates"}})
    assert b2["fields"]["focus"]["stale"] and not b2["fields"]["retention"]["stale"] and not b2["fields"]["doc.scope_in"]["stale"]
    assert {e["key"] for e in ev if e["op"] == "stale"} == {"focus"}
    b3, _ = B.apply_patch(b2, {"rev": 2, "set": {"focus": "Potency"}})                   # re-confirming clears it
    assert not b3["fields"]["focus"]["stale"]


def test_a_build_from_an_early_page_lists_what_was_not_decided():
    b, _ = B.apply_patch(B.new_brief("sop_annexes", {"title_en": "Sampling"}),
                         {"rev": 0, "set": {"doc.purpose": "Sample FP"}, "suggest": {"doc.scope_in": "FP"},
                          "annexes": [{"title_en": "Sampling plan"}]})
    p = B.to_workflow(b, [{"name": "old.docx", "text": "x" * 20_000}, {"name": "scan.pdf", "text": ""}])
    assert p["questionnaire"] == "sop_qc" and p["meta"]["title_mk"] == "Sampling"
    assert "title_mk (copied from title_en)" in p["assumed"] and "doc.scope_in (suggested, not confirmed)" in p["assumed"]
    assert "focus" in p["assumed"] and "Purpose: Sample FP" in p["instructions"]
    assert p["package_annexes"] == [{"title_en": "Sampling plan", "title_mk": "", "kind": "form"}]
    assert [t["name"] for t in p["source_texts"]] == ["old.docx"] and len(p["source_texts"][0]["text"]) == 12_000
    with pytest.raises(B.BriefError):
        B.to_workflow(B.new_brief("sop"), [])                                 # a title is the one thing required


def _docx(text: str) -> bytes:
    from docx import Document
    d, buf = Document(), io.BytesIO()
    d.add_paragraph(text)
    t = d.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "ЦЕЛ", "PURPOSE"
    d.save(buf)
    return buf.getvalue()


def test_uploaded_sources_keep_their_text():
    assert "Мострирање" in sources.extract("a.docx", _docx("Мострирање на примероци")) and "ЦЕЛ | PURPOSE" in sources.extract("a.docx", _docx("x"))
    assert sources.extract("a.md", "# Title\nbody".encode()) == "# Title\nbody"
    assert '"a": 1' in sources.extract("a.json", b'{"a":1}')
    for name, blob in (("a.exe", b"x"), ("a.json", b"{not json"), ("a.pdf", b"not a pdf")):
        with pytest.raises(sources.SourceError):
            sources.extract(name, blob)


def test_the_brief_reaches_the_authors_with_its_source_files():
    text = pipeline._brief("sop_qc", {"x": "y"}, "DOCUMENT BRIEF", ["old.docx"], [{"name": "old.docx", "text": "Old SOP"}])
    assert "SOURCE FILE «old.docx»" in text and "Old SOP" in text


@pytest.fixture
def api(monkeypatch):
    """The brief routes over an in-memory store (the SQL itself is covered by test_db.py against Postgres)."""
    monkeypatch.setattr(settings, "api_key", KEY)
    store, srcs, started = {}, {}, []

    async def create(b, by=""):
        bid = "11111111-1111-1111-1111-11111111111" + str(len(store))
        store[bid] = {**b, "id": bid}
        srcs[bid] = []
        return bid

    async def get(bid):
        return {**store[bid], "sources": [{k: v for k, v in s.items() if k != "text"} for s in srcs[bid]]} if bid in store else None

    async def save(b, old, events):
        if store[b["id"]]["rev"] != old:
            return False
        store[b["id"]] = {k: v for k, v in b.items() if k != "sources"}
        return True

    async def add(bid, rev, actor, s):
        if store[bid]["status"] != "open":
            return None
        srcs[bid].append({**s, "id": "22222222-2222-2222-2222-22222222222" + str(len(srcs[bid])), "chars": len(s["text"])})
        store[bid]["rev"] += 1
        return srcs[bid][-1]["id"]

    async def lst(bid, with_text=True):
        return srcs[bid]

    async def jcreate(kind, payload, by=""):
        started.append(payload)
        return "33333333-3333-3333-3333-333333333333"

    async def nothing(*a):
        return None

    async def taken():
        return {"QCSOP026"}

    async def of_brief(bid):
        return set()

    async def jupdate(jid, **f):
        failed.append((jid, f))

    failed = []
    for name, fn in (("brief_create", create), ("brief_get", get), ("brief_save", save), ("brief_source_add", add),
                     ("brief_sources", lst), ("job_create", jcreate), ("codes_in_use", taken),
                     ("codes_of_brief", of_brief), ("job_update", jupdate)):
        monkeypatch.setattr(db, name, fn)
    monkeypatch.setattr(db, "ready", lambda: True)
    monkeypatch.setattr(main, "run_workflow", nothing)
    monkeypatch.setattr(main.LettaClient, "configured", property(lambda self: True))
    with TestClient(app) as c:
        yield c, store, started


def h():
    return {"X-API-Key": KEY}


def test_one_workflow_from_first_page_to_build(api):
    c, store, started = api
    assert c.post("/briefs", json={"package": "sop"}).status_code == 401
    b = c.post("/briefs", json={"package": "sop_annexes", "meta": {"title_en": "Sampling", "code": "qcsop 11"}}, headers=h()).json()
    assert b["meta"]["code"] == "QCSOP011" and b["rev"] == 0
    b = c.patch(f"/briefs/{b['id']}", json={"rev": 0, "set": {"doc.purpose": "Sample FP"}, "lock": ["doc.purpose"],
                                            "annexes": [{"title_en": "Sampling plan"}]}, headers=h()).json()
    assert b["rev"] == 1 and b["fields"]["doc.purpose"]["state"] == "locked"
    assert c.patch(f"/briefs/{b['id']}", json={"rev": 0, "set": {"doc.use": "x"}}, headers=h()).status_code == 409
    assert c.patch(f"/briefs/{b['id']}", json={"rev": 1, "set": {"doc.purpose": "x"}}, headers=h()).status_code == 409
    up = c.post(f"/briefs/{b['id']}/sources", json={"name": "old.docx", "content_b64": base64.b64encode(_docx("Old SOP text")).decode()},
                headers=h()).json()
    assert up["chars"] > 0 and up["brief"]["rev"] == 2 and up["brief"]["sources"][0]["name"] == "old.docx"
    assert c.post(f"/briefs/{b['id']}/sources", json={"name": "x.exe", "content_b64": "eA=="}, headers=h()).status_code == 400
    r = c.post(f"/briefs/{b['id']}/build", json={"rev": 2}, headers=h()).json()
    assert r["code"] == "QCSOP011" and r["status"] == "queued" and "doc.scope_in" in r["assumed"]
    assert started[0]["package_annexes"][0]["title_en"] == "Sampling plan" and started[0]["source_texts"][0]["name"] == "old.docx"
    assert store[b["id"]]["job_ids"] == ["33333333-3333-3333-3333-333333333333"] and store[b["id"]]["rev"] == 3


def test_a_build_without_a_code_takes_the_next_free_one(api):
    c, store, started = api
    b = c.post("/briefs", json={"package": "sop", "meta": {"title_mk": "Мострирање"}}, headers=h()).json()
    r = c.post(f"/briefs/{b['id']}/build", json={}, headers=h()).json()
    assert r["code"] == "QCSOP027"                                           # QCSOP026 held by an open job


def test_package_annexes_start_after_the_sop_and_number_under_it(monkeypatch):
    made, started = [], []

    async def taken():
        return {"QCSOP011A01"}

    async def jcreate(kind, payload, by=""):
        made.append(payload)
        return f"job-{len(made)}"

    async def fake_run(jid):
        started.append(jid)

    monkeypatch.setattr(db, "codes_in_use", taken)
    monkeypatch.setattr(db, "job_create", jcreate)
    monkeypatch.setattr(pipeline, "run_workflow", fake_run)
    p = {"package_annexes": [{"title_en": "Plan"}, {"title_mk": "Дневник", "kind": "log"}], "instructions": "BRIEF"}

    async def go():
        out = await pipeline._start_package_annexes("sop-job", p, {"code": "QCSOP011", "version": "03"})
        await asyncio.sleep(0)
        return out
    out = asyncio.run(go())
    assert [o["code"] for o in out] == ["QCSOP011A10", "QCSOP011A11"]     # A01–A09 held by the register and jobs
    assert made[0]["parent_sop"] == "QCSOP011" and made[1]["meta"]["title_en"] == "Дневник" and started == ["job-1", "job-2"]


# ---------- phase 3: the questions react to the answers (question_rules) and the advisor proposes (advisor) ----------
from app import advisor, question_rules as QR  # noqa: E402


def test_an_answer_recommends_options_of_later_questions_and_says_why():
    b, _ = B.apply_patch(fresh(), {"rev": 0, "set": {"focus": "Potency"}})
    g = QR.guide(b, {"techniques": ["HPLC-DAD", "GC-MS"], "reference_standards": ["CRS", "SST", "Internal"]})
    assert g["questions"]["techniques"]["recommended"] == [{"value": "HPLC-DAD", "because": ["focus: Potency"]}]
    assert [r["value"] for r in g["questions"]["reference_standards"]["recommended"]] == ["CRS", "SST"]
    assert g["questions"]["techniques"]["depends_on"] == ["focus", "sample_types"]


def test_a_changed_answer_makes_only_its_dependents_stale_and_never_a_locked_one():
    b, _ = B.apply_patch(fresh(), {"rev": 0, "set": {"focus": "Potency", "techniques": ["HPLC-DAD"],
                                                       "acceptance": ["LOD NMT 12 %"], "performer": "QC Analyst"}})
    b, _ = B.apply_patch(b, {"rev": 1, "lock": ["acceptance"]})
    b, ev = B.apply_patch(b, {"rev": 2, "set": {"focus": "Contaminants"}})
    assert b["fields"]["techniques"]["stale"] and not b["fields"]["performer"]["stale"]
    assert b["fields"]["acceptance"]["state"] == "locked" and not b["fields"]["acceptance"]["stale"]
    assert [e["key"] for e in ev if e["op"] == "stale"] == ["techniques"]


def test_a_transport_title_points_to_the_transport_question_bank():
    b = B.new_brief("sop", {"title_en": "Transport of bulk cannabis"})
    assert QR.guide(b, {})["questionnaire_hint"] == "sop_transport"
    assert QR.guide(B.new_brief("sop", {"title_en": "Sampling"}), {})["questionnaire_hint"] == ""


def test_the_advisor_answer_is_kept_to_the_questions_asked():
    asked = {"techniques": {"label": "Techniques", "multi": True, "options": ["HPLC-DAD"]},
             "doc.purpose": {"label": "Purpose", "multi": False, "options": []}}
    raw = {"proposals": [
        {"key": "techniques", "options": [{"value": "GC-MS", "why": "terpenes", "source": "E1"}, {"value": ""}],
         "pick": ["GC-MS", "invented"]},
        {"key": "doc.purpose", "options": [{"value": "To define sampling", "why": "x"}], "pick": "not offered"},
        {"key": "made_up", "options": [{"value": "x"}]}]}
    out = advisor.clean(raw, asked)
    assert set(out) == {"techniques", "doc.purpose"}
    assert out["techniques"] == {"options": [{"value": "GC-MS", "why": "terpenes", "source": "E1"}], "pick": ["GC-MS"]}
    assert out["doc.purpose"]["pick"] == "" and out["doc.purpose"]["options"][0]["source"] == "house practice"
    assert advisor._first_json('Here: ```json\n{"proposals": []}\n```') == {"proposals": []}
    # review 19.10.2026: the greedy match ran to the LAST brace, so an answer that mentioned another {…} was no JSON
    assert advisor._first_json('Here: {"proposals": []} and also {"x": 1}') == {"proposals": []}
    assert advisor._first_json('no object {here, then: {"proposals": [{"key": "a"}]} done') == {"proposals": [{"key": "a"}]}
    with pytest.raises(ValueError):
        advisor._first_json("nothing here")
    text = advisor.prompt(fresh(), asked, "[E1] house", "[R1] reg", [{"name": "a.md", "text": "src"}])
    assert "techniques: Techniques (several may be chosen)" in text and "SOURCE FILE «a.md»" in text


def test_suggest_keeps_proposals_and_never_overwrites_a_decision(api, monkeypatch):
    c, store, _ = api
    seen = {}

    async def fake(brief, asked, sources, client=None, rag=None):
        seen["asked"] = set(asked)
        return {"proposals": {k: {"options": [{"value": "HPLC-DAD", "why": "potency", "source": "E1"}],
                                  "pick": ["HPLC-DAD"] if asked[k]["multi"] else "HPLC-DAD"} for k in asked},
                "notes": []}
    monkeypatch.setattr(advisor, "propose", fake)
    b = c.post("/briefs", json={"package": "sop", "meta": {"title_en": "Potency"}}, headers=h()).json()
    assert "guide" in b and b["proposals"] == {}
    b = c.patch(f"/briefs/{b['id']}", json={"rev": 0, "set": {"focus": "Potency"}}, headers=h()).json()
    assert b["guide"]["questions"]["techniques"]["recommended"][0]["value"] == "HPLC-DAD"
    assert c.post(f"/briefs/{b['id']}/suggest", json={"rev": 0}, headers=h()).status_code == 409
    assert c.post(f"/briefs/{b['id']}/suggest", json={"rev": 1, "keys": ["nope"]}, headers=h()).status_code == 400
    r = c.post(f"/briefs/{b['id']}/suggest", json={"rev": 1, "keys": ["focus", "techniques"]}, headers=h()).json()
    assert seen["asked"] == {"techniques"}                     # focus is decided: not asked about
    nb = r["brief"]
    assert nb["fields"]["focus"] == {**b["fields"]["focus"]}
    assert nb["fields"]["techniques"]["state"] == "suggested" and nb["fields"]["techniques"]["value"] == ["HPLC-DAD"]
    assert nb["proposals"]["techniques"]["options"][0]["source"] == "E1" and nb["rev"] == 2


# ---------- phase 4: one package view and one merged PDF ----------
from app import package as PKG  # noqa: E402


def _pdf(pages: int) -> bytes:
    import io
    from pypdf import PdfWriter
    w = PdfWriter()
    for _ in range(pages):
        w.add_blank_page(width=595, height=842)
    out = io.BytesIO()
    w.write(out)
    return out.getvalue()


def test_the_package_merges_the_sop_then_its_annexes():
    from pypdf import PdfReader
    import io
    assert len(PdfReader(io.BytesIO(PKG.merge([_pdf(3), _pdf(1), _pdf(2)]))).pages) == 6


def test_a_package_is_complete_only_when_every_member_is_built(tmp_path):
    f = tmp_path / "QCSOP026.docx"
    f.write_bytes(b"x")
    done = PKG.member({"status": "done", "payload": {"meta": {"code": "QCSOP026"}}, "result": {"document_id": "d"}},
                      {"path": str(f)}, "j1")
    review = PKG.member({"status": "awaiting_review", "payload": {"meta": {"code": "QCSOP026A01"}},
                         "result": {"artifact": {"path": str(f)}, "review_reason": "§6A FIX"}}, None, "j2")
    running = PKG.member({"status": "running", "stage": "generate", "payload": {"meta": {"code": "QCSOP026A02"}},
                          "result": {}}, None, "j3")
    assert done["file"] and review["file"] and not running["file"]
    s = PKG.summary([done, review, running])
    assert not s["complete"] and not s["pdf_ready"] and s["running"] == ["QCSOP026A02"]
    assert s["awaiting_review"] == ["QCSOP026A01"]
    assert PKG.summary([done, review])["pdf_ready"]
    assert PKG.member(None, None, "gone")["status"] == "missing"
    assert PKG.pending({"title_en": "Log"}, "")["status"] == "pending"


def test_package_route_lists_the_build_its_annexes_and_merges_them(api, monkeypatch, tmp_path):
    c, store, _ = api
    from app import db, main
    f = tmp_path / "x.docx"
    f.write_bytes(b"x")
    jobs = {
        "33333333-3333-3333-3333-333333333333": {"status": "done", "payload": {"meta": {"code": "QCSOP026"},
                                                  "package_annexes": [{"title_en": "Log"}, {"title_en": "Checklist"}]},
                                                  "result": {"document_id": "44444444-4444-4444-4444-444444444444",
                                                             "package_jobs": [{"code": "QCSOP026A01", "job_id": "a1"}]}},
        "a1": {"status": "awaiting_review", "payload": {"meta": {"code": "QCSOP026A01", "title_en": "Log"}},
               "result": {"artifact": {"path": str(f)}}}}

    async def job_get(j):
        return jobs.get(j)

    async def doc_get(d):
        return {"path": str(f)}

    rendered = []

    async def render(path, label):
        rendered.append(label)
        return _pdf(1)
    monkeypatch.setattr(db, "job_get", job_get)
    monkeypatch.setattr(db, "document_get", doc_get)
    monkeypatch.setattr(main, "_render_pdf", render)
    monkeypatch.setattr(main.settings, "gotenberg_url", "http://render")
    b = c.post("/briefs", json={"package": "sop_annexes", "meta": {"title_en": "Sampling"}}, headers=h()).json()
    assert c.get(f"/briefs/{b['id']}/package", headers=h()).json()["members"] == []
    store[b["id"]]["job_ids"] = ["33333333-3333-3333-3333-333333333333"]
    p = c.get(f"/briefs/{b['id']}/package", headers=h()).json()
    assert [m["code"] or m["title_en"] for m in p["members"]] == ["QCSOP026", "QCSOP026A01", "Checklist"]
    assert [m["status"] for m in p["members"]] == ["done", "awaiting_review", "pending"]
    assert not p["complete"] and "file" not in p["members"][0] and p["members"][0]["has_file"]
    r = c.get(f"/briefs/{b['id']}/package.pdf", headers=h())
    assert r.status_code == 409 and "Checklist" in r.json()["detail"]
    jobs["33333333-3333-3333-3333-333333333333"]["payload"]["package_annexes"] = [{"title_en": "Log"}]
    r = c.get(f"/briefs/{b['id']}/package.pdf", headers=h())
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert rendered == ["QCSOP026", "QCSOP026A01"] and 'QCSOP026_package.pdf' in r.headers["content-disposition"]


# ---------- the approved, in-use house documents are a source, not only a style sample (19.10.2026) ----------
from app import house  # noqa: E402


def test_the_in_use_list_keeps_the_newest_version_of_each_code(monkeypatch):
    import asyncio
    from app import db

    async def refs(kind=""):
        assert kind == "approved"
        return [{"id": "1", "code": "QCSOP011", "version": "01", "title_en": "Sampling", "title_mk": "Земање мостри"},
                {"id": "2", "code": "QCSOP011", "version": "02", "title_en": "Sampling", "title_mk": "Земање мостри"},
                {"id": "3", "code": "QCT017", "version": "01", "title_en": "Sample log", "title_mk": ""}]

    async def get(rid):
        return {"markdown": f"text {rid}"}
    monkeypatch.setattr(db, "ready", lambda: True)
    monkeypatch.setattr(db, "reference_list", refs)
    monkeypatch.setattr(db, "reference_get", get)
    docs = asyncio.run(house.in_use_list())
    assert [(d["code"], d["version"]) for d in docs] == [("QCSOP011", "02"), ("QCT017", "01")]
    block = house.list_block(docs)
    assert "never invent a code" in block and "- QCSOP011 v02 — Sampling | Земање мостри" in block
    assert "- QCT017 v01 — Sample log" in block
    assert asyncio.run(house.approved_text("QCSOP011")) == {"markdown": "text 2", "version": "02", "id": "2"}
    assert asyncio.run(house.approved_text("QCSOP099")) is None



# ---------- the review of 19.10.2026: builds, suggestions and numbering under real use ----------
def test_a_build_racing_an_edit_fails_its_job_instead_of_orphaning_it(api, monkeypatch):
    c, store, started = api
    from app import db
    b = c.post("/briefs", json={"package": "sop", "meta": {"title_en": "Sampling"}}, headers=h()).json()
    real_save = db.brief_save

    async def lose(new, old, events):          # an edit commits between the build's read and its save
        store[new["id"]]["rev"] += 1
        return await real_save(new, old, events)
    monkeypatch.setattr(db, "brief_save", lose)
    r = c.post(f"/briefs/{b['id']}/build", json={"rev": 0}, headers=h())
    assert r.status_code == 409 and "reload" in r.json()["detail"]
    assert store[b["id"]]["job_ids"] == []


def test_a_closed_brief_is_not_built_suggested_or_given_sources(api, monkeypatch):
    c, store, _ = api
    b = c.post("/briefs", json={"package": "sop", "meta": {"title_en": "Sampling"}}, headers=h()).json()
    assert c.patch(f"/briefs/{b['id']}", json={"rev": 0, "status": "archived"}, headers=h()).status_code == 400
    b = c.patch(f"/briefs/{b['id']}", json={"rev": 0, "status": "closed"}, headers=h()).json()
    assert c.post(f"/briefs/{b['id']}/build", json={"rev": 1}, headers=h()).status_code == 409
    assert c.post(f"/briefs/{b['id']}/suggest", json={"rev": 1}, headers=h()).status_code == 409
    assert c.post(f"/briefs/{b['id']}/sources", json={"name": "a.md", "content_b64": "eA=="}, headers=h()).status_code == 409


def test_an_unknown_department_is_a_400_not_a_500(api):
    c, _, _ = api
    b = c.post("/briefs", json={"package": "sop", "meta": {"title_en": "Sampling", "department": "Quality"}}, headers=h()).json()
    r = c.post(f"/briefs/{b['id']}/build", json={"rev": 0}, headers=h())
    assert r.status_code == 400 and "cannot number" in r.json()["detail"]


def test_only_questions_of_the_brief_are_stored_and_answers_are_bounded():
    with pytest.raises(B.BriefError):
        B.apply_patch(fresh(), {"rev": 0, "set": {"made_up": "x"}})
    with pytest.raises(B.BriefError):
        B.apply_patch(fresh(), {"rev": 0, "set": {"doc.purpose": "x" * (B.MAX_VALUE_CHARS + 1)}})
    b, _ = B.apply_patch(fresh(), {"rev": 0, "set": {"focus": "Potency"}})
    b, _ = B.apply_patch(b, {"rev": 1, "questionnaire": "sop_transport"})
    b, ev = B.apply_patch(b, {"rev": 2, "set": {"focus": None}})      # an answer of the old bank can still be cleared
    assert "focus" not in b["fields"] and ev[0]["op"] == "clear"


def test_the_not_decided_instruction_survives_a_long_brief():
    b, _ = B.apply_patch(fresh(), {"rev": 0, "set": {k: "y" * 19_000 for k in B.TEXT_FIELDS}})
    out = B.to_workflow(b, [])["instructions"]
    assert out.startswith("Not decided by the requester") and len(out) <= 12_000


def test_package_annexes_number_from_a01_on_a_rebuild_and_by_position_when_unconfirmed(monkeypatch):
    import asyncio
    from app import db, pipeline
    created = []

    async def in_use():
        return {"QCSOP026", "QCSOP026A01", "QCSOP026A02", "QASOPXXXA01"}   # A01, A02: the first build's own

    async def of_brief(bid):
        return {"QCSOP026", "QCSOP026A01", "QCSOP026A02"} if bid == "brief-1" else set()

    async def jcreate(kind, payload, by=""):
        created.append(payload["meta"]["code"])
        return f"j{len(created)}"

    async def nothing(jid):
        return None
    monkeypatch.setattr(db, "codes_in_use", in_use)
    monkeypatch.setattr(db, "codes_of_brief", of_brief)
    monkeypatch.setattr(db, "job_create", jcreate)
    monkeypatch.setattr(pipeline, "run_workflow", nothing)
    p = {"package_annexes": [{"title_en": "Log"}, {"title_en": "Checklist"}], "brief_id": "brief-1"}

    async def go():
        await pipeline._start_package_annexes("sop", p, {"code": "QCSOP026"})
        await pipeline._start_package_annexes("sop", {**p, "brief_id": "brief-2"}, {"code": "QASOPXXX"})
        await asyncio.sleep(0)
    asyncio.run(go())
    assert created == ["QCSOP026A01", "QCSOP026A02", "QASOPXXXA01", "QASOPXXXA02"]


# ---------- phase 5: the signed-in person is recorded (X-Actor, set by the PP Suite proxy) ----------
def test_the_signed_in_person_is_recorded_on_every_change(api, monkeypatch):
    c, store, started = api
    from app import db
    seen = {}
    real_create, real_save = db.brief_create, db.brief_save

    async def create(b, by=""):
        seen["created_by"] = by
        return await real_create(b, by)

    async def save(new, old, events):
        seen.setdefault("actors", []).extend(e.get("actor") for e in events)
        return await real_save(new, old, events)
    monkeypatch.setattr(db, "brief_create", create)
    monkeypatch.setattr(db, "brief_save", save)
    hx = {**h(), "X-Actor": "b.nikolov"}
    b = c.post("/briefs", json={"package": "sop", "meta": {"title_en": "Sampling"}, "requested_by": "ppdocwiz-session"},
               headers=hx).json()
    assert seen["created_by"] == "b.nikolov"
    c.patch(f"/briefs/{b['id']}", json={"rev": 0, "actor": "someone-else", "set": {"focus": "Potency"}}, headers=hx)
    r = c.post(f"/briefs/{b['id']}/build", json={"rev": 1}, headers=hx).json()
    assert set(seen["actors"]) == {"b.nikolov"} and started[-1]["requested_by"] == "b.nikolov" and r["status"] == "queued"
    c.patch(f"/briefs/{b['id']}", json={"rev": 2, "actor": "direct-caller", "set": {"focus": "Purity"}}, headers=h())
    assert seen["actors"][-1] == "direct-caller"          # without the header, the body's own actor


# ---------- 19.10.2026: an annex brief kept the SOP question set and was written as a nine-section SOP ----------
def test_changing_the_package_takes_its_question_set_and_a_mismatch_is_never_built():
    b = B.new_brief("sop_annexes", {"title_en": "Request for sampling"})
    assert b["questionnaire"] == "sop_qc"
    b, ev = B.apply_patch(b, {"rev": 0, "package": "annex", "parent_sop": "QCSOP001"})
    assert b["package"] == "annex" and b["questionnaire"] == "annex_form"
    assert {e["op"] for e in ev} >= {"package", "questionnaire"}
    b, _ = B.apply_patch(b, {"rev": 1, "package": "form"})            # annex -> form: the form set already fits
    assert b["questionnaire"] == "annex_form"
    b, _ = B.apply_patch(b, {"rev": 2, "package": "sop"})
    assert b["questionnaire"] == "sop_qc"
    with pytest.raises(B.BriefError):
        B.apply_patch(b, {"rev": 3, "questionnaire": "annex_form"})    # a form set on an SOP brief
    with pytest.raises(B.BriefError):
        B.new_brief("annex", questionnaire="sop_qc")
    stuck = {**B.new_brief("sop_annexes", {"title_en": "x"}), "package": "annex"}   # a brief saved before the fix
    with pytest.raises(B.BriefError, match="choose the package again"):
        B.to_workflow(stuck, [])


# ---------- review 19.10.2026: the brief path is guarded as POST /workflows is ----------
def test_a_brief_whose_code_and_package_disagree_is_refused_before_a_job_exists(api, monkeypatch):
    """repro 1 and 2: an SOP brief with the code QCSOP001A01 was built as an SOP (no _check_doctype on the brief path);
    an annex brief for QCSOP001 numbered QCSOP026 was built as an annex of nothing."""
    c, store, started = api
    r = c.post("/briefs", json={"package": "sop", "meta": {"title_en": "Request for Sampling", "code": "QCSOP001A01"}}, headers=h())
    assert r.status_code == 422 and "annex" in r.json()["detail"]
    r = c.post("/briefs", json={"package": "annex", "parent_sop": "QCSOP001", "meta": {"title_en": "Form", "code": "QCSOP026"}}, headers=h())
    assert r.status_code == 422 and "no annex number" in r.json()["detail"]
    ok = c.post("/briefs", json={"package": "annex", "parent_sop": "QCSOP001", "meta": {"title_en": "Form", "code": "QCSOP001A01"}}, headers=h())
    assert ok.status_code == 200
    assert c.post("/briefs", json={"package": "form", "meta": {"title_en": "Form", "code": "QCT026"}}, headers=h()).status_code == 200
    # a patch that makes them disagree
    b = c.post("/briefs", json={"package": "sop", "meta": {"title_en": "S", "code": "QCSOP030"}}, headers=h()).json()
    assert c.patch(f"/briefs/{b['id']}", json={"rev": 0, "meta": {"code": "QCSOP001A01"}}, headers=h()).status_code == 422
    assert c.patch(f"/briefs/{b['id']}", json={"rev": 0, "package": "annex", "parent_sop": "QCSOP001"}, headers=h()).status_code == 422
    assert store[b["id"]]["rev"] == 0                                               # nothing was saved
    # the build path itself, for a brief saved before the guard existed
    store[b["id"]]["meta"]["code"] = "QCSOP001A01"
    r = c.post(f"/briefs/{b['id']}/build", json={"rev": 0}, headers=h())
    assert r.status_code == 422 and "annex" in r.json()["detail"] and started == []
    store[b["id"]]["meta"]["code"] = "QCSOP030"
    store[b["id"]]["package"], store[b["id"]]["questionnaire"] = "annex", "annex_form"
    store[b["id"]]["parent_sop"] = "QCSOP001"
    r = c.post(f"/briefs/{b['id']}/build", json={"rev": 0}, headers=h())
    assert r.status_code == 422 and "no annex number" in r.json()["detail"] and started == []
    store[b["id"]]["meta"]["code"] = ""                                             # numbered under the SOP at build
    r = c.post(f"/briefs/{b['id']}/build", json={"rev": 0}, headers=h())
    assert r.status_code == 200 and r.json()["code"] == "QCSOP001A01" and started[-1]["questionnaire"] == "annex_form"


def test_a_brief_seeded_from_a_reference_is_guarded_too(api, monkeypatch):
    c, store, started = api
    ref = {"id": "55555555-5555-5555-5555-555555555555", "kind": "approved", "code": "QCSOP001A01", "parent": "QCSOP001",
           "version": "01", "title_en": "Request for sampling", "title_mk": "", "markdown": "x" * 30, "sha256": "s"}

    async def get(rid):
        return ref if rid == ref["id"] else None
    monkeypatch.setattr(db, "reference_get", get)
    r = c.post("/briefs", json={"package": "sop", "from_reference": ref["id"]}, headers=h())
    assert r.status_code == 422 and "QCSOP001A01" in r.json()["detail"]
    b = c.post("/briefs", json={"package": "annex", "from_reference": ref["id"]}, headers=h()).json()
    assert b["meta"]["code"] == "QCSOP001A01" and b["parent_sop"] == "QCSOP001" and b["sources"][0]["name"].startswith("QCSOP001A01")
    # a revision of an approved document is the next version (Head of QC, 19.10.2026: QCT002 v02 → v03)
    assert b["meta"]["version"] == "02"
    ref["kind"] = "wip"
    assert c.post("/briefs", json={"package": "annex", "from_reference": ref["id"]}, headers=h()).json()["meta"]["version"] == "01"


def test_locking_a_suggestion_is_a_decision():
    """repro 4: a locked field in state 'suggested' was hidden from `assumed` without anyone deciding it."""
    b, ev = B.apply_patch(fresh(), {"rev": 0, "suggest": {"focus": "Potency"}, "lock": ["focus"]}, "b.nikolov")
    assert [e["op"] for e in ev] == ["suggest", "set", "lock"]
    assert b["fields"]["focus"] == {"value": "Potency", "state": "locked", "stale": False, "at": ev[1]["at"], "by": "b.nikolov"}
    assert ev[1]["before"] == ev[1]["after"] == "Potency" and ev[1]["actor"] == "b.nikolov"
    assert not any(a.startswith("focus") for a in B.to_workflow(b, [])["assumed"])
    b2, _ = B.apply_patch(fresh(), {"rev": 0, "suggest": {"focus": "Potency"}})
    assert "focus (suggested, not confirmed)" in B.to_workflow(b2, [])["assumed"]    # still open: still assumed


def test_wrong_types_in_meta_and_annexes_are_a_400_never_a_500(api):
    """repro 3: a non-string department was stored and crashed the numbering at build."""
    c, store, started = api
    assert c.post("/briefs", json={"package": "sop", "meta": {"title_en": "S", "department": 5}}, headers=h()).status_code == 400
    b = c.post("/briefs", json={"package": "sop_annexes", "meta": {"title_en": "S"}}, headers=h()).json()
    for bad in ({"meta": {"department": ["QC"]}}, {"meta": {"orient": "a\nb"}}, {"meta": {"version": 1}},
                {"annexes": [{"title_en": ["Log"]}]}, {"annexes": [{"title_en": "Log", "kind": 3}]}):
        r = c.patch(f"/briefs/{b['id']}", json={"rev": 0, **bad}, headers=h())
        assert r.status_code == 400, bad
    store[b["id"]]["meta"]["department"] = 5                                        # saved before the check existed
    r = c.post(f"/briefs/{b['id']}/build", json={}, headers=h())
    assert r.status_code == 400 and "department" in r.json()["detail"] and started == []
    with pytest.raises(B.BriefError):
        B.apply_patch(fresh(), {"rev": 0, "meta": {"department": ["QC"]}})
    with pytest.raises(B.BriefError):
        B.apply_patch(fresh(), {"rev": 0, "annexes": [{"title_mk": {"x": 1}}]})


def test_a_house_form_may_belong_to_an_sop_under_its_own_number(api, monkeypatch):
    """Head of QC, 19.10.2026: QCT002 (Request for analysis) is revised to v03 under QCSOP001 — a form or logbook of an
    SOP keeps its own family number; only an SOP-family code without an annex number is refused with a parent."""
    c, store, started = api
    r = c.post("/briefs", json={"package": "form", "parent_sop": "QCSOP001",
                                "meta": {"title_en": "Request for sampling and/or analysis", "code": "QCT002", "version": "03"}}, headers=h())
    assert r.status_code == 200
    r = c.post(f"/briefs/{r.json()['id']}/build", json={"rev": 0}, headers=h())
    assert r.status_code == 200 and r.json()["code"] == "QCT002" and started[-1]["parent_sop"] == "QCSOP001"
