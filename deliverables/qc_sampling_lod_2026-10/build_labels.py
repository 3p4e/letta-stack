#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Print-ready label sheets for the sampling execution package of PP-QC-SP-002/26 (QASOP_031).

  QASOP_031_A05_v1  "МОСТРИРАНО | SAMPLED" bag label — A4 sheet of 52.5 x 33 mm perforated labels
                    (4 x 9), one label per sampled bag, prefilled with batch, P lot and bag ID K#B#.
  QASOP_031_A07     sample label — A4 sheet 4 x 2 (105 x 74.25 mm), one label per composite sample,
                    prefilled with the material, strain, batch, P lot and the owner's harvest and
                    packaging dates; the sample code, masses, date and sampler are written at sampling.

The fields and wording are the templates' own (QASOP_031_Axx Sampled Label, QASOP_031_A07 Sample
Label FULL_A4-4x2); only the values are added. These sheets carry no running header: they are
labels, printed edge to edge on perforated stock, not documents.

    python3 build_labels.py
"""
import csv
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_campaign_docs as bc   # noqa: E402  (data loader, glyph audit, engine assets on sys.path)
import campaign_data as cd         # noqa: E402  (owner's workbook path, batch key)
from docx import Document          # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn        # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

OUT_S = os.path.join(bc.OUT, "2_SAMPLING_EXECUTION")
FONT = "Calibri"                   # renders as Carlito, the engine's metric-compatible house face
BLACK = RGBColor(0, 0, 0)
GREY = RGBColor(0x59, 0x59, 0x59)
NAVY = RGBColor(0x2B, 0x54, 0x7E)
SID = "____/26_SFR-PC-___"


def bl(l):
    return l["batch"] if l["p_lot"] in ("", "—", None) else "%s · %s" % (l["batch"], l["p_lot"])


def harvest_dates():
    out = {}
    with open(cd.BATCH_DATES_CSV, encoding="utf-8") as fh:
        for row in csv.DictReader(fh):
            a, b = (row.get("harvest_from") or "").strip(), (row.get("harvest_to") or "").strip()
            v = a if (a == b or not b) else "%s – %s" % (a, b)
            if v:
                out[cd.batch_key(row["batch"])] = v
                if (row.get("p_batch") or "").strip():
                    out[row["p_batch"].strip()] = v
    return out


def sheet_doc():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(s, side, Cm(0))
    s.header_distance = s.footer_distance = Cm(0)
    st = d.styles["Normal"]
    st.font.name = FONT
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.font.size = Pt(7)
    pf = st.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    return d


def run(p, text, size, bold=False, color=BLACK, ital=False):
    r = p.add_run(text)
    r.font.name = FONT
    r._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = ital
    r.font.color.rgb = color
    return r


def para(cell, first=False, align=WD_ALIGN_PARAGRAPH.LEFT, before=0.0):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(0)
    return p


def label_grid(d, cols, rows, w_cm, h_cm, labels, fill):
    """One table per sheet; a tiny trailing paragraph keeps each sheet on its own page."""
    per = cols * rows
    for start in range(0, len(labels), per):
        chunk = labels[start:start + per]
        t = d.add_table(rows=rows, cols=cols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        tblPr = t._tbl.tblPr
        lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
        mar = OxmlElement("w:tblCellMar")
        for side, v in (("top", 85), ("left", 113), ("bottom", 57), ("right", 85)):
            e = OxmlElement("w:" + side); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); mar.append(e)
        tblPr.append(mar)
        for gc in t._tbl.tblGrid.gridCol_lst:
            gc.set(qn("w:w"), str(int(w_cm * 567)))
        for r in t.rows:
            r.height = Cm(h_cm)
            r.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
            trPr = r._tr.get_or_add_trPr()
            cs = OxmlElement("w:cantSplit"); trPr.append(cs)
            for c in r.cells:
                c.width = Cm(w_cm)
        for i, lab in enumerate(chunk):
            fill(t.cell(i // cols, i % cols), lab)
        if start + per < len(labels):
            p = d.add_paragraph()
            p.paragraph_format.line_spacing = Pt(1)
            run(p, "", 1)
            p.add_run().add_break(WD_BREAK.PAGE)
        else:
            p = d.add_paragraph()
            p.paragraph_format.line_spacing = Pt(1)
    return d


def fill_sampled(cell, lab):
    p = para(cell, first=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    run(p, "МОСТРИРАНО", 11, bold=True, color=NAVY)
    run(p, "  |  SAMPLED", 8, bold=True, color=NAVY)
    p = para(cell, before=2)
    run(p, "Датум на мострир. | Date of sampling: ", 6.5)
    run(p, "____________", 7)
    p = para(cell, before=2)
    run(p, "Прием / Серија / Лот бр. | Receipt Bill No. / Batch no / Lot No.", 6, color=GREY)
    p = para(cell)
    run(p, lab["batch_lot"], 8, bold=True)
    p = para(cell, before=2)
    # a bag chosen at sampling (06.10.2026) is a write-in, K___B___, set smaller so the line still fits
    if "_" in lab["bag"]:
        run(p, lab["bag"], 10, bold=True)
        run(p, "   Потпис | Signature: ", 6.5)
        run(p, "_________", 7)
    else:
        run(p, lab["bag"], 13, bold=True)
        run(p, "     Потпис | Signature: ", 6.5)
        run(p, "__________", 7)
    p = para(cell, before=1, align=WD_ALIGN_PARAGRAPH.RIGHT)
    run(p, "QASOP_031_A05_v1 · %s" % lab["day"], 6, color=GREY)


def fill_sample(cell, lab):
    def field(mk_en, value=None, size=7.5, bold=False):
        p = para(cell, before=1.5)
        run(p, mk_en + (": " if value is not None else ""), 6.5, color=GREY)
        if value is not None:
            run(p, value, size, bold=bold)

    p = para(cell, first=True)
    run(p, "Намена на мостра, анализа | Sample intended use, assay", 6.5, color=GREY)
    p = para(cell)
    run(p, "Мостра за испитување на Губиток при сушење", 9, bold=True, color=NAVY)
    run(p, " | Loss on Drying Test Sample", 7.5, bold=True, color=NAVY)
    field("Опис на Материјал | Material Description")
    p = para(cell)
    run(p, "Сув Цвет од Канабис за медицинска употреба | Dry Cannabis Flower for Medical Use", 7.5)
    field("Сорта и/или ID код на растение | Strain and/or Plant ID", lab["strain"], bold=True)
    field("ID код на мостра (sID) | Sample ID code (sID)", SID, size=8.5, bold=True)
    field("Серија / Лот бр. | Batch / Lot No.", lab["batch_lot"], size=8.5, bold=True)
    field("Датум на берба | Date of harvest", lab["harvest"])
    field("Датум на пакување | Date of packaging", lab["packaging"])
    p = para(cell, before=1.5)
    run(p, "Нето (g) | Net (g): ", 6.5, color=GREY); run(p, "________", 7.5)
    run(p, "     Бруто (g) | Gross (g): ", 6.5, color=GREY); run(p, "________", 7.5)
    field("Услови за чување | Storage condition",
          "Да се чува на суво и темно место на температура од 15-25°C | Store in a dry, dark place at 15-25°C", size=6.5)
    p = para(cell, before=1.5)
    run(p, "Датум на мострирање | Date of sampling: ", 6.5, color=GREY); run(p, "__________", 7.5)
    run(p, "   Мострирано од | Sampled by: ", 6.5, color=GREY); run(p, "__________", 7.5)
    p = para(cell, before=1, align=WD_ALIGN_PARAGRAPH.RIGHT)
    run(p, "QASOP_031_A07 · %s" % lab["day"], 6, color=GREY)


def build(day, lots, bags, hv):
    ls = [l for l in lots if l["day"] == day]
    tag = "%s · Ден | Day %d" % (bc.CODE, day)
    sampled = []
    for l in ls:
        for b in [x for x in bags if int(x["day"]) == day and x["batch"] == l["batch"]]:
            sampled.append({"batch_lot": bl(l), "bag": b["bag_id"], "day": tag})
    samples = []
    for l in ls:
        samples.append({"batch_lot": bl(l), "strain": l["strain"],
                        "harvest": hv.get(cd.batch_key(l["batch"])) or hv.get(l["p_lot"]) or "____________",
                        "packaging": l["packaging_date"] or "____________", "day": tag})
    d1 = label_grid(sheet_doc(), 4, 9, 5.25, 3.29, sampled, fill_sampled)
    d2 = label_grid(sheet_doc(), 2, 4, 10.5, 7.41, samples, fill_sample)
    return (d1, len(sampled)), (d2, len(samples))


STEMS = ("S%d-4_QASOP031_A05_SAMPLED_Bag_Labels_Day%d", "S%d-5_QASOP031_A07_Sample_Labels_Day%d")


def main():
    os.makedirs(OUT_S, exist_ok=True)
    lots, bags = bc.load()
    hv = harvest_dates()
    for day in (1,):                 # one sampling day (Head of QC, 06.10.2026)
        for (d, n), stem in zip(build(day, lots, bags, hv), STEMS):
            bc.glyph_audit(d)
            path = os.path.join(OUT_S, (stem % (day, day)) + ".docx")
            d.save(path)
            print("%4d labels  %s" % (n, os.path.basename(path)))


if __name__ == "__main__":
    main()
