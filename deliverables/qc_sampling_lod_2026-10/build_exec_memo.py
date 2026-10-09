#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""PP-QC-SP-002/26-INF-01 — official information note from the Head of QC to the CEO and the executive
management, 07.10.2026: the deviation in the initial weighing of the loss-on-drying samples of run 1
(LOD-01, DEV-01), the measures taken before the 24-hour weighing, what the run-1 values show, and the
repeat of all 46 analyses on the AUW220D (LOD-01R) as the only official, documented LoD evidence before
shipment. Built with the Purely Plant document engine (pp-document-suite) through the campaign helpers.

Head of QC, 07.10.2026: "Create the document using the PP document engine. Don't create the Word files."
The DOCX the engine writes is therefore built in a working folder outside the repository, checked
(glyph audit, pp_verify), converted to PDF with LibreOffice as the campaign does, and only the PDF is
copied to out/3_LOD_ANALYSIS_EXECUTION/.

Every run-1 number is computed here from run1_LOD-01_2026-10-05/RUN1_CHECK.tsv (check_run1.py, the
laboratory's workbook recalculated from its raw weighings) and asserted against that file; the tranche
counts come from SAMPLING_PLAN_T1_T2_2026-10.tsv. The workbook's batch column is empty and its sample IDs
are S1-S46, so no run-1 value is tied to a batch name.

    python3 build_exec_memo.py [--work DIR]      # DIR: where the DOCX is built (default: a temp folder)
"""
import argparse
import csv
import math
import os
import shutil
import statistics as st
import subprocess
import sys
import tempfile
from decimal import Decimal as D

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_campaign_docs as bc                         # noqa: E402  (engine on sys.path, constants)
import build_execution_packages as bx                    # noqa: E402  (grid, row_height, codes)
from build_campaign_docs import pf, pr                   # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH            # noqa: E402
from docx.oxml import OxmlElement                        # noqa: E402

CODE = bc.CODE + "-INF-01"                 # INF = information note; the campaign's -LOD-/-DEV- scheme
DATE = "07.10.2026"
STATUS = bc.STATUS                         # the campaign rule: IN REVIEW until the Head of QC approves
LOD, LODR, DEV = bx.lod_code(1), bx.lod_code(1) + "R", bx.dev_code()
OUT_DIR = os.path.join(bc.OUT, "3_LOD_ANALYSIS_EXECUTION")
STEM = "PP-QC-SP-002_26-INF-01_Information_Note_LoD_Deviation_Repeat"
RUN1 = os.path.join(HERE, "run1_LOD-01_2026-10-05", "RUN1_CHECK.tsv")
HEAD_QC = "B. Nikolov, M.Pharm."           # as on every campaign document (PREPARED / APPROVED)
W = 18.46


# ----------------------------------------------------------------------------- data
def mk(x, dec=2, sign=False):
    """Macedonian number: decimal comma."""
    return ((("%+." if sign else "%.") + str(dec) + "f") % x).replace(".", ",").replace("-", "−")


def en(x, dec=2, sign=False):
    return ((("%+." if sign else "%.") + str(dec) + "f") % x).replace("-", "−")


def run1():
    """RUN1_CHECK.tsv -> one dict per sample, recomputed from the masses and asserted against the file."""
    with open(RUN1, encoding="utf-8") as fh:
        rows = list(csv.DictReader((l for l in fh if not l.startswith("#")), delimiter="\t"))
    out = []
    for r in rows:
        mb, m0, g1, gp, ga = (D(r[k]) for k in ("m_B_g", "m0_g", "G1_g", "G2_precision_g", "G2_AUW220D_g"))
        assert g1 == mb + m0, r["sample_id"]
        lp, la = (g1 - gp) / m0 * 100, (g1 - ga) / m0 * 100
        assert "%.2f" % lp == r["LoD_pct_precision_pair"] and "%.2f" % la == r["LoD_pct_G2_on_AUW220D"], r["sample_id"]
        s = dict(sid=r["sample_id"].rsplit("_", 1)[1], m0=m0, lp=float(lp), la=float(la),
                 lp2=D(r["LoD_pct_precision_pair"]), la2=D(r["LoD_pct_G2_on_AUW220D"]), dg=float((ga - gp) * 1000))
        s["crit"] = [n for n, hit in (
            (1, s["lp2"] >= D("10.00")),                                  # highest results, nearest the limit
            (2, not D("0.900") <= m0 <= D("1.100")),                      # Ph. Eur.: +/-10 % on 1.000 g
            (3, m0.as_tuple().exponent < -3),                             # 4 decimals: not a precision reading
            (4, abs(s["la2"] - s["lp2"]) >= D("0.10")),                   # moves >= 0.10 % between balances
        ) if hit]
        out.append(s)
    assert len(out) == 46 and [s["sid"] for s in out] == ["S%d" % i for i in range(1, 47)]
    return out


def stats(ss):
    lp = [s["lp"] for s in ss]
    lo = min(ss, key=lambda s: s["lp"])
    hi = max(ss, key=lambda s: s["lp"])
    bins = [("< 7,00", "< 7.00", lambda v: v < 7), ("7,00–7,99", "7.00–7.99", lambda v: 7 <= v < 8),
            ("8,00–8,99", "8.00–8.99", lambda v: 8 <= v < 9), ("9,00–9,99", "9.00–9.99", lambda v: 9 <= v < 10),
            ("10,00–10,99", "10.00–10.99", lambda v: 10 <= v < 11), ("≥ 11,00", "≥ 11.00", lambda v: v >= 11)]
    dist = [(a, b, sum(1 for s in ss if f(float(s["lp2"])))) for a, b, f in bins]
    assert sum(n for _, _, n in dist) == len(ss)
    dg = [s["dg"] for s in ss]
    shift = [float(s["la2"] - s["lp2"]) for s in ss]
    m0 = [s["m0"] for s in ss]
    # DEV-01 section C: U (k = 2) = 2 * sqrt(2) * sqrt(s^2 + d^2 / 12) / m0 * 100, s = d = 1 mg, m0 = 1.000 g
    u = 2 * math.sqrt(2) * math.sqrt(1.0 + 1.0 / 12) / 1000 * 100
    return dict(n=len(ss), lo=lo, hi=hi, mean=st.mean(lp), sd=st.stdev(lp), median=st.median(lp),
                mean_a=st.mean(s["la"] for s in ss), dist=dist, n_ok=sum(1 for v in lp if v <= 12.0),
                dg_mean=st.mean(dg), dg_sd=st.stdev(dg), dg_min=min(dg), dg_max=max(dg),
                shift_max=max(abs(x) for x in shift), m0_min=min(m0), m0_max=max(m0), u=u)


def tranches():
    lots, _ = bc.load()
    t = [sum(1 for l in lots if l["tranche"] == k) for k in (1, 2)]
    assert sum(t) == len(lots) == 46, t
    return t


# ----------------------------------------------------------------------------- layout helpers
def section(d, num_, mk_, en_):
    """Memo-sized numbered heading: the engine's subsection face (MK 16 | EN 12), kept with what follows."""
    p = pr.subsec(d, num_ + ".", mk_, en_)
    p.paragraph_format.keep_with_next = True
    return p


def caption(d, mk_, en_):
    p = d.add_paragraph()
    pr.sp(p, 6, 3)
    p.paragraph_format.keep_with_next = True
    pr.rin(p, mk_, 10, pr.NAVY, bold=True)
    pr.rin(p, "  |  ", 9, pr.GREY)
    pr.rin(p, en_, 9, pr.GREY, ital=True)
    return p


def title_block(d):
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; pr.sp(p, 4, 2)
    pr.rin(p, "СЛУЖБЕНА ИНФОРМАЦИЈА", 13, pr.GREY, bold=True); pr.rin(p, "  |  ", 11, pr.GREY)
    pr.rin(p, "OFFICIAL INFORMATION NOTE", 11, pr.GREY)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; pr.sp(p, 4, 2)
    pr.rin(p, "Отстапување при почетното мерење на примероците за губиток при сушење и повторување на "
              "сите 46 анализи", 17, pr.NAVY, bold=True)
    p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; pr.sp(p, 0, 8)
    pr.rin(p, "Deviation in the initial weighing of the loss-on-drying samples and repeat of all 46 analyses",
           12, pr.NAVY)
    pf.status_band(d, STATUS, bc.VERSION)
    pr.gap(d, 8)


def sign_block(d):
    rows = [[(x, 1, "h") for x in ["Улога | Role", "Име | Name", "Датум | Date", "Потпис | Signature"]],
            [("Од: Раководител на КК | From: Head of QC", 1, "l"), HEAD_QC, "", ""],
            [("Примил: Главен извршен директор | Received: CEO", 1, "l"), "", "", ""]]
    rows += [[("Примил: Извршен менаџмент | Received: Executive management", 1, "l"), "", "", ""] for _ in range(2)]
    t = bx.grid(d, [6.86, 5.2, 2.8, 3.6], rows, sz=9, head=1, heights={i: 0.85 for i in range(1, len(rows))})
    for r in t.rows:                     # the sign-off block never splits across a page
        r._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for c in r.cells:
            for p in c.paragraphs:
                p.paragraph_format.keep_with_next = True


# ----------------------------------------------------------------------------- the note
def build(ss, t1, t2):
    S = stats(ss)
    by = {s["sid"]: s for s in ss}
    crit = {n: [s for s in ss if n in s["crit"]] for n in (1, 2, 3, 4)}
    flagged = [s for s in ss if s["crit"]]
    double = [s for s in ss if len(s["crit"]) >= 2]
    hi, lo = S["hi"], S["lo"]
    assert hi["sid"] == "S39" and lo["sid"] == "S5"                       # as stated in README / RUN1_CHECK.tsv

    d = bc.new_doc(CODE, "СЛУЖБЕНА ИНФОРМАЦИЈА — ГУБИТОК ПРИ СУШЕЊЕ, ТРАНША 1 И 2",
                   "OFFICIAL INFORMATION NOTE — LOSS ON DRYING, T1 AND T2")
    cp = d.core_properties
    cp.title = "%s — Official information note: LoD deviation and repeat, Tranches 1 and 2" % CODE
    cp.subject = "Loss on drying before shipment, PP-QC-SP-002/26"
    cp.keywords = "%s; %s; %s; %s" % (CODE, LOD, DEV, LODR)
    title_block(d)
    bc.kv_table(d, [
        ("Документ бр. | Document No.", CODE),
        ("До | To", "Главен извршен директор и извршен менаџмент | CEO and executive management"),
        ("Од | From", "%s — Раководител на КК | Head of QC" % HEAD_QC),
        ("Датум | Date", DATE),
        ("Предмет | Subject", "Губиток при сушење пред испорака, Транша 1 и 2 (%d серии): отстапување %s и повторување %s | "
                              "Loss on drying before shipment, Tranches 1 and 2 (%d batches): deviation %s and repeat %s"
         % (S["n"], DEV, LODR, S["n"], DEV, LODR)),
        ("Намена | Purpose", "За информација; потврда за прием во делот 7 | For information; acknowledgement of receipt in section 7"),
    ])

    # 1
    section(d, "1", "ЦЕЛ", "Purpose")
    pr.body(d, "Со оваа информација раководството се известува за отстапувањето при почетното мерење на примероците во "
               "првото сушење за губиток при сушење (ГпС) на Транша 1 и 2, за мерките преземени пред мерењето по 24 h, за "
               "тоа што покажуваат пресметаните вредности и зошто сите %d анализи треба да се повторат за да се добијат "
               "службени, документирани резултати пред испорака." % S["n"],
            "This note informs management of the deviation in the initial weighing of the samples in the first loss-on-drying "
            "(LoD) run of Tranches 1 and 2, the measures taken before the 24-h weighing, what the calculated values show, and "
            "why all %d analyses must be repeated to obtain official, documented results before shipment." % S["n"])

    # 2
    section(d, "2", "ШТО СЕ СЛУЧИ", "What Happened")
    pr.body(d, "Пред испораката на Транша 1 (%d серии) и Транша 2 (%d серии), ГпС на сите %d серии се определува во КК "
               "лабораторијата по методот a02.2: по една тест порција од околу 1,000 g од секоја серија, сите во едно сушење "
               "во вакуумската печка VO29 на 40 °C и 20 ± 2 mbar над молекуларно сито R во тек на 24 h, ладење најмалку 30 min "
               "во ексикатор; ГпС %% = (G1 − G2) ÷ m₀ × 100. Методот бара вага што чита најмалку четири децимали (0,1 mg) — "
               "аналитичката вага Shimadzu AUW220D." % (t1, t2, S["n"]),
            "Before Tranche 1 (%d batches) and Tranche 2 (%d batches) ship, the LoD of all %d batches is determined in the QC "
            "laboratory by method a02.2: one test portion of about 1.000 g per batch, all in one run in vacuum oven VO29 at "
            "40 °C and 20 ± 2 mbar over molecular sieve R for 24 h, cooled at least 30 min in a desiccator; LoD %% = (G1 − G2) "
            "÷ m₀ × 100. The method requires a balance reading at least four decimals (0.1 mg): the Shimadzu AUW220D "
            "analytical balance." % (t1, t2, S["n"]))
    pr.gap(d, 4)
    for mk_, en_ in [
        ("Отстапување: во првото сушење (%s, започнато на 05.10.2026) празниот сад (m_B) и тест порцијата од околу 1 g (G1) "
         "на сите %d порции се измерени на прецизната вага, која чита три децимали (d = 1 mg), наместо на AUW220D."
         % (LOD, S["n"]),
         "Deviation: in the first run (%s, started 05.10.2026) the empty bottle (m_B) and the ~1 g test portion (G1) of all "
         "%d portions were weighed on the precision balance, which reads three decimals (d = 1 mg), instead of the AUW220D."
         % (LOD, S["n"])),
        ("Откриено на 06.10.2026, за време на сушењето — пред мерењето по 24 h и пред каква било пресметка на резултат — и "
         "заведено како отстапување %s." % DEV,
         "Found on 06.10.2026, during drying — before the 24-h weighing and before any result was calculated — and recorded "
         "as deviation %s." % DEV),
    ]:
        pr.bullet(d, mk_, en_)

    # 3
    section(d, "3", "ВЕДНАШНИ МЕРКИ ПРЕД МЕРЕЊЕТО ПО 24 h", "Immediate Measures")
    for mk_, en_ in [
        ("Мерењето по 24 h (G2) е направено на истата прецизна вага, така што сите три мерења во резултатот (m_B, G1, G2) "
         "се од една вага и разликата меѓу двете ваги не влегува во резултатот.",
         "The 24-h weighing (G2) was made on the same precision balance, so that all three weighings in the result (m_B, G1, "
         "G2) come from one balance and the offset between the two balances does not enter the result."),
        ("Веднаш потоа секој сад е измерен и на аналитичката вага AUW220D; двете мерења за сите %d садови се запишани "
         "(Прилог 1 кон %s; табелата на лабораторијата)." % (S["n"], LOD),
         "Straight after, every bottle was also weighed on the AUW220D analytical balance; both readings for all %d bottles "
         "are recorded (Attachment 1 to %s; the laboratory workbook)." % (S["n"], LOD)),
        ("Остатоците од примероците се чуваат затворени и означени, за анализата да може да се повтори од истиот материјал.",
         "The sample remainders are kept closed and labelled, so that the analysis can be repeated on the same material."),
    ]:
        pr.bullet(d, mk_, en_)

    # 4
    section(d, "4", "ШТО ПОКАЖУВААТ РЕЗУЛТАТИТЕ ОД СУШЕЊЕТО 1", "Run-1 Results")
    pr.body(d, "Лабораторијата ги пресмета вредностите во својата табела (T1_and_T2_LOD_Analysis.xlsx); секој ред е повторно "
               "пресметан од суровите мерења (check_run1.py → RUN1_CHECK.tsv) и се согласува за сите %d реда. Табела 1 ги "
               "дава клучните бројки, табела 2 вредностите што бараат внимание, Прилогот A сите %d вредности." % (S["n"], S["n"]),
            "The laboratory calculated the values in its workbook (T1_and_T2_LOD_Analysis.xlsx); every row was recalculated "
            "from the raw weighings (check_run1.py → RUN1_CHECK.tsv) and agrees on all %d rows. Table 1 gives the key figures, "
            "Table 2 the values that call for caution, Annex A all %d values." % (S["n"], S["n"]))
    caption(d, "Табела 1 — Клучни бројки, сушење 1 (05–06.10.2026)", "Table 1 — Key figures, run 1 (05–06.10.2026)")
    dist = " · ".join("%s: %d" % (a, n) for a, _, n in S["dist"])
    bx.grid(d, [6.4, 12.06], [
        [("Показател | Indicator", 1, "h"), ("Вредност | Value", 1, "h")],
        [("Порции | Portions", 1, "l"), "%d · S1–S%d · по една за серија | one per batch" % (S["n"], S["n"])],
        [("ГпС %, прецизна вага | LoD %, precision balance", 1, "l"),
         "%s (%s) – %s (%s) · средна %s · SD %s · медијана %s | mean %s · SD %s · median %s"
         % (mk(lo["lp"]), lo["sid"], mk(hi["lp"]), hi["sid"], mk(S["mean"]), mk(S["sd"]), mk(S["median"]),
            en(S["mean"]), en(S["sd"]), en(S["median"]))],
        [("Во критериумот ≤ 12,0 % | Within ≤ 12.0 %", 1, "l"), ("%d од %d | %d of %d" % (S["n_ok"], S["n"], S["n_ok"], S["n"]), 1, "b")],
        [("Распределба, ГпС % (бр.) | Distribution, LoD % (no.)", 1, "l"), dist],
        [("G2 (AUW220D) − G2 (прецизна) | G2 (AUW220D) − G2 (precision)", 1, "l"),
         "средна %s mg · SD %s mg · опсег %s до %s mg (n = %d) | mean %s mg · SD %s mg · range %s to %s mg"
         % (mk(S["dg_mean"], sign=True), mk(S["dg_sd"]), mk(S["dg_min"], sign=True), mk(S["dg_max"], sign=True), S["n"],
            en(S["dg_mean"], sign=True), en(S["dg_sd"]), en(S["dg_min"], sign=True), en(S["dg_max"], sign=True))],
        [("ГпС со G2 од AUW220D | LoD with G2 from the AUW220D", 1, "l"),
         "промена најмногу %s %% апсолутно · средна %s %% | change at most %s %% absolute · mean %s %%"
         % (mk(S["shift_max"]), mk(S["mean_a"]), en(S["shift_max"]), en(S["mean_a"]))],
        [("Тест порција m₀ | Test portion m₀", 1, "l"),
         "%s–%s g · %d надвор од 0,900–1,100 g | %d outside 0.900–1.100 g"
         % (mk(S["m0_min"], 3), mk(S["m0_max"], 3), len(crit[2]), len(crit[2]))],
        [("U (k = 2) на еден резултат, DEV-01 дел C | U (k = 2) of one result, DEV-01 section C", 1, "l"),
         "≈ %s %% апсолутно (s = 1 mg) | ≈ %s %% absolute (s = 1 mg)" % (mk(S["u"]), en(S["u"]))],
    ], sz=8, head=1)

    caption(d, "Табела 2 — Вредности што бараат внимание", "Table 2 — Values that call for caution")
    lst = lambda ss_, f: " · ".join("%s %s" % (s["sid"], f(s)) for s in ss_)
    top = sorted(crit[1], key=lambda s: -s["lp"])
    rows = [[(x, 1, "h") for x in ["№", "Критериум | Criterion", "Примероци | Samples", "Бр. | No."]],
            ["1", ("ГпС ≥ 10,00 % — најблиску до границата 12,0 % | LoD ≥ 10.00 % — nearest the 12.0 % limit", 1, "n"),
             lst(top, lambda s: mk(s["lp"])) + " (%)", str(len(crit[1]))],
            ["2", ("m₀ надвор од 0,900–1,100 g (Ph. Eur.: ± 10 % од 1,000 g) | m₀ outside 0.900–1.100 g "
                   "(Ph. Eur.: ± 10 % of 1.000 g)", 1, "n"),
             lst(crit[2], lambda s: mk(s["m0"], 3)) + " (g)", str(len(crit[2]))],
            ["3", ("m₀ со четири децимали — не е отчитување од прецизната вага | m₀ to four decimals — "
                   "not a precision-balance reading", 1, "n"),
             lst(crit[3], lambda s: mk(s["m0"], 4) + " g"), str(len(crit[3]))],
            ["4", ("ГпС се менува ≥ 0,10 % со G2 од AUW220D | LoD moves ≥ 0.10 % with G2 from the AUW220D", 1, "n"),
             lst(crit[4], lambda s: "%s → %s" % (mk(s["lp2"]), mk(s["la2"]))) + " (%)", str(len(crit[4]))],
            ["5", ("Идентитет на серијата | Batch identity", 1, "n"),
             ("Колоната за серија во табелата на лабораторијата е празна; ознаките се S1–S%d, па вредностите не можат да се "
              "поврзат со серии само од датотеката | The workbook's batch column is empty; the IDs are S1–S%d, so the values "
              "cannot be tied to batches from the file alone" % (S["n"], S["n"]), 1, "n"), str(S["n"])]]
    bx.grid(d, [0.7, 7.6, 8.96, 1.2], rows, sz=8, head=1)
    pr.note(d, "Со барем еден критериум 1–4: %d примероци (%s); со два: %s."
               % (len(flagged), ", ".join(s["sid"] for s in flagged),
                  " и ".join("%s (%s)" % (s["sid"], " и ".join(map(str, s["crit"]))) for s in double)),
            "Meeting at least one of criteria 1–4: %d samples (%s); meeting two: %s."
            % (len(flagged), ", ".join(s["sid"] for s in flagged),
               " and ".join("%s (%s)" % (s["sid"], " and ".join(map(str, s["crit"]))) for s in double)))
    pr.body(d, "Оценка на Раководителот на КК: според пресметаните вредности, резултатите од сушењето 1 се добри — сите %d се "
               "под 12,0 %%, а највисокиот (%s, %s %%) останува под границата и со неизвесноста од %s (%s + %s = %s %%) — освен "
               "можеби три или четири резултати меѓу оние во табела 2. Кои се тие не се одлучува со оваа информација: "
               "повторувањето (делот 5) дава резултат за секоја серија."
               % (S["n"], hi["sid"], mk(hi["lp"]), DEV, mk(hi["lp"]), mk(S["u"]), mk(round(hi["lp"], 2) + round(S["u"], 2))),
            "Assessment of the Head of QC: on the calculated values the run-1 results are good — all %d are below 12.0 %%, and "
            "the highest (%s, %s %%) stays below the limit even with the uncertainty from %s (%s + %s = %s %%) — except perhaps "
            "three or four results among those in Table 2. Which ones is not decided in this note: the repeat (section 5) "
            "gives a result for every batch."
            % (S["n"], hi["sid"], en(hi["lp"]), DEV, en(hi["lp"]), en(S["u"]), en(round(hi["lp"], 2) + round(S["u"], 2))))

    # 5
    section(d, "5", "ОДЛУКА И ПРЕПОРАКА", "Decision and Recommendation")
    for mk_, en_ in [
        ("Сушењето 1 не може да служи како службен, документиран доказ за ГпС при пуштање пред испорака: мерењата пред "
         "сушење (m_B, G1) не се на вага што чита четири децимали, како што бара методот, и ниедно подоцнежно мерење не може "
         "тоа да го поправи (%s, дел D)." % DEV,
         "Run 1 cannot serve as official, documented LoD evidence for release before shipment: the pre-drying weighings (m_B, "
         "G1) are not on a balance reading four decimals, as the method requires, and no later weighing can correct that "
         "(%s, section D)." % DEV),
        ("Одлука на Раководителот на КК (06.10.2026): сушењето 1 се поништува; сите %d серии се повторуваат по записот %s, "
         "кој е подготвен." % (S["n"], LODR),
         "Decision of the Head of QC (06.10.2026): run 1 is invalidated; all %d batches are repeated under record %s, which is "
         "prepared." % (S["n"], LODR)),
        ("Препорака: за службени и валидни резултати за секоја серија пред испорака, поткрепени со записи, мерења и правилно "
         "извршување на постапката, сите %d анализи се изведуваат одново, од почеток до крај — нова тест порција од 1,000 g од "
         "остатокот на секој примерок; m_B, G1 и G2 на аналитичката вага AUW220D поставена на четири децимали (0,1 mg); "
         "едно сушење во VO29 по a02.2." % S["n"],
         "Recommendation: for official, valid results for every batch before shipment, backed by records, weighings and proper "
         "execution of the procedure, all %d analyses are performed again, from start to end — a new 1.000 g test portion from "
         "each sample remainder; m_B, G1 and G2 on the AUW220D analytical balance set to four decimals (0.1 mg); one run in "
         "VO29 per a02.2." % S["n"]),
        ("Време: уште 24 h сушење, плус мерењето пред сушење, ладењето од најмалку 30 min и мерењето по сушењето (и секое "
         "понатамошно мерење до константна маса по %s)." % LODR,
         "Time: a further 24 h of drying, plus the weighing before drying, the cooling of at least 30 min and the weighing "
         "after drying (and any further weighing to constant mass under %s)." % LODR),
        ("За секоја серија се известува резултатот од %s. Податоците од сушењето 1 (05–06.10.2026: %d резултати и %d парови "
         "G2 на двете ваги) остануваат во записот како дополнителни, вредни податоци само за споредба и не се известуваат."
         % (LODR, S["n"], S["n"]),
         "For every batch the %s result is reported. The run-1 data (05–06.10.2026: %d results and %d G2 pairs on the two "
         "balances) stay in the record as additional, valuable data for comparison only and are not reported."
         % (LODR, S["n"], S["n"])),
    ]:
        pr.bullet(d, mk_, en_)

    # 6
    section(d, "6", "РЕФЕРЕНЦИ", "References")
    bx.grid(d, [5.2, 13.26], [
        [("Ознака | Code", 1, "h"), ("Документ | Document", 1, "h")],
        [(bc.CODE, 1, "b"), ("План за земање примероци и ГпС, Транша 1 и 2 (изменет 06.10.2026) | Sampling plan and LoD, "
                              "Tranches 1 and 2 (amended 06.10.2026)", 1, "n")],
        [(LOD, 1, "b"), ("Извршен запис, сушење 1 (поништено) | Execution record, run 1 (invalidated)", 1, "n")],
        [(LOD + "/A1", 1, "b"), ("Прилог 1 кон %s — мерења на две ваги | Attachment 1 to %s — weighings on two balances"
                                 % (LOD, LOD), 1, "n")],
        [(DEV, 1, "b"), ("Извештај за отстапување — мерења пред сушење на прецизна вага | Deviation report — pre-drying "
                         "weighings on a precision balance", 1, "n")],
        [(LODR, 1, "b"), ("Извршен запис — повторување на сите %d серии на AUW220D | Execution record — repeat of all %d "
                          "batches on the AUW220D" % (S["n"], S["n"]), 1, "n")],
        [("Сушење 1 | Run 1", 1, "b"), ("run1_LOD-01_2026-10-05/: T1_and_T2_LOD_Analysis.xlsx (табела на лабораторијата | "
                                       "laboratory workbook) · check_run1.py → RUN1_CHECK.tsv (повторна пресметка | "
                                       "recalculation)", 1, "n")],
        [("Метод | Method", 1, "b"), ("a02.2 · Ph. Eur. 2.2.32 · монографија 3028 Cannabis flos · QCSP 001 (ГпС ≤ 12,0 %) | "
                                     "monograph 3028 Cannabis flos · QCSP 001 (LoD ≤ 12.0 %)", 1, "n")],
    ], sz=8, head=1)

    # 7
    section(d, "7", "ПОТПИС И ПОТВРДА ЗА ПРИЕМ", "Signature and Acknowledgement of Receipt")
    sign_block(d)

    # Annex A
    d.add_page_break()
    section(d, "A", "ПРИЛОГ — ВРЕДНОСТИ ОД СУШЕЊЕТО 1", "Annex — Run-1 Values")
    pr.note(d, "Само за споредба; не се известуваат (%s, дел D). Критериумите 1–4 се од табела 2." % DEV,
            "For comparison only; not reported (%s, section D). Criteria 1–4 are those of Table 2." % DEV)
    head = ["Пр. | ID", "m₀ g", "ГпС % | LoD %", "ГпС %, G2 AUW | LoD %, G2 AUW", "Крит. | Crit."]
    half = (len(ss) + 1) // 2
    rows = [[(x, 1, "h") for x in head + head]]
    for i in range(half):
        row = []
        for s in (ss[i], ss[i + half] if i + half < len(ss) else None):
            if s is None:
                row += ["", "", "", "", ""]
            else:
                row += [(s["sid"], 1, "b"), mk(s["m0"], 4 if s["m0"].as_tuple().exponent < -3 else 3), mk(s["lp2"]),
                        mk(s["la2"]), ", ".join(map(str, s["crit"])) or "—"]
        rows.append(row)
    bx.grid(d, [1.3, 1.75, 1.75, 2.45, 1.98] * 2, rows, sz=8, head=1)
    pr.note(d, "ГпС %% = (G1 − G2) ÷ m₀ × 100, G1 = m_B + m₀, со m_B, G1 и G2 на прецизната вага. „G2 AUW“: истата пресметка "
               "со G2 од AUW220D (Прилог 1 кон %s). Вредностите се од RUN1_CHECK.tsv. Ознаките S1–S%d не се поврзани со "
               "серии (табела 2, ред 5)." % (LOD, S["n"]),
            "LoD %% = (G1 − G2) ÷ m₀ × 100, G1 = m_B + m₀, with m_B, G1 and G2 on the precision balance. “G2 AUW”: the same "
            "calculation with G2 from the AUW220D (Attachment 1 to %s). Values from RUN1_CHECK.tsv. The IDs S1–S%d are not "
            "tied to batches (Table 2, row 5)." % (LOD, S["n"]))
    return d


# ----------------------------------------------------------------------------- build, convert, check
def pdf_checks(pdf):
    info = subprocess.run(["pdfinfo", pdf], capture_output=True, text=True).stdout
    pages = int([l for l in info.splitlines() if l.startswith("Pages:")][0].split()[1])
    size = [l for l in info.splitlines() if l.startswith("Page size:")][0]
    assert "(A4)" in size, size
    fonts = subprocess.run(["pdffonts", pdf], capture_output=True, text=True).stdout.splitlines()[2:]
    not_emb = [l for l in fonts if l.split()[-5] != "yes"]
    assert not not_emb, not_emb
    return pages, [l.split()[0] for l in fonts]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--work", help="folder for the DOCX and the conversion (outside the repository)")
    a = ap.parse_args()
    work = os.path.abspath(a.work) if a.work else tempfile.mkdtemp(prefix="pp_inf01_")
    if os.path.commonpath([work, bc.REPO]) == bc.REPO:
        raise SystemExit("--work must be outside the repository (no DOCX in the repo)")
    os.makedirs(work, exist_ok=True)
    t1, t2 = tranches()
    d = build(run1(), t1, t2)
    bc.glyph_audit(d)
    docx = os.path.join(work, STEM + ".docx")
    pf.save(d, docx)
    if not bc.verify(docx):
        raise SystemExit("pp_verify FAIL")
    profile = "file://" + os.path.join(work, "lo_profile")     # own LibreOffice profile: safe beside other sessions
    subprocess.run(["soffice", "-env:UserInstallation=" + profile, "--headless", "--convert-to", "pdf", "--outdir", work,
                    docx], check=True, capture_output=True)
    pdf = os.path.join(work, STEM + ".pdf")
    pages, fonts = pdf_checks(pdf)
    os.makedirs(OUT_DIR, exist_ok=True)
    target = os.path.join(OUT_DIR, STEM + ".pdf")
    shutil.copyfile(pdf, target)
    print("%s  %d pages, A4, fonts embedded: %s" % (CODE, pages, ", ".join(fonts)))
    print("PDF :", target)
    print("DOCX:", docx, "(not copied to the repository)")


if __name__ == "__main__":
    main()
