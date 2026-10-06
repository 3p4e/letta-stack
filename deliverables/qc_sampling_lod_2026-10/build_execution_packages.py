#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Build the two execution packages of PP-QC-SP-002/26 with the Purely Plant document engine.

Head of QC, 06.10.2026: the sampling of the bags and the loss-on-drying analysis are two separate
documentation packages, each on the company's own forms:

  Package 2 — sampling execution, per sampling day (out/2_SAMPLING_EXECUTION/):
    QCT 024        transfer, secure warehouse -> sampling room (receipt of the selected bags)
    QCSOP 011_A03  visual inspection on opening, one form per batch
    QCT 021        material quantity review: before / sampled / after, per bag
    QASOP_031_A05  "SAMPLED" bag labels and QASOP_031_A07 sample labels (build_labels.py)
    QCT 024        transfer, sampling room -> secure warehouse (return of the sampled bags)
    QCT 024        transfer, sampling room -> QC laboratory (the composite samples)
  Package 3 — loss-on-drying analysis execution, one record per group of samples analysed
    together (out/3_LOD_ANALYSIS_EXECUTION/): receipt, equipment, homogenisation and test
    portions, weighings to constant mass, results, deviations, sign-offs.

The forms reproduce the July 2026 layouts field for field (QCT 024 v01, QCSOP 011_A03 v7.0,
QCT 021 v01); only the values are prefilled. Every number comes from SAMPLING_PLAN_T1_T2_2026-10.tsv
and bag_selection.tsv (campaign_data.py). Each DOCX must pass pp_verify.py.

    python3 build_execution_packages.py
"""
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
import build_campaign_docs as bc                       # noqa: E402  (engine on sys.path, helpers, constants)
from build_campaign_docs import pf, pr, num, ranges, BLANK  # noqa: E402
from docx.enum.table import WD_TABLE_ALIGNMENT          # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH           # noqa: E402
from docx.oxml import OxmlElement                       # noqa: E402
from docx.oxml.ns import qn                             # noqa: E402
from docx.table import _Cell                            # noqa: E402

OUT_S = os.path.join(bc.OUT, "2_SAMPLING_EXECUTION")
OUT_L = os.path.join(bc.OUT, "3_LOD_ANALYSIS_EXECUTION")
CODE = bc.CODE
PORTRAIT_W, LANDSCAPE_W = 18.46, 27.16

WAREHOUSE = "Сеф-магацин | Secure warehouse"
QC = "Контрола на квалитет | QC"
SAMPLING_ROOM = "Просторија за земање примероци | Sampling room"
QC_LAB = "КК лабораторија | QC laboratory"
SID = "____/26_SFR-PC-___"          # QCSOP 011 v03 §6.2.1 sample code, written at sampling


def rec(day, kind):
    return "%s-D%d-%s" % (CODE, day, kind)


def lod_code(day):
    return "%s-LOD-%02d" % (CODE, day)


def day_lots(lots, day):
    return [l for l in lots if l["day"] == day]


def batch_lot(l):
    """Batch with its P lot; a lot with no P lot on record prints the batch alone."""
    return l["batch"] if l["p_lot"] in ("", "—", None) else "%s · %s" % (l["batch"], l["p_lot"])


def chapter(d, num_, mk, en):
    """Engine chapter heading kept on the page of the table or text that follows it."""
    p = pr.chapter(d, num_, mk, en)
    p.paragraph_format.keep_with_next = True
    return p


def warehouses(ls):
    out = []
    for l in ls:
        if l["warehouse"] not in out:
            out.append(l["warehouse"])
    return out


# ----------------------------------------------------------------------------- table helpers
KIND = {
    "h": (pr.NAVYF, pr.WHITE, True, False),     # navy header
    "l": (pr.LBL, pr.BLACK, True, False),       # label
    "v": (None, pr.BLACK, False, False),        # value / write-in
    "b": (None, pr.BLACK, True, False),         # bold value
    "s": ("DDE7F1", pr.BLACK, True, False),     # batch sub-header
    "t": (pr.LBL, pr.BLACK, True, False),       # total
    "n": (None, pr.BLACK, False, True),         # left-aligned note
}


def grid(d, widths, rows, sz=8, head=0, mode="full", heights=None):
    """A table from row specs. A cell is 'text', (text, span) or (text, span, kind); spans of a row
    must add up to the column count. `head` rows repeat on every page the table spans."""
    ncol = len(widths)
    t = d.add_table(rows=len(rows), cols=ncol)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        tr = t._tbl.tr_lst[i]
        tcs = list(tr.tc_lst)            # direct row access: table.cell() rescans the whole table per call
        j = 0
        for c in row:
            if isinstance(c, str):
                c = (c, 1, "v")
            elif len(c) == 2:
                c = (c[0], c[1], "v")
            text, span, kind = c
            tc = tcs[j]
            if span > 1:                 # horizontal merge: one cell spanning `span` grid columns
                for extra in tcs[j + 1:j + span]:
                    tr.remove(extra)
                tc.grid_span = span
            cell = _Cell(tc, t)
            fill, col, bold, left = KIND[kind]
            pr.cellfmt(cell, text, None, sz, col, bold=bold, fill=fill)
            if left:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            j += span
        if j != ncol:
            raise ValueError("row %d spans %d of %d columns: %r" % (i, j, ncol, row))
    pr.fixed(t, widths, mode=mode, header_repeat=False)
    if head:
        pr._repeat_header(t, head)
    pr.borders(t)
    for i, h in (heights or {}).items():
        row_height(t.rows[i], h)
    return t


def row_height(row, cm):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement("w:trHeight")
    e.set(qn("w:val"), str(int(cm * 567)))
    e.set(qn("w:hRule"), "atLeast")
    trPr.append(e)


def new_form(code, version, mk, en, orient="portrait"):
    pr.PAGE_W = PORTRAIT_W if orient == "portrait" else LANDSCAPE_W
    return pf.new_annex(code=code, version=version, mk_title=mk, en_title=en, orient=orient, status=bc.STATUS)


def caption(d, mk, en):
    """One small line naming which movement or day this copy of the form records."""
    p = d.add_paragraph()
    pr.sp(p, 2, 4)
    pr.rin(p, mk, 10, pr.NAVY, bold=True)
    pr.rin(p, "  |  ", 9, pr.GREY)
    pr.rin(p, en, 9, pr.GREY, ital=True)


def page_break(d):
    d.add_page_break()


# ----------------------------------------------------------------------------- QCT 024 (July layout)
def qct024(record, caption_mk, caption_en, frm, to, ref, batch_lot, note, items, total,
           operator=("Оператор | Operator", "Супервизор | Supervised by"),
           confirm=("Доставено од | Delivered by", "Прием од | Receipt by")):
    d = new_form("QCT 024", "01", "Запис за пренос на материјал", "Material transfer record")
    caption(d, caption_mk, caption_en)
    w6 = [PORTRAIT_W / 6] * 6
    grid(d, w6, [
        [("Запис број | Record Number", 2, "l"), (record, 4, "b")],
        [("ПРЕДАВАЊЕ ОД | HANDOVER FROM", 3, "h"), ("ПРИЕМ ВО | HANDOVER TO", 3, "h")],
        [("Оддел | Department", 1, "l"), ("Просторија | Room", 1, "l"), ("Процес | Process", 1, "l")] * 2,
        list(frm) + list(to),
    ])
    pr.gap(d, 6)
    grid(d, [2.3, 2.9, 2.9, 4.86, 2.3, 3.2], [
        [("Информации за преносот | Transfer information", 6, "h")],
        [("Датум | Date", 1, "l"), BLANK, ("Референтен документ | Reference document", 1, "l"), ref,
         ("Серија / Лот | Batch / Lot", 1, "l"), batch_lot],
        [("Забелешка | Note", 1, "l"), (note, 5, "n")],
    ])
    pr.gap(d, 6)
    grid(d, [4.6, 4.63, 4.6, 4.63], [
        [(operator[0], 2, "h"), (operator[1], 2, "h")],
        [("Име Презиме | Name Surname", 1, "l"), ("Датум и потпис | Date and signature", 1, "l")] * 2,
        ["", "", "", ""],
    ], heights={2: 0.9})
    pr.gap(d, 6)
    rows = [[("Попис на контејнери/кутии/кеси — нето/бруто | Container/box/bag inventory — net/gross", 5, "h")],
            [("№", 1, "l"), ("Контејнер/Кутија/Кеса бр. | Container/Box/Bag No.", 1, "l"),
             ("Нето (g) | Net (g)", 1, "l"), ("Бруто (g) | Gross (g)", 1, "l"), ("Забелешка | Remark", 1, "l")]]
    rows += items
    rows.append(total)
    grid(d, [1.0, 6.2, 2.8, 2.8, 5.66], rows, sz=8, head=2)
    pr.gap(d, 6)
    grid(d, [3.1, 3.1, 3.03, 3.1, 3.1, 3.03], [
        [("Потврда | Confirmation", 6, "h")],
        [(confirm[0], 3, "l"), (confirm[1], 3, "l")],
        [("Име Презиме | Name Surname", 1, "l"), ("Потпис | Signature", 1, "l"), ("Датум | Date", 1, "l")] * 2,
        ["", "", "", "", "", ""],
    ], heights={3: 0.9})
    return d


def bag_items(ls, bags, day, net, remark=""):
    items, i = [], 0
    for l in ls:
        items.append([("%s · %s · %s · магацин/store %s · картони/cartons %s · n = %d"
                       % (batch_lot(l), l["strain"], l["grade"] or "—", l["warehouse"],
                          ranges(l["carton_list"]), l["n"]), 5, "s")])
        for b in [x for x in bags if int(x["day"]) == day and x["batch"] == l["batch"]]:
            i += 1
            items.append([str(i), b["bag_id"], net, "", remark])
    return items, i


def build_receipt(day, lots, bags):
    ls = day_lots(lots, day)
    items, nb = bag_items(ls, bags, day, "400,0")
    whs = " · ".join(warehouses(ls))
    return qct024(
        rec(day, "RCPT"),
        "ПРИЕМ — Ден %d: сеф-магацин → просторија за земање примероци" % day,
        "RECEIPT — Day %d: secure warehouse → sampling room" % day,
        (WAREHOUSE, "Магацин | Store %s" % whs, "Складирање | Storage"),
        (QC, SAMPLING_ROOM, "Земање примероци | Sampling"),
        "%s · PP-QC-MHR-___/26" % CODE,
        "%d серии | %d batches" % (len(ls), len(ls)),
        "Картоните на секоја серија се повлекуваат; избраните кеси (K{картон}B{кеса}, план 4b) се попишани поединечно. "
        "Нето = етикетата на примарното пакување (400,0 g); бруто го мери КК при прием. Ако документираниот број кеси се "
        "разликува од N на планот, се запишува тука и изборот се пресметува повторно пред отворање. | The cartons of each "
        "batch are retrieved; the selected bags (K{carton}B{bag}, plan 4b) are listed one by one. Net = the primary-pack "
        "label (400.0 g); gross is weighed by QC at receipt. If the documented bag count differs from the plan's N, it is "
        "written here and the selection is recomputed before opening.",
        items,
        [("ВКУПНО | TOTAL", 2, "t"), (num(nb * 400, 1), 1, "t"), ("", 1, "t"), ("%d кеси / bags" % nb, 1, "t")],
        confirm=("Доставено од (магацин) | Delivered by (warehouse)", "Прием од (КК) | Receipt by (QC)"))


def build_return(day, lots, bags):
    ls = day_lots(lots, day)
    items, nb = bag_items(ls, bags, day, "")
    return qct024(
        rec(day, "RET"),
        "ВРАЌАЊЕ — Ден %d: просторија за земање примероци → сеф-магацин" % day,
        "RETURN — Day %d: sampling room → secure warehouse" % day,
        (QC, SAMPLING_ROOM, "Земање примероци | Sampling"),
        (WAREHOUSE, "Магацин | Store %s" % " · ".join(warehouses(ls)), "Складирање | Storage"),
        "%s · QCT 021 Ден | Day %d" % (CODE, day),
        "%d серии | %d batches" % (len(ls), len(ls)),
        "Враќање: ново нето и бруто по кеса од QCT 021 (после мострирање); кесите се затворени и означени „МОСТРИРАНО“ "
        "(QASOP_031_A05_v1). Картоните се враќаат на истата позиција. | Return: new net and gross per bag from "
        "QCT 021 (after sampling); the bags are closed and labelled “SAMPLED” (QASOP_031_A05_v1). The cartons go back to "
        "the same position.",
        items,
        [("ВКУПНО | TOTAL", 2, "t"), ("", 1, "t"), ("", 1, "t"), ("%d кеси / bags" % nb, 1, "t")],
        confirm=("Доставено од (КК) | Delivered by (QC)", "Прием од (магацин) | Receipt by (warehouse)"))


def build_sample_transfer(day, lots):
    ls = day_lots(lots, day)
    items = []
    for i, l in enumerate(ls, start=1):
        items.append([str(i), "%s · %s" % (l["batch"], SID), "", "",
                      "k = %d · %d цвета / flowers" % (l["k"], l["n"])])
    return qct024(
        rec(day, "SMP"),
        "ПРИМЕРОЦИ — Ден %d: просторија за земање примероци → КК лабораторија" % day,
        "SAMPLES — Day %d: sampling room → QC laboratory" % day,
        (QC, SAMPLING_ROOM, "Земање примероци | Sampling"),
        (QC, QC_LAB, "Губиток при сушење | Loss on drying"),
        "%s · QCT 021 Ден | Day %d → %s" % (CODE, day, lod_code(day)),
        "%d серии | %d batches" % (len(ls), len(ls)),
        "Збирни примероци, по еден за серија (сите n цвета на серијата), во затворени чисти сади означени по "
        "QASOP_031_A07; нето = Σ мострирано на серијата од QCT 021. Се анализираат заедно на %s. | Composite samples, one "
        "per batch (all n flowers of the batch), in closed clean containers labelled per QASOP_031_A07; net = the batch's "
        "Σ sampled from QCT 021. They are analysed together on %s." % (lod_code(day), lod_code(day)),
        items,
        [("ВКУПНО | TOTAL", 2, "t"), ("", 1, "t"), ("", 1, "t"), ("%d примероци / samples" % len(ls), 1, "t")],
        confirm=("Доставено од (земање) | Delivered by (sampling)", "Прием од (КК лабораторија) | Receipt by (QC laboratory)"))


# ----------------------------------------------------------------------------- QCSOP 011_A03 (v7.0 layout)
CRITERIA = [
    ("Боја и изглед | Color & Appearance", "Здрава, без дисколорација | Healthy, no discoloration"),
    ("Мирис | Odor", "Карактеристичен, без непријатни | Characteristic, no off-odors"),
    ("Структура | Flower Structure", "Густи, добро формирани | Dense, well-formed buds"),
    ("Трихоми | Trichome Integrity", "Целосни и видливи | Intact and visible"),
    ("Мувла/Пепелница | Mold/Mildew", "Не е детектирано | None detected"),
    ("Штетници | Pests", "Не е детектирано | None detected"),
    ("Семки | Seeds", "Не е детектирано (ако е лек) | None (if medicine)"),
    ("Страни материи | Foreign Matter", "≤ 2% (Ph. Eur. 2.8.2) | ≤ 2% (Ph. Eur. 2.8.2)"),
    ("Контаминација | Contamination", "Без прашина, влакна, честички | No dust, fibers, particles"),
]


def build_a03(day, lots, bags):
    ls = day_lots(lots, day)
    d = new_form("QCSOP 011_A03", "7.0", "Формулар за визуелна инспекција на канабис цвет",
                 "Cannabis flower visual inspection form")
    for idx, l in enumerate(ls):
        if idx:
            page_break(d)
        caption(d, "Ден %d · серија %d од %d · %s" % (day, idx + 1, len(ls), CODE),
                "Day %d · batch %d of %d · %s" % (day, idx + 1, len(ls), CODE))
        grid(d, [2.6, 3.6, 2.2, 4.06, 2.0, 4.0], [
            [("Идентификација на инспекцијата | Inspection Identification", 6, "h")],
            [("Запис бр. | Record No.", 1, "l"), ("QCSOP 011_A03-___/26", 5, "v")],
            [("Серија | Batch", 1, "l"), (batch_lot(l), 1, "b"), ("Сорта | Strain", 1, "l"),
             "%s · %s" % (l["strain"], l["grade"] or "—"), ("Датум | Date", 1, "l"), BLANK],
            [("Просторија | Room", 1, "l"), SAMPLING_ROOM, ("Инспектирал | Inspector", 1, "l"), ("", 3, "v")],
            [("Фаза на инспекција | Inspection stage", 6, "h")],
            [("[ ] Берба   [ ] Триминг   [ ] Пред пакување   [X] По отворање на секоја кеса (SP-12) | "
              "[ ] Harvest   [ ] Trimming   [ ] Pre-packaging   [X] Post-opening of each bag (SP-12)", 6, "v")],
        ])
        pr.gap(d, 5)
        rows = [[("Критериуми на инспекција (ниво на серија) | Inspection criteria (batch level)", 3, "h")],
                [("Параметар | Parameter", 1, "l"), ("Критериум | Criteria", 1, "l"), ("Резултат | Result", 1, "l")]]
        rows += [[(p, 1, "l"), c, "[ ] Одговара | Conforms"] for p, c in CRITERIA]
        grid(d, [5.2, 8.26, 5.0], rows, sz=8)
        pr.note(d, "За секоја избрана кеса, пред земање, изврши и евидентирај визуелна инспекција. Кое било „Не одговара“ → "
                   "карантин и отстапување (PP-QA-SOP-003).",
                "For each selected bag, before sampling, perform and record a visual inspection. Any “Does not conform” → "
                "quarantine and deviation (PP-QA-SOP-003).")
        rows = [[("Поединечна инспекција при отворање — по кеса (SP-12) | Per-bag inspection on opening (SP-12)", 9, "h")],
                [("№", 1, "l"), ("Кеса бр. | Bag No.", 1, "l"), ("Боја/мирис | Color/Odor", 1, "l"),
                 ("Мувла | Mould", 1, "l"), ("Штетници | Pests", 1, "l"), ("Страни м./семки | FM/Seeds", 1, "l"),
                 ("Оштет. пак. | Pack dmg", 1, "l"), ("Севкупно | Overall", 1, "l"), ("Инсп. | Init.", 1, "l")]]
        bl = [b for b in bags if int(b["day"]) == day and b["batch"] == l["batch"]]
        rows += [[str(i), (b["bag_id"], 1, "b"), "", "", "", "", "", "", ""] for i, b in enumerate(bl, start=1)]
        rows.append([("Одлука | Decision", 2, "l"),
                     ("[ ] Одобрено за следна фаза   [ ] Потребна доработка   [ ] Одбиено | "
                      "[ ] Approved for next stage   [ ] Requires rework   [ ] Rejected", 7, "v")])
        grid(d, [1.0, 2.2, 2.3, 1.9, 2.0, 2.5, 2.2, 2.3, 2.06], rows, sz=8, head=2)
        pr.gap(d, 5)
        grid(d, [5.0, 6.06, 3.4, 4.0], [
            [("Улога | Role", 1, "h"), ("Име/Позиција | Name/Position", 1, "h"), ("Датум | Date", 1, "h"),
             ("Потпис | Signature", 1, "h")],
            [("Инспектирал | Inspected by", 1, "l"), "", "", ""],
            [("Одобрил (КК) | Approved by (QC)", 1, "l"), "", "", ""],
        ], heights={1: 0.8, 2: 0.8})
    return d


# ----------------------------------------------------------------------------- QCT 021 (v01 layout)
def build_qct021(day, lots, bags):
    ls = day_lots(lots, day)
    d = new_form("QCT 021", "01", "Преглед на количини на материјали и дневно усогласување на магацинската состојба",
                 "Material quantity review & daily warehouse reconciliation", orient="landscape")
    caption(d, "ПРЕД И ПОСЛЕ МОСТРИРАЊЕ — Ден %d · %d серии · %s" % (day, len(ls), CODE),
            "BEFORE AND AFTER SAMPLING — Day %d · %d batches · %s" % (day, len(ls), CODE))
    grid(d, [2.6, 4.0, 3.6, 10.36, 2.6, 4.0], [
        [("MLR №", 1, "l"), "", ("Референтен документ | Reference document", 1, "l"),
         "%s · QCT 024 %s" % (CODE, rec(day, "RCPT")), ("Датум | Date", 1, "l"), BLANK],
    ], heights={0: 0.8})
    pr.gap(d, 5)
    W = [3.0, 2.0, 2.2, 2.2, 4.6, 2.2, 2.2, 2.36, 2.0, 2.2, 2.2]
    rows = [[("№", 1, "h"), ("Пред мострирање (g) | Before sampling", 3, "h"), ("Мострирано (g) | Sampled", 4, "h"),
             ("После мострирање (g) | After sampling", 3, "h")],
            [("Серија | Batch", 1, "l"), ("Кеса бр. / ID код | Bag No. / ID code", 1, "l"), ("Нето | Net (g)", 1, "l"),
             ("Бруто | Gross (g)", 1, "l"), ("ID код на мостра | Sample ID Code (sID)", 1, "l"), ("Нето | Net (g)", 1, "l"),
             ("Бруто | Gross (g)", 1, "l"), ("Намена | Use", 1, "l"), ("Кеса бр. / ID код | Bag No. / ID code", 1, "l"),
             ("Нето | Net (g)", 1, "l"), ("Бруто | Gross (g)", 1, "l")]]
    nb = 0
    for l in ls:
        bl = [b for b in bags if int(b["day"]) == day and b["batch"] == l["batch"]]
        for b in bl:
            nb += 1
            rows.append([l["batch"], (b["bag_id"], 1, "b"), "400,0", "", "→ збирен / composite", "", "", "ГпС / LoD",
                         (b["bag_id"], 1, "b"), "", ""])
        rows.append([("ВКУПНО / TOTAL %s" % l["batch"], 1, "t"), ("%d кеси / bags" % len(bl), 1, "t"),
                     (num(len(bl) * 400, 1), 1, "t"), ("", 1, "t"), (SID, 1, "t"), ("", 1, "t"), ("", 1, "t"),
                     ("ГпС / LoD · k = %d" % l["k"], 1, "t"), ("%d кеси / bags" % len(bl), 1, "t"), ("", 1, "t"),
                     ("", 1, "t")])
    rows.append([("ВКУПНО (g) | TOTAL (g)", 2, "t"), (num(nb * 400, 1), 1, "t"), ("", 1, "t"),
                 ("%d збирни примероци / composite samples" % len(ls), 1, "t"), ("", 1, "t"), ("", 1, "t"), ("", 1, "t"),
                 ("%d кеси / bags" % nb, 1, "t"), ("", 1, "t"), ("", 1, "t")])
    grid(d, W, rows, sz=7, head=2)
    pr.note(d, "Пред: нето и бруто се препишуваат од QCT 024 (прием). Мострирано: еден цвет по кеса, најголемиот, во збирниот "
               "примерок на серијата; по серија редот ВКУПНО ја носи шифрата на збирниот примерок и неговото нето. После: бруто "
               "се мери по затворање на кесата; нето после = нето пред − мострирано. Каде кесата има картон за кеса "
               "(PO_SOP_007_A14-02), мострирањето се запишува и таму (дејство 01).",
            "Before: net and gross transcribed from QCT 024 (receipt). Sampled: one flower per bag, the largest, into the "
            "batch composite; per batch the TOTAL row carries the composite's sample code and its net. After: gross weighed "
            "after the bag is closed; net after = net before − sampled. Where the bag carries its bag card "
            "(PO_SOP_007_A14-02), the sampling is entered there too (action 01).")
    grid(d, [7.0, 4.0, 6.0], [
        [("Параметар | Parameter", 1, "l"), ("m (g)", 1, "l"), ("Потпис | Signature", 1, "l")],
        [("Отпад | Waste", 1, "l"), "", ""],
        [("Изгубено во процес | Lost in process", 1, "l"), "", ""],
        [("Добиено во процес | Gained in process", 1, "l"), "", ""],
    ], mode="compact")
    pr.gap(d, 6)
    grid(d, [3.2, 2.7, 3.0, 2.0, 2.68, 3.2, 2.7, 3.0, 2.0, 2.68], [
        [("Подготвено од | Prepared by:", 1, "l"), ("Функција | Function:", 1, "l"), ("Име и Презиме | Name & Surname", 1, "l"),
         ("Датум | Date:", 1, "l"), ("Потпис | Signature", 1, "l"), ("Проверено и Одобрено од | Checked & Approved by:", 1, "l"),
         ("Функција | Function:", 1, "l"), ("Име и Презиме | Name & Surname", 1, "l"), ("Датум | Date:", 1, "l"),
         ("Потпис | Signature", 1, "l")],
        ["", "QC персонал | QC Personnel", "", "", "", "", "QC Менаџер | QC Manager", "", "", ""],
    ], heights={1: 0.9})
    return d


# ----------------------------------------------------------------------------- package index
def build_index(lots):
    d = new_form(CODE + "-SMP", bc.VERSION, "Пакет за извршување — земање примероци",
                 "Execution package — sampling")
    chapter(d, "1", "СОДРЖИНА НА ПАКЕТОТ", "Package Contents")
    pr.body(d, "Овој пакет го документира земањето примероци на Транша 1 и 2 според %s, на обрасците на QCSOP 011 и "
               "QASOP_031, по еден комплет за секој ден на земање. Анализата за губиток при сушење е посебен пакет (%s, %s)."
               % (CODE, lod_code(1), lod_code(2)),
            "This package documents the sampling of Tranches 1 and 2 under %s, on the QCSOP 011 and QASOP_031 forms, one set "
            "per sampling day. The loss-on-drying analysis is a separate package (%s, %s)." % (CODE, lod_code(1), lod_code(2)))
    rows = [[("Чекор | Step", 1, "h"), ("Документ | Document", 1, "h"), ("Образец | Form", 1, "h"),
             ("Запис бр. Ден 1 | Record No. Day 1", 1, "h"), ("Запис бр. Ден 2 | Record No. Day 2", 1, "h")]]
    steps = [
        ("1", "Пренос сеф-магацин → просторија за земање | Transfer secure warehouse → sampling room", "QCT 024 v01",
         rec(1, "RCPT"), rec(2, "RCPT")),
        ("2", "Визуелна инспекција при отворање, по серија | Visual inspection on opening, per batch", "QCSOP 011_A03 v7.0",
         "QCSOP 011_A03-___/26 × %d" % len(day_lots(lots, 1)), "QCSOP 011_A03-___/26 × %d" % len(day_lots(lots, 2))),
        ("3", "Мерење пред и после мострирање, по кеса | Weighing before and after sampling, per bag", "QCT 021 v01",
         "MLR № ___ (Ден | Day 1)", "MLR № ___ (Ден | Day 2)"),
        ("4", "Етикети „МОСТРИРАНО“ на кесите | “SAMPLED” labels on the bags", "QASOP_031_A05_v1",
         "%d" % sum(l["n"] for l in day_lots(lots, 1)), "%d" % sum(l["n"] for l in day_lots(lots, 2))),
        ("5", "Етикети на збирните примероци | Labels on the composite samples", "QASOP_031_A07",
         "%d" % len(day_lots(lots, 1)), "%d" % len(day_lots(lots, 2))),
        ("6", "Враќање на кесите во сеф-магацин | Return of the bags to the secure warehouse", "QCT 024 v01",
         rec(1, "RET"), rec(2, "RET")),
        ("7", "Пренос на примероците во КК лабораторија | Transfer of the samples to the QC laboratory", "QCT 024 v01",
         rec(1, "SMP"), rec(2, "SMP")),
    ]
    rows += [[s, (doc, 1, "n"), f, r1, r2] for s, doc, f, r1, r2 in steps]
    grid(d, [1.3, 6.6, 3.0, 3.78, 3.78], rows, sz=8, head=1)
    chapter(d, "2", "РЕДОСЛЕД НА РАБОТА", "Order of Work")
    for mk, en in [
        ("Пред земање: RQS е регистриран (QCSOP 011 v03 §6.1.1); просторијата и приборот се чисти и суви; вагата е проверена.",
         "Before sampling: the RQS is registered (QCSOP 011 v03 §6.1.1); the room and tools are clean and dry; the balance is checked."),
        ("Картоните на серијата се примаат на QCT 024 (прием) со бруто по избрана кеса. Кесата се отвора, се прегледува на A03, "
         "се зема еден цвет, најголемиот, во збирниот сад на серијата, кесата се затвора и се мери (QCT 021).",
         "The batch cartons are received on QCT 024 (receipt) with the gross of each selected bag. The bag is opened, inspected "
         "on A03, one flower, the largest, goes into the batch's composite container, and the bag is closed and weighed (QCT 021)."),
        ("Секоја мострирана кеса добива етикета „МОСТРИРАНО | SAMPLED“; збирниот сад добива етикета QASOP_031_A07 со шифрата "
         "на примерокот. Кесите се враќаат на QCT 024 (враќање), примероците одат во КК лабораторија на QCT 024 (примероци).",
         "Each sampled bag gets a “SAMPLED” label; the composite container gets a QASOP_031_A07 label with the sample code. "
         "The bags go back on QCT 024 (return); the samples go to the QC laboratory on QCT 024 (samples)."),
        ("Истиот ден примероците влегуваат во анализата %s (Ден 1) или %s (Ден 2)." % (lod_code(1), lod_code(2)),
         "The same day the samples enter analysis %s (Day 1) or %s (Day 2)." % (lod_code(1), lod_code(2))),
        ("Документација: истовремено, трајно сино мастило, без празни полиња („N/A“), поправка со една линија, иницијали и "
         "датум (ALCOA+).",
         "Documentation: contemporaneous, permanent blue ink, no blank fields (“N/A”), single-line corrections with initials "
         "and date (ALCOA+)."),
    ]:
        pr.bullet(d, mk, en)
    return d


# ----------------------------------------------------------------------------- LoD execution record
def build_lod(day, lots):
    ls = day_lots(lots, day)
    code = lod_code(day)
    sum_k = sum(l["k"] for l in ls)
    d = bc.new_doc(code, "ИЗВРШЕН ЗАПИС — ГУБИТОК ПРИ СУШЕЊЕ — ПРИМЕРОЦИ ОД ДЕН %d" % day,
                   "EXECUTION RECORD — LOSS ON DRYING — SAMPLES OF DAY %d" % day)
    bc.cover(d, "Извршен запис — губиток при сушење, примероци од ден на земање %d" % day,
             "Execution record — loss on drying, samples of sampling day %d" % day,
             [("Запис бр. | Record No.", code),
              ("План | Plan", CODE + " · Транша 1 и 2 | Tranches 1 and 2"),
              ("Примероци | Samples", "QCT 024 %s · %d збирни примероци | %d composite samples" % (rec(day, "SMP"), len(ls), len(ls))),
              ("Тест порции Σk | Test portions Σk", "%d" % sum_k),
              ("Метод | Method", "Ph. Eur. 2.2.32 (3028) · SAM_a02.2 · 40 °C · 15–25 mbar · 24 h"),
              ("Критериум | Criterion", "≤ 12,0 % w/w (QCSP 001) | ≤ 12.0 % w/w (QCSP 001)"),
              ("Влез во печка / мерење 24 h / второ мерење | Oven in / 24-h weighing / second weighing",
               "%s / %s / %s" % (BLANK, BLANK, BLANK)),
              ("Аналитичар | Analyst", bc.ANALYST)],
             "ИЗВРШЕН ЗАПИС | EXECUTION RECORD", "Loss on Drying — samples of sampling Day %d" % day,
             "Сите серии анализирани заедно", "All batches analysed together")

    chapter(d, "A", "ПРИЕМ НА ПРИМЕРОЦИТЕ", "Receipt of the Samples")
    rows = [["№", "Серија | Batch", "P лот | P lot", "Сорта | Strain", "Шифра на примерок | Sample code",
             "Збирен g (QCT 021) | Composite g", "k", "Состојба | Condition", "Иниц./час | Init./time"]]
    rows = [[(x, 1, "h") for x in rows[0]]]
    for i, l in enumerate(ls, start=1):
        rows.append([str(i), (l["batch"], 1, "b"), l["p_lot"], l["strain"], SID, "", str(l["k"]),
                     "[ ] затворен, означен | closed, labelled", ""])
    grid(d, [0.8, 2.4, 1.7, 2.6, 3.0, 2.0, 0.9, 3.0, 2.06], rows, sz=7, head=1)
    pr.note(d, "Примероците се примени на QCT 024 %s. Масата на збирниот примерок се препишува од QCT 021 (ВКУПНО по серија)."
            % rec(day, "SMP"),
            "The samples are received on QCT 024 %s. The composite mass is transcribed from QCT 021 (TOTAL per batch)."
            % rec(day, "SMP"))
    pr.step_signoff(d, "Потпис за делот A (прием) | Sign-off for section A (receipt)", None)

    chapter(d, "B", "ОПРЕМА, МАТЕРИЈАЛИ И УСЛОВИ", "Equipment, Materials and Conditions")
    bc.kv_table(d, [
        ("Вакуумска печка — ID / статус на квалификација | Vacuum oven — ID / qualification status", BLANK + " / " + BLANK),
        ("Поставени услови | Set conditions", "40 °C · 15–25 mbar · 24 h   (постигнато | achieved: ____ °C · ____ mbar)"),
        ("Аналитичка вага — ID / калибрација до / дневна проверка | Analytical balance — ID / calibration until / daily check",
         BLANK + " / " + BLANK + " / [ ]"),
        ("Контролен тег — номинално / измерено | Check weight — nominal / found", BLANK + " / " + BLANK),
        ("Ексикатор / ексикант | Desiccator / desiccant", BLANK + " / [ ] активен | active"),
        ("Садови за мерење — бр. / претходно исушени под условите на методот | Weighing bottles — nos. / previously dried "
         "under the method conditions", BLANK + " / [ ]"),
        ("Прибор за сечење / подлога | Cutting tools / tray", "нерѓосувачки челик, чисти и суви [ ] | stainless steel, clean and dry [ ]"),
        ("Амбиентални услови (°C, % RH) | Ambient conditions (°C, % RH)", BLANK),
    ])
    pr.step_signoff(d, "Потпис за делот B | Sign-off for section B", None)

    chapter(d, "C", "ХОМОГЕНИЗАЦИЈА И ТЕСТ ПОРЦИИ", "Homogenisation and Test Portions")
    for mk, en in [
        ("Сите цветови на серијата се сечат грубо со чисти ножици на чиста подлога и се мешаат со четвртирање, двапати "
         "(израмнување, четири четвртини, спротивните четвртини се спојуваат).",
         "All flowers of the batch are coarsely cut with clean scissors on a clean tray and mixed by quartering, twice "
         "(flatten, four quarters, opposite quarters recombined)."),
        ("Од различни четвртини се земаат k тест порции од околу 1,000 g во претходно исушени тарирани садови (m0) и веднаш "
         "се мерат (G1).",
         "From different quarters, k test portions of about 1.000 g are taken into previously dried, tared bottles (m0) and "
         "weighed at once (G1)."),
        ("Остатокот се чува затворен и означен до одобрувањето на овој запис.",
         "The remainder is kept closed and labelled until this record is approved."),
    ]:
        pr.bullet(d, mk, en)
    rows = [[(x, 1, "h") for x in ["№", "Серија | Batch", "Сечење | Cutting", "Четвртирање ×2 | Quartering ×2",
                                   "Порции k | Portions k", "Остаток затворен, означен | Remainder closed, labelled",
                                   "Иниц./час | Init./time"]]]
    for i, l in enumerate(ls, start=1):
        rows.append([str(i), (l["batch"], 1, "b"), "[ ]", "[ ]", str(l["k"]), "[ ]", ""])
    grid(d, [0.8, 2.8, 2.2, 3.0, 2.2, 4.4, 3.06], rows, sz=8, head=1)
    pr.step_signoff(d, "Потпис за делот C | Sign-off for section C", None)

    chapter(d, "D", "МЕРЕЊА И РЕЗУЛТАТИ ПО ТЕСТ ПОРЦИЈА", "Weighings and Results per Test Portion")
    rows = [[(x, 1, "h") for x in ["№", "Серија | Batch", "Порција | Portion", "Сад бр. | Bottle No.", "m0 g", "G1 g",
                                   "G2 (24 h) g", "G2 (второ) g | G2 (second) g", "Δ mg", "ГпС % | LoD %",
                                   "Средно % | Mean %", "≤ 12,0 % | ≤ 12.0 %"]]]
    r = 0
    for l in ls:
        for p in range(1, l["k"] + 1):
            r += 1
            rows.append([str(r), (l["batch"], 1, "b"), "%d/%d" % (p, l["k"]), "", "", "", "", "", "", "", "", "[ ]"])
    grid(d, [0.8, 2.3, 1.3, 1.3, 1.6, 1.6, 1.7, 1.9, 1.2, 1.6, 1.6, 1.56], rows, sz=7, head=1)
    pr.note(d, "ГпС % = (G1 − G2) ÷ (G1 − m0) × 100 по порција, со G2 = последното мерење до константна маса (две последователни "
               "мерења се разликуваат за не повеќе од 0,5 mg; Δ = разлика помеѓу 24-часовното и второто мерење). Средната вредност "
               "од k порции е резултатот на серијата. Резултат надвор од спецификација → QCSOP 014 и дел F.",
            "LoD % = (G1 − G2) ÷ (G1 − m0) × 100 per portion, with G2 = the last weighing to constant mass (two consecutive "
            "weighings differ by not more than 0.5 mg; Δ = difference between the 24-h and the second weighing). The mean of k "
            "portions is the batch result. An out-of-specification result → QCSOP 014 and section F.")
    pr.step_signoff(d, "Потпис за делот D (определување) | Sign-off for section D (determination)", None)

    chapter(d, "E", "РЕЗУЛТАТИ ПО СЕРИЈА", "Results per Batch")
    rows = [[(x, 1, "h") for x in ["№", "Серија | Batch", "P лот | P lot", "k", "Средно ГпС % | Mean LoD %",
                                   "Критериум | Criterion", "Одговара | Conforms", "Сушење од–до | Drying from–to",
                                   "Иниц. | Init."]]]
    for i, l in enumerate(ls, start=1):
        rows.append([str(i), (l["batch"], 1, "b"), l["p_lot"], str(l["k"]), "", "≤ 12,0 %", "[ ] да | yes  [ ] не | no", "", ""])
    grid(d, [0.8, 2.6, 1.8, 0.8, 2.3, 2.0, 3.0, 3.3, 1.86], rows, sz=7, head=1)
    pr.note(d, "Резултатот и датумите на сушење влегуваат во регистарот на сертификати по правилата на бирото за сертификати.",
            "The result and the drying dates enter the certificate register under the certificate desk's rules.")

    chapter(d, "F", "ОТСТАПУВАЊА / OOS", "Deviations / OOS")
    rows = [[(x, 1, "h") for x in ["№", "Серија / порција | Batch / portion", "Опис | Description", "Дејство | Action",
                                   "Иниц./датум | Init./date"]]]
    rows += [[str(i), "", "", "", ""] for i in range(1, 7)]
    grid(d, [0.8, 3.4, 6.4, 5.0, 2.86], rows, sz=8, head=1)
    pr.note(d, "„Нема отстапувања | No deviations“ се запишува ако табелата останува празна. Отстапувањата се водат по "
               "PP-QA-SOP-003; OOS по QCSOP 014.",
            "“No deviations” is written if the table stays empty. Deviations follow PP-QA-SOP-003; OOS follows QCSOP 014.")

    d.add_page_break()
    chapter(d, "G", "ЗАКЛУЧНИ ПОТПИСИ", "Closing Sign-offs")
    pr.execution_signoff(d, mk_exec="Извршил (КК аналитичар) | Executed (QC Analyst)", reviewer="J. Romevska",
                         approver="B. Nikolov, M.Pharm. (Раководител на КК оддел | QC Department Manager)")
    return d


# ----------------------------------------------------------------------------- main
SAMPLING_DOCS = [  # (step, stem template, builder)
    ("1", "S%d-1_QCT024_Transfer_Warehouse_to_Sampling_Day%d", lambda day, lots, bags: build_receipt(day, lots, bags)),
    ("2", "S%d-2_QCSOP011_A03_Visual_Inspection_Day%d", lambda day, lots, bags: build_a03(day, lots, bags)),
    ("3", "S%d-3_QCT021_Before_After_Sampling_Day%d", lambda day, lots, bags: build_qct021(day, lots, bags)),
    ("6", "S%d-6_QCT024_Transfer_Return_to_Warehouse_Day%d", lambda day, lots, bags: build_return(day, lots, bags)),
    ("7", "S%d-7_QCT024_Transfer_Samples_to_QC_Lab_Day%d", lambda day, lots, bags: build_sample_transfer(day, lots)),
]
INDEX_STEM = "S0_Package_Index_Sampling_Execution"
LOD_STEM = "PP-QC-SP-002_26-LOD-%02d_LoD_Execution_Record_Day%d_Samples"


def main():
    os.makedirs(OUT_S, exist_ok=True)
    os.makedirs(OUT_L, exist_ok=True)
    lots, bags = bc.load()
    docs = [(os.path.join(OUT_S, INDEX_STEM + ".docx"), build_index(lots))]
    for day in (1, 2):
        for _, stem, fn in SAMPLING_DOCS:
            docs.append((os.path.join(OUT_S, (stem % (day, day)) + ".docx"), fn(day, lots, bags)))
        docs.append((os.path.join(OUT_L, (LOD_STEM % (day, day)) + ".docx"), build_lod(day, lots)))
    ok = True
    for path, d in docs:
        bc.glyph_audit(d)
        pf.save(d, path)
        ok = bc.verify(path) and ok
    if not ok:
        raise SystemExit("pp_verify FAIL")
    print("written:", OUT_S, OUT_L)


if __name__ == "__main__":
    main()
