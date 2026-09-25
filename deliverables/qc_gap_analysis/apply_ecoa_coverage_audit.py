#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""The eCoA Coverage Audit: every external certificate in the renamed eCoA_DATABASE,
held against CoQ_Analysis_Master_v48 to see whether the desk's own workbook accounts
for it — cited as a Certificate of Quality's source, on file but not credited,
mentioned only in passing, or on Drive and in no record of the desk at all.

    python3 deliverables/qc_gap_analysis/apply_ecoa_coverage_audit.py

OI-42 (16.09.2026) asked this question of the pre-rename corpus (490 scans: 449
laboratory, 41 in-house) and found forty-four unrecorded. The corpus has grown since
— the 20.09.2026 rename covers 520 files, 479 of them laboratory-issued — and this is
that same check run fresh against it, checked twice: independently verified against
its own first pass (a citation with a trailing Macedonian annotation that broke an
exact-string match, and a Reference-sheet mention applied to four sibling
certificates but missed on a fifth, both corrected here rather than carried over).

This writes a STANDALONE companion workbook and deliberately does not modify
CoQ_Analysis_Master. Two attempts at modifying it established why:

  * Inserting the SHEETS / VERSION HISTORY / OPEN ITEMS lines into their proper
    blocks shifted every section beneath them by two rows and broke 38 of
    verify_workbook.py's deeper checks, on a master that verifies clean — exactly
    as the master's own CONVENTIONS section warns ("insert a row and every code
    beneath moves by one"). A pure openpyxl round-trip with no edits verifies
    clean, so the hazard is moving rows, not the round-trip.
  * Appending them below those blocks shifts nothing and keeps the deeper checks
    clean, but then verify_workbook's Read Me check fails: every sheet in the
    master must be described inside the Read Me section, and that section is a
    FIXED-BOUNDS block (rows 2-83, from the `_fold_Read_Me` defined name written
    by fold_reference_sheet). It is full — its only free rows are the blank
    gutters between its subsections — and Open Items (846-905) has no free row
    at all. A thirteenth sheet cannot be added from outside without growing those
    ranges, and the only thing that grows them coherently is
    build_tracker_v8.py's own fold machinery.

So the master keeps its twelve sheets and its clean verification, and the audit
ships beside it. Folding this into the master properly is a build_tracker_v8.py
change plus a full rebuild — worth doing, but it is a rebuild of a controlled
record and not something to bolt on from outside.
"""
import datetime as dt
import json
import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "tracker", "CoQ_Analysis_Master_v48.xlsx")
OUT = os.path.join(HERE, "tracker", "eCoA_Coverage_Audit_2026-09-20.xlsx")
DATA = os.path.join(HERE, "tracker", "ecoa_coverage_status_2026-09-20.json")

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
STATUS_FILL = {
    "cited": PatternFill("solid", fgColor="C6EFCE"),          # matches the tracker's ✓ green
    "also_on_file": PatternFill("solid", fgColor="FFF2CC"),   # matches the amber "on file, not credited"
    "other_sheet_only": PatternFill("solid", fgColor="EDEDED"),
    "absent": PatternFill("solid", fgColor="F4CCCC"),         # matches the tracker's ✗ red
}
STATUS_LABEL = {
    "cited": "Cited as CoQ source",
    "also_on_file": "Also on file, not credited",
    "other_sheet_only": "Mentioned elsewhere only",
    "absent": "Not in the workbook",
}
COLS = ["Lab", "P Batch", "CU Batch", "Doc Code", "Date", "Status", "Evidence", "Evidence Detail", "Filename"]


def build_sheet(wb, rows):
    if "eCoA Coverage Audit" in wb.sheetnames:
        del wb["eCoA Coverage Audit"]
    ws = wb.create_sheet("eCoA Coverage Audit")
    ws.append(COLS)
    for c in ws[1]:
        c.font = HEADER_FONT
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = "A2"

    for r in rows:
        ws.append([
            r["lab"], r["p_batch"] or "", r["cu_batch"] or "", r["doc_code"], r["date"],
            STATUS_LABEL[r["status"]], r["evidence"], r["evidence_detail"], r["filename"],
        ])
        row_idx = ws.max_row
        fill = STATUS_FILL[r["status"]]
        for col in range(1, len(COLS) + 1):
            ws.cell(row=row_idx, column=col).fill = fill

    widths = [10, 12, 16, 24, 12, 24, 34, 44, 44]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(COLS)), ws.max_row)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = 8  # A3, matching the workbook's other findings sheets
    ws.page_setup.fitToWidth = 1
    ws.print_options.horizontalCentered = False
    return ws


def counts(rows):
    c = {}
    for r in rows:
        c[r["status"]] = c.get(r["status"], 0) + 1
    return c


def find_row(ws, col0_value):
    for row in ws.iter_rows(min_row=1):
        if row[0].value == col0_value:
            return row[0].row
    return None


def read_me(wb, rows, c, master):
    """This companion's own Read Me: what the audit sheet is, and the three entries
    the master's Reference sheet should carry once build_tracker_v8.py folds this
    finding in — drafted here so the next rebuild has text to place rather than
    text to write."""
    ws = wb.create_sheet("Read Me", 0)

    sheets_text = (
        "One row per external (non-in-house) laboratory certificate in the renamed "
        "eCoA_DATABASE (%d certificates, the 41 in-house reports excluded), held against "
        "CoQ_Analysis_Master_v48: whether it is cited as a Certificate of Quality's source "
        "(Document column, CoQ Compilation (long)), listed only under Also on file, mentioned "
        "only in another sheet, or absent from that workbook entirely. %d cited, %d also on file, %d "
        "mentioned elsewhere only, %d absent. Checked twice: an independent second pass "
        "against the first found and corrected two misclassifications (a citation whose cell "
        "carried a trailing Macedonian annotation, and a Reference-sheet mention applied to "
        "four sibling certificates but missed on a fifth) before either number was written here."
    ) % (len(rows), c.get("cited", 0), c.get("also_on_file", 0),
         c.get("other_sheet_only", 0), c.get("absent", 0))

    # highest existing OI number, so the new one continues the series
    max_n = 0
    for row in master["Reference"].iter_rows():
        v = row[0].value
        if isinstance(v, str) and v.startswith("OI-"):
            try:
                max_n = max(max_n, int(v.split("-")[1]))
            except (ValueError, IndexError):
                pass
    n = max_n + 1

    absent = [r for r in rows if r["status"] == "absent"]
    absent_list = "; ".join(sorted(
        "%s %s (%s, %s)" % (r["lab"], r["doc_code"], r["cu_batch"] or r["p_batch"], r["date"])
        for r in absent))
    what_found = (
        "OI-42's count (44 unrecorded, of a 490-scan pre-rename corpus) refreshed against the "
        "20.09.2026 renamed eCoA_DATABASE (520 files, 479 external): %d of those 479 do not "
        "appear anywhere in CoQ_Analysis_Master_v48, checked certificate by certificate against "
        "all twelve of its sheets, not the Document column alone. %d are cited as a CoQ's source, %d are on file "
        "but not credited, %d are named only in passing. An earlier pass over the same corpus "
        "counted 20 absent; two of those twenty are demonstrably present on closer reading "
        "(one cited five times with a trailing Macedonian annotation on the cell, one named in "
        "Reference's own OI-42 text for a sibling certificate but not this one) and are not "
        "repeated here."
    ) % (len(absent), c.get("cited", 0), c.get("also_on_file", 0), c.get("other_sheet_only", 0))
    what_desk_did = (
        "Every one of the %d identified by laboratory doc code and cross-checked against the "
        "Document column, the Also on file column, and every other sheet, rather than assumed "
        "absent from one column's silence. Full per-certificate result on the eCoA Coverage "
        "Audit sheet, one row each, filterable on Status."
    ) % len(absent)
    what_needed = (
        "For each of the %d: either it was replaced by a later retest and belongs in Result "
        "Supersession, or it is a result nobody has reviewed and needs a word from the Head of "
        "QC before the batches it belongs to can be called evaluated. The largest single group "
        "(7 of %d) is Institute of Public Health microbiology of 24.06.2026 — microbiology is "
        "where the master's other findings also concentrate."
    ) % (len(absent), len(absent))

    vh_text = (
        "The eCoA Coverage Audit: the 479 external certificates in the 20.09.2026 renamed "
        "eCoA_DATABASE held against CoQ_Analysis_Master_v48, one row each — %d cited as a CoQ's "
        "source, %d on file but not credited, %d mentioned elsewhere only, %d in no record of "
        "the desk (OI-%d). Carried in a companion workbook, not as a thirteenth sheet of the "
        "master: the master's Read Me and Open Items blocks are fixed-bounds ranges written by "
        "fold_reference_sheet, both full, and growing them from outside either shifts every "
        "section beneath (38 of verify_workbook.py's deeper checks broke that way) or leaves the "
        "new sheet undescribed. The three entries here are drafted for build_tracker_v8.py to "
        "place at the next rebuild."
    ) % (c.get("cited", 0), c.get("also_on_file", 0), c.get("other_sheet_only", 0),
         len(absent), n)

    # Column A carries the key and column B the text — the same shape the master's
    # own Reference blocks use, so these three entries can be lifted across as they
    # stand when build_tracker_v8.py folds this finding in.
    ws.append([None])
    ws.append([None])
    ws.append(["ECOA COVERAGE AUDIT — %s" % dt.date.today().strftime("%d.%m.%Y")])
    ws.append([None, "A companion to CoQ_Analysis_Master_v48.xlsx, which this audit reads and "
                     "does not modify. The three entries below are drafted for the master's "
                     "SHEETS, VERSION HISTORY and OPEN ITEMS blocks and are written in those "
                     "blocks' own shape, ready to be placed by build_tracker_v8.py at the next "
                     "rebuild. They are not in the master today: those blocks are fixed-bounds "
                     "ranges (Read Me rows 2-83, Open Items 846-905, from the _fold_* defined "
                     "names) with no free row, and the master's own CONVENTIONS section warns "
                     "that inserting one moves every code beneath it — which is exactly what "
                     "broke 38 of verify_workbook.py's deeper checks on the first attempt."])
    ws.append([None])
    ws.append(["eCoA Coverage Audit", sheets_text])          # for the master's SHEETS block
    ws.append([None])
    # No version number: the master stays at v48 and this ships beside it, so
    # naming a v49 here would claim a workbook that does not exist. The line
    # belongs to whichever version first carries the sheet.
    ws.append(["(next master version)", vh_text])            # for the master's VERSION HISTORY block
    ws.append([None])
    ws.append(["Ref", "Area", "State", "Item", "What was found",
               "What the desk did", "What is needed", "Evidence"])
    ws.append(["OI-%d" % n, "Record integrity", "open",
               "%d of 479 external eCoA certificates are on Drive and in no record of the "
               "desk (OI-42, refreshed)" % len(absent),
               what_found, what_desk_did, what_needed,
               "eCoA Coverage Audit sheet; tracker/ecoa_coverage_status_2026-09-20.json; "
               + absent_list[:400]])

    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 120
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws["A3"].font = Font(bold=True, size=12)


def main():
    rows = json.load(open(DATA, encoding="utf-8"))
    if len(rows) != 479:
        raise SystemExit("expected 479 certificates, found %d" % len(rows))
    bad = [r for r in rows if r["status"] not in STATUS_LABEL]
    if bad:
        raise SystemExit("unrecognised status on %d rows: %r" % (len(bad), bad[:3]))

    # the master is opened read-only, only to read the highest OI number off it,
    # so the new Open Item continues that series rather than restarting it
    master = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    c = counts(rows)
    print("counts:", c, "sum:", sum(c.values()))

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    build_sheet(wb, rows)
    read_me(wb, rows, c, master)
    wb.save(OUT)
    print("wrote", OUT, os.path.getsize(OUT), "bytes")
    print("CoQ_Analysis_Master_v48.xlsx: not modified")


if __name__ == "__main__":
    main()
