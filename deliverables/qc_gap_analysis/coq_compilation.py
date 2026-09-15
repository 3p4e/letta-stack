#!/usr/bin/env python3
"""The certificate-of-quality compilation — the owner's first request (31.08.2026):

    "for all certificates of quality, for all batches, for all initial and retest
    certificates of quality, a table containing all information that is needed for the
    certificate of quality template, but most importantly, from parameter 1 to 12, the
    document codes and date of issuing of the certificate of analysis from external
    laboratories — and the analysis results for each of those parameters — for every
    certificate of quality document code individually, for every single parameter."

Two tables from one source, `coq_artifact_data.json` (the export the certificates are
compiled from), with the CoQ code the CoQ Register states:

* **wide** — one row per certificate of quality: the template's header fields (code,
  date of issue, the certificate it supersedes, batch, P lot, strain, harvest and
  packaging, the internal certificate it cites, the Total THC result and its
  certificate, grade, potency window, product and specification codes, the
  specification's bands), then for every determination #1 … #12 — sub-determinations
  included, 23 in all — four columns: the result the certificate prints, the document
  it rests on, that document's date of issue and its laboratory.
* **long** — one row per certificate of quality and determination, the same four
  fields with the method, the acceptance criterion, the laboratory's receipt date of
  the sample, the desk's status for the row, the route where nothing is on file yet,
  and the other documents on file that also carry the result.

A determination the certificate does not print as a result (not tested at release,
upon request, to be performed in house) shows that status in the result column and no
document; a reissue's row carried from the initial testing shows the initial
certificate's document and says so in the status.

    python3 coq_compilation.py [--src coq_artifact_data.json] [--out tracker/CoQ_compilation_v33]

writes <out>.xlsx (sheets "CoQ Compilation" and "CoQ Compilation (long)"), <out>_wide.csv
and <out>_long.csv. The workbook build calls build()/fill_wide()/fill_long() with its own
CoQ Register codes, so the tabs and the files carry one code per certificate.
"""
import collections
import csv
import json
import os
import re
import sys

G = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(G, "tracker"))
sys.path.insert(0, G)
import tracker_data as T          # noqa: E402
import receipt_dates as RD        # noqa: E402

DETS = ["1", "2", "3", "4", "5", "6", "7", "8", "9.1", "9.2", "9.3", "9.4", "9.5", "9.6", "9.7",
        "10.1", "10.2", "10.3", "11.1", "11.2", "11.3", "11.4", "12"]
HEAD = ["CoQ code", "Series", "Register status", "Date of issue", "Supersedes", "Plan no.",
        "Batch (cultivation)", "P lot", "Strain", "Harvest", "Packaging",
        "Manuf. date (template)", "Pack. date (template)",
        "iCoA code", "iCoA issued", "iCoA tested / sampled", "Retest campaign",
        "Total THC (%)", "THC certificate", "Grade", "Potency window",
        "Product code", "Specification code", "Specification status",
        "Phenotype", "Chemotype", "Processing", "Dominance", "Packaging (spec.)",
        "Latest external certificate"]
LONG = ["CoQ code", "Series", "Batch (cultivation)", "P lot", "Strain", "Date of issue", "#",
        "Parameter", "Method", "Acceptance criterion", "Result", "Document", "Issued",
        "Laboratory", "Received by the laboratory", "Status", "Route", "Also on file"]
NAVY, RED, AMBER, GREY = "1F3864", "F4B6B6", "FDE9D9", "EDEDED"


def series(t):
    if t.startswith("additional"):
        return "reissue — 12-month retest"
    return "release"


def status_word(st):
    """The result column where the certificate prints no result: what the row says."""
    st = st or ""
    if st.startswith("upon request"):
        return "upon request — not required for release"
    if st.startswith("not tested"):
        return "not tested"
    if st.startswith("to be performed"):
        return "to be performed (in house)"
    if st.startswith("awaiting the cannabinoid"):
        return "awaiting the Farmahem cannabinoid certificate"
    if st.startswith("awaiting the mycotoxin"):
        return "awaiting the Farmahem mycotoxin certificate"
    if st.startswith("outside the retest scope"):
        return "outside the retest scope"
    if st.startswith("in-house CoA only"):
        return "in-house record only"
    return st[:60] if st else "—"


def build(src=None, codes=None, batch_dates=None):
    """(wide rows, long rows). `codes` — {(batch key, 'I'|'R'): CoQ code} — overrides the
    export's code, as the workbook build passes its own CoQ Register."""
    src = src or os.path.join(G, "coq_artifact_data.json")
    d = json.load(open(src, encoding="utf-8"))
    dets = {x["no"]: x for x in d["dets"]}
    bl = {}
    bd = batch_dates or os.path.join(G, "tracker", "batch_dates.csv")
    if os.path.exists(bd):
        for r in csv.DictReader(open(bd, encoding="utf-8")):
            for k in (r.get("cu_batch"), r.get("p_batch")):
                if k:
                    bl.setdefault(T.batch_key(k), r)

    def span(a, b):
        a, b = (a or "").strip(), (b or "").strip()
        return a if not b or b == a else f"{a} – {b}"

    wide, long_ = [], []
    for c in d["coqs"]:
        code = c.get("regcode") if str(c.get("regcode", "")).startswith("CoQ-PP_26-") else ""
        sfx = "R" if c["t"].startswith("additional") else "I"
        if codes:
            for name in (c.get("pp"), c.get("cb")):
                if name and (T.batch_key(name), sfx) in codes:
                    code = codes[(T.batch_key(name), sfx)]
                    break
        code = code or "— at issue —"
        rows = {str(r["no"]): r for r in c["rows"]}
        r4 = rows.get("4") or {}
        sup = c.get("supersedes") or {}
        spc = c.get("spc") or {}
        b = bl.get(T.batch_key(c.get("pp") or "")) or bl.get(T.batch_key(c.get("cb") or "")) or {}
        issue = c.get("issue") or ""
        issue = issue if re.match(r"^\d{2}\.\d{2}\.\d{4}$", issue) else (issue and f"{issue} (floor — not yet dated)")
        head = {
            "CoQ code": code, "Series": series(c["t"]),
            "Register status": {"yes": "issuable", "allocated": "allocated — code reserved, date provisional"}.get(c.get("reg_issuable"), "not yet issuable"),
            "Date of issue": issue or "—",
            "Supersedes": (f"{sup['code']} of {sup.get('date', '')}".strip() if sup.get("code") else "n/a — release certificate" if sfx == "I" else "— initial certificate not yet numbered —"),
            "Plan no.": c.get("n") or "—",
            "Batch (cultivation)": c.get("cb") or "—", "P lot": c.get("pp") or "— no P lot assigned —",
            "Strain": c.get("strain") or "—",
            "Harvest": span(b.get("harvest_from"), b.get("harvest_to")) or b.get("harvest_printed") or "—",
            "Packaging": span(b.get("packaging_from"), b.get("packaging_to")) or b.get("packaging_printed") or "—",
            "Manuf. date (template)": c.get("md") or "—", "Pack. date (template)": c.get("pk") or "—",
            "iCoA code": c.get("icoa_code") or c.get("ic") or "—", "iCoA issued": c.get("icoa_issue") or "—",
            "iCoA tested / sampled": c.get("icoa_tested") or "—",
            "Retest campaign": {"197": "Tranche 1", "220": "Tranche 2", "227": "Tranche 3"}.get(c.get("icoa_campaign"), c.get("icoa_campaign") or ("—" if sfx == "I" else "no campaign on file")),
            "Total THC (%)": (r4.get("res") if r4.get("res") not in (None, "", "—") else "—"),
            "THC certificate": (r4.get("doc") if r4.get("doc") not in (None, "", "—") else "—"),
            "Grade": c.get("grade") or "—",
            "Potency window": (r4.get("crit") or "").split("(")[0].strip() or "—",
            "Product code": c.get("pcode") or "—", "Specification code": c.get("spec") or "—",
            "Specification status": c.get("spec_status") or "—",
            "Phenotype": spc.get("pheno") or "—", "Chemotype": spc.get("chemo") or "—",
            "Processing": spc.get("proc") or "—", "Dominance": spc.get("dominance") or "—",
            "Packaging (spec.)": spc.get("pack") or "—",
            "Latest external certificate": c.get("last_external") or "—",
        }
        for no in DETS:
            r = rows.get(no) or {}
            res = (r.get("res") or "").strip()
            doc = (r.get("doc") or "").strip()
            has_doc = bool(doc) and doc != "—"
            result = res if res and res != "—" else status_word(r.get("st"))
            head[f"#{no} result"] = result
            head[f"#{no} document"] = doc if has_doc else "—"
            head[f"#{no} issued"] = (r.get("dd") or "—") if has_doc else "—"
            head[f"#{no} laboratory"] = (r.get("lab") or "—") if has_doc else "—"
            dd = dets.get(no, {})
            long_.append({
                "CoQ code": code, "Series": head["Series"], "Batch (cultivation)": head["Batch (cultivation)"],
                "P lot": head["P lot"], "Strain": head["Strain"], "Date of issue": head["Date of issue"],
                "#": no, "Parameter": dd.get("en") or "—",
                "Method": r.get("mth") or dd.get("method") or "—",
                "Acceptance criterion": r.get("crit") or dd.get("crit") or "—",
                "Result": result, "Document": head[f"#{no} document"], "Issued": head[f"#{no} issued"],
                "Laboratory": head[f"#{no} laboratory"],
                "Received by the laboratory": (RD.received(doc) or "—") if has_doc and not doc.startswith("iCoA") else "—",
                "Status": r.get("st") or "—", "Route": r.get("route") or "", "Also on file": r.get("also") or "",
            })
        wide.append(head)

    def sk(r):
        m = re.match(r"CoQ-PP_26-(\d+)", r["CoQ code"])
        return (0, int(m.group(1)), "") if m else (1, 0, r["Batch (cultivation)"] + r["P lot"])
    wide.sort(key=sk)
    order = {r["CoQ code"] + r["Batch (cultivation)"] + r["P lot"] + r["Series"]: i for i, r in enumerate(wide)}
    long_.sort(key=lambda r: (order.get(r["CoQ code"] + r["Batch (cultivation)"] + r["P lot"] + r["Series"], 9999), DETS.index(r["#"])))
    return wide, long_


WIDE_COLS = HEAD + [f"#{no} {k}" for no in DETS for k in ("result", "document", "issued", "laboratory")]


def _style(ws, widths, freeze):
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF", size=9)
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[1].height = 30
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(size=9)
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = freeze
    ws.auto_filter.ref = ws.dimensions


def fill_wide(ws, rows):
    from openpyxl.styles import PatternFill
    ws.append(WIDE_COLS)
    for r in rows:
        ws.append([r.get(c, "") for c in WIDE_COLS])
    widths = [15, 22, 22, 12, 24, 16, 16, 12, 18, 18, 18, 12, 12, 15, 11, 11, 12, 10, 14, 6, 18, 18, 22, 40, 12, 10, 16, 14, 40, 12]
    widths += [14, 16, 11, 24] * len(DETS)
    _style(ws, widths, "D2")
    grey, amber = PatternFill("solid", fgColor=GREY), PatternFill("solid", fgColor=AMBER)
    first = len(HEAD) + 1
    for row in ws.iter_rows(min_row=2, min_col=first):
        for c in row:
            if (c.column - first) % 4 == 0:            # a result column
                v = str(c.value or "")
                if v.startswith(("not tested", "upon request", "to be performed", "awaiting", "outside", "in-house record", "—")):
                    c.fill = grey if v.startswith(("upon request", "—")) else amber
    return len(rows)


def fill_long(ws, rows):
    from openpyxl.styles import PatternFill
    ws.append(LONG)
    for r in rows:
        ws.append([r.get(c, "") for c in LONG])
    _style(ws, [15, 22, 16, 12, 18, 12, 6, 30, 34, 34, 22, 18, 11, 30, 11, 40, 30, 30], "H2")
    grey, amber = PatternFill("solid", fgColor=GREY), PatternFill("solid", fgColor=AMBER)
    for row in ws.iter_rows(min_row=2):
        v = str(row[10].value or "")
        if v.startswith(("not tested", "upon request", "to be performed", "awaiting", "outside", "in-house record", "—")):
            row[10].fill = grey if v.startswith(("upon request", "—")) else amber
    return len(rows)


def write_files(wide, long_, out):
    import openpyxl
    with open(out + "_wide.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=WIDE_COLS)
        w.writeheader()
        w.writerows(wide)
    with open(out + "_long.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=LONG)
        w.writeheader()
        w.writerows(long_)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "CoQ Compilation"
    fill_wide(ws, wide)
    fill_long(wb.create_sheet("CoQ Compilation (long)"), long_)
    wb.save(out + ".xlsx")


def main():
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument("--src", default=os.path.join(G, "coq_artifact_data.json"))
    ap.add_argument("--out", default=os.path.join(G, "tracker", "CoQ_compilation"))
    a = ap.parse_args()
    wide, long_ = build(a.src)
    write_files(wide, long_, a.out)
    n_doc = sum(1 for r in long_ if r["Document"] != "—")
    print("%s.xlsx / _wide.csv / _long.csv: %d certificates × %d determinations = %d rows, %d with a document"
          % (a.out, len(wide), len(DETS), len(long_), n_doc))
    print(collections.Counter(r["Series"] for r in wide))
    print(collections.Counter(r["Register status"] for r in wide))


if __name__ == "__main__":
    main()
