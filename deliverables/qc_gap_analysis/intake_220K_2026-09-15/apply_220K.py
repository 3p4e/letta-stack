#!/usr/bin/env python3
"""Write the Tranche 2 Farmahem cannabinoid certificates (220-1-К/26 … 220-32-К/26) into the
owner's release register, in the shape of the Tranche 1 rows (197-n-К/26).

    python3 deliverables/qc_gap_analysis/intake_220K_2026-09-15/apply_220K.py
        [--register PATH]   the register to extend  (default: …SUBLOT_2026-09-01.xlsx)
        [--out PATH]        where to write          (default: in place)

Idempotent: a certificate whose code is already in the register is skipped. The desk
does not write the owner's register on its own; this script is the write, ready to run.

The thirty-two were in the owner's eCoA database (Drive folder eCoA_DATABASE, files
250826_220-n-K-26_… and 260826_220-n-K-26_…) from 09.09.2026 and were never read into
the desk — which is why every Tranche 2 reissue stood "retest assay pending" on the v29
CoQ Register (owner, 15.09.2026: "here are all retest results that I have on file and
correct yourself"). Received by the laboratory 17.08.2026, analysed 24/25.08.2026, issued
25.08.2026 (220-1 … -16) and 26.08.2026 (220-17 … -32); Total CBD and Total CBN "< LOQ"
on every certificate, Total Δ9-THC as printed with its expanded uncertainty.

Two reads must agree before a certificate is written: reads_220K.json (the page read of
15.09.2026, from rendered page crops — no text layer, no classical OCR) and
readB2_220K.json (an independent second transcription of the same pages by a separate
reader session that saw no other transcription; the runner's Gemini reader was refused on
15.09.2026 — every key 503 or 403 — so the second read is not Gemini's), compared on the
batch printed, the date of issue, the date of analysis, the laboratory sample number, the
three results and the expanded uncertainty printed beside each. One disagreement stops the write.

Row by row: the block by the batch as printed — the P lot for 28, the cultivation batch
for GG1024, JD042601, FB042601 and CC042601 — else the cultivation batch the Head of QC's
batch list (tracker/batch_dates.csv) gives for that P lot, which is how the seven lots
whose label row carries no P number (FB012601/1, GRC102501-2, JD012603-02, JD012603-02V,
KC102501, PM112501, SCR112501) are found; every one of the 32 has a block after the
intakes of 14.09 and 15.09, so this script opens none and stops if one is missing. The
row goes after the block's last real certificate row: THC % as printed, "≥ 5.00 %" in
THC spec as the 197 rows carry it, CBD % and CBN % "<LOQ" as the 197 rows write it, "/"
in every column the certificate does not report, the code with the Cyrillic К every
197-n-К/26 row uses, the date of issue, "Farmahem", "Open".
"""
import argparse, copy, csv, json, os, re, sys
import openpyxl
from openpyxl.utils import range_boundaries

HERE = os.path.dirname(os.path.abspath(__file__))
GAP = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(GAP)), 'ingestion', 'common'))
from batch_id import batch_key  # noqa: E402

PLACEHOLDER = {'', '(not numbered)', 'n/a'}
THC_SPEC = '≥ 5.00 %'
RESULTS = ('Total THC', 'Total CBD', 'Total CBN')


def reg_value(v):
    """A printed result in the register's own spelling: '< LOQ' is written '<LOQ'."""
    return str(v).replace('< LOQ', '<LOQ').replace('<  LOQ', '<LOQ').strip()


def reg_code(code):
    """The code in the register's own spelling: Cyrillic К, as every 197-n-К/26 row."""
    return str(code).strip().replace('-K/', '-К/')


def _val(x):
    """A result from either read's shape — a string, or {'result': …}."""
    if isinstance(x, dict):
        x = x.get('result', x.get('value', ''))
    return reg_value(x).replace('%w/w', '').replace('%', '').strip()


def _find(d, *needles):
    """The value of the first key of d containing every needle (case-insensitive)."""
    for k, v in (d or {}).items():
        kk = k.lower()
        if all(n.lower() in kk for n in needles):
            return v
    return None


def disagreements(m, b):
    """Where the page read and the second read differ; empty when they agree; None without a second read."""
    if not b:
        return None
    out = []
    lbl = b.get('client_sample_label') or {}
    bb = str((lbl.get('batch') if isinstance(lbl, dict) else '') or b.get('batch') or _find(b, 'batch') or '').strip()
    if batch_key(bb) != batch_key(m['batch_printed']):
        out.append('batch %r vs %r' % (bb, m['batch_printed']))
    for key, mine, label in (('issue_date', 'date_of_issue', 'issue date'), ('analysis_date', 'date_analysed', 'analysis date'),
                             ('received_date', 'date_received', 'receipt date'), ('lab_sample_number', 'lab_sample_number', 'sample number')):
        theirs = str(b.get(key) or _find(b, key.split('_')[0]) or '').strip()
        if theirs.replace(' ', '') != str(m[mine]).replace(' ', ''):
            out.append('%s %r vs %r' % (label, theirs, m[mine]))
    res = b.get('results') or b
    for name, needles in (('Total THC', ('thc',)), ('Total CBD', ('cbd',)), ('Total CBN', ('cbn',))):
        theirs = _find(res, *needles)
        if theirs is None:
            out.append('%s not in the second read' % name); continue
        if _val(theirs) != _val(m['results_flat'][name]):
            out.append('%s %r vs %r' % (name, _val(theirs), m['results_flat'][name]))
        u_theirs = str(theirs.get('U', '')).strip() if isinstance(theirs, dict) else ''
        if u_theirs and u_theirs != str(m['uncertainty'][name]).strip():
            out.append('%s uncertainty %r vs %r' % (name, u_theirs, m['uncertainty'][name]))
    return out


def load_reads():
    R = json.load(open(os.path.join(HERE, 'reads_220K.json'), encoding='utf-8'))
    pb = os.path.join(HERE, 'readB2_220K.json')
    B = json.load(open(pb, encoding='utf-8')) if os.path.exists(pb) else {}
    bad, single = {}, []
    for code, m in R.items():
        b = B.get(code) or B.get(code.replace('-K-26', '-K/26')) or B.get(m['cert_code'])
        d = disagreements(m, b)
        if d is None:
            single.append(m['cert_code'])
        elif d:
            bad[code] = d
    if bad:
        for code, d in bad.items():
            print('DISAGREE', code, '—', '; '.join(d))
        sys.exit('%d certificate(s) where the two reads differ — nothing written' % len(bad))
    if single:
        sys.exit('%d certificate(s) with no second read — nothing written: %s' % (len(single), ', '.join(single)))
    return R, single


def cu_from_list(p):
    """The cultivation batch the Head of QC's batch list gives for a P lot, or ''."""
    with open(os.path.join(GAP, 'tracker', 'batch_dates.csv'), encoding='utf-8-sig') as fh:
        for r in csv.DictReader(fh):
            if (r.get('p_batch') or '').strip() == p:
                return (r.get('cu_batch') or '').strip()
    return ''


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--register', default=os.path.join(GAP, 'PP_Batch_Release_QC_Register_SUBLOT_2026-09-01.xlsx'))
    ap.add_argument('--out', default=None)
    a = ap.parse_args()
    out = a.out or a.register
    R, _ = load_reads()
    wb = openpyxl.load_workbook(a.register)
    ws = wb['Batch Release QC']
    present = {reg_code(ws.cell(r, 23).value or '') for r in range(6, ws.max_row + 1)}
    real = lambda r: str(ws.cell(r, 23).value or '').strip() not in PLACEHOLDER

    labels = []
    for r in range(6, ws.max_row + 1):
        v = ws.cell(r, 2).value
        if v:
            labels.append((r, str(v).strip(), str(ws.cell(r, 3).value or '').strip()))
        elif str(ws.cell(r, 1).value or '').upper() == 'LEGEND':
            labels.append((r, None, None)); break
    blocks = {}
    for i, (r, b, p) in enumerate(labels):
        if b:
            blk = {'first': r, 'last': labels[i + 1][0] - 1, 'label': b}
            blocks[batch_key(b)] = blk
            if p.startswith('P'):
                blocks.setdefault(batch_key(p), blk)
    tmpl_row = next(r for r in range(6, ws.max_row + 1) if str(ws.cell(r, 23).value or '') == '197-17-К/26')
    assert str(ws.cell(tmpl_row, 6).value).strip() == THC_SPEC, ws.cell(tmpl_row, 6).value
    last_data = max(r for r in range(6, ws.max_row + 1) if ws.cell(r, 23).value not in (None, ''))

    def values(m):
        v = [None] * 26
        r = m['results_flat']
        v[4], v[5], v[6], v[7] = reg_value(r['Total THC']), THC_SPEC, reg_value(r['Total CBD']), reg_value(r['Total CBN'])
        for c in range(9, 23): v[c - 1] = '/'
        v[22], v[23], v[24], v[25] = reg_code(m['cert_code']), m['date_of_issue'], 'Farmahem', 'Open'
        return v

    def style_from(src, dst):
        for c in range(1, 27):
            s, d = ws.cell(src, c), ws.cell(dst, c)
            d.font, d.fill, d.border, d.alignment, d.number_format = copy.copy(s.font), copy.copy(s.fill), copy.copy(s.border), copy.copy(s.alignment), s.number_format
        ws.row_dimensions[dst].height = ws.row_dimensions[src].height

    todo = sorted(R.values(), key=lambda m: m['n'])
    todo = [m for m in todo if reg_code(m['cert_code']) not in present]
    into, missing, via_list = [], [], []
    for m in todo:
        cands = [m['batch_printed']]
        if m['p_number']:
            cu = cu_from_list(m['p_number'])
            if cu:
                cands.append(cu)
        blk = next((blocks[batch_key(c)] for c in cands if c and batch_key(c) in blocks), None)
        if blk is None:
            missing.append(m['cert_code']); continue
        if batch_key(m['batch_printed']) not in blocks:
            via_list.append('%s = %s' % (m['p_number'], blk['label']))
        into.append((m, blk))
    if missing:
        sys.exit('no register block for %s — nothing written (the intakes of 14.09 and 15.09 open every block these need)' % ', '.join(missing))
    inserts = sorted(((max(r for r in range(blk['first'], blk['last'] + 1) if real(r)) + 1, m) for m, blk in into), key=lambda t: -t[0])
    merged_below = [str(mr) for mr in ws.merged_cells.ranges if mr.min_row > last_data]
    for mr in merged_below: ws.unmerge_cells(mr)
    for at, m in inserts:
        ws.insert_rows(at); style_from(tmpl_row + (1 if tmpl_row >= at else 0), at)
        for c, v in enumerate(values(m), 1): ws.cell(at, c).value = v
    last_data += len(inserts)
    for mr in merged_below:
        c1, r1, c2, r2 = range_boundaries(mr); ws.merge_cells(start_row=r1 + len(inserts), start_column=c1, end_row=r2 + len(inserts), end_column=c2)
    ws.auto_filter.ref = 'A4:Z%d' % last_data
    wb.save(out)
    print('%s: %d row(s) into existing blocks, 0 new blocks, %d already present, legend shifted by %d'
          % (os.path.basename(out), len(inserts), 32 - len(todo), len(inserts)))
    if via_list:
        print('block found through the batch list (label row carries no P lot): ' + ', '.join(via_list))


if __name__ == '__main__':
    main()
