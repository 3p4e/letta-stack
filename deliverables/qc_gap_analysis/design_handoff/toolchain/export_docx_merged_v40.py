#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Each merged tranche document of certificates of quality, as one Word file.

    python3 design_handoff/toolchain/export_docx_merged_v40.py
    python3 design_handoff/toolchain/export_docx_merged_v40.py --dpi 200   # smaller files

The Head of QC, 18.09.2026, asked for Tranche 1 and Tranche 2 "PDF and Word" on both
fleets. The internal certificates already had it (`icoa_handoff/v3/export_docx_icoa.py`);
the certificates of quality had only per-certificate Word files inside the distribution
archives, which is not the same thing as the one document the tranche PDF is.

Same rule as everywhere else on this desk: the page is carried as it prints. Every page of
the merged tranche PDF is rendered at 300 dpi and placed edge to edge on an A4 sheet with
zero margins, one certificate per page, in the order the merge put them. The text is not
editable in Word — the editable source is the HTML in `out/`, the record is the PDF — but
the document cannot drift from the PDF, because it is made from it.

Output: design_handoff/pdf/CoQ_Tranche_<n>.docx, beside the PDF it is made from.
"""
import argparse
import glob
import io
import os
import re
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
PDFDIR = os.path.join(os.path.dirname(HERE), "pdf")


def build(dpi, pattern):
    import pymupdf
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Mm

    made = 0
    for src in sorted(glob.glob(os.path.join(PDFDIR, pattern))):
        stem = os.path.splitext(os.path.basename(src))[0]
        book = pymupdf.open(src)
        doc = Document()
        s = doc.sections[0]
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Mm(0)
        s.header_distance = s.footer_distance = Mm(0)
        for i, page in enumerate(book):
            pix = page.get_pixmap(dpi=dpi, colorspace=pymupdf.csRGB, alpha=False)
            para = doc.add_paragraph()
            para.paragraph_format.space_before = para.paragraph_format.space_after = Mm(0)
            run = para.add_run()
            if i:
                run.add_break(WD_BREAK.PAGE)
            run.add_picture(io.BytesIO(pix.tobytes("png")), width=Mm(210), height=Mm(297))
        n = book.page_count
        book.close()
        core = doc.core_properties
        core.title = stem.replace("_", " ")
        core.subject = "Certificate of Quality — Purely Plant GmbH"
        core.comments = ("%d certificates, each page rendered from the controlled vector "
                         "PDF at %d dpi; the PDF is the record." % (n, dpi))
        out = os.path.join(PDFDIR, stem + ".docx")
        doc.save(out)
        made += 1
        print("  %-28s %3d certificate(s)  %.1f MiB"
              % (stem + ".docx", n, os.path.getsize(out) / 1048576))
    return made


def main(argv):
    ap = argparse.ArgumentParser()
    ap.add_argument("--dpi", type=int, default=300)
    ap.add_argument("--pattern", default="CoQ_Tranche_[0-9].pdf",
                    help="which merged documents to convert (default: the whole-tranche ones)")
    a = ap.parse_args(argv[1:])
    made = build(a.dpi, a.pattern)
    print("Word documents written: %d" % made)
    return 0 if made else 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
