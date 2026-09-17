#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Every certificate of quality as a Word document.

    python3 design_handoff/toolchain/export_docx_v40.py            # all 172
    python3 design_handoff/toolchain/export_docx_v40.py --dpi 300

The Head of QC, 17.09.2026: "include every certificate of quality as a Word document."

A certificate's page is an intricate fixed layout — 56 stacked style layers, faded bands,
embedded signatures — that no HTML-to-Word conversion reproduces faithfully, and a Word
rebuild of the template would be a second template to keep in step with the first. So each
.docx carries the certificate as the page **exactly as it prints**: the vector PDF page
rendered at 300 dpi and placed edge to edge on an A4 page with zero margins. It opens in
Word, previews, prints and forwards like any Word file, and it cannot drift from the PDF
because it is made from it.

What it is not: the text inside is not editable in Word. The editable source of every
certificate is its HTML in `out/`, which is also in the package; the controlled record is
the PDF. Nothing here is signed by the desk — the signature scans travel as on the PDF.

Output: design_handoff/docx/<same folders as out/>/<same stem>.docx
"""
import argparse, glob, io, os, sys
HERE = os.path.dirname(os.path.abspath(__file__)); HANDOFF = os.path.dirname(HERE)
PAGES = os.path.join(HANDOFF, "pdf", "pages"); OUT = os.path.join(HANDOFF, "out"); DOCX = os.path.join(HANDOFF, "docx")


def folder_of(stem):
    for html in glob.glob(os.path.join(OUT, "**", stem + ".html"), recursive=True):
        return os.path.relpath(os.path.dirname(html), OUT)
    return "."


def build(dpi):
    import pymupdf
    from docx import Document
    from docx.shared import Mm
    n = 0
    for pdf in sorted(glob.glob(os.path.join(PAGES, "*.pdf"))):
        stem = os.path.splitext(os.path.basename(pdf))[0]
        d = pymupdf.open(pdf); pg = d[0]
        pix = pg.get_pixmap(dpi=dpi, colorspace=pymupdf.csRGB, alpha=False)
        png = pix.tobytes("png"); d.close()
        doc = Document()
        s = doc.sections[0]
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Mm(0)
        s.header_distance = s.footer_distance = Mm(0)
        para = doc.add_paragraph(); para.paragraph_format.space_before = para.paragraph_format.space_after = 0
        para.add_run().add_picture(io.BytesIO(png), width=Mm(210), height=Mm(297))
        core = doc.core_properties
        core.title = stem; core.subject = "Certificate of Quality — Purely Plant GmbH"
        core.comments = "Page rendered from the controlled vector PDF at %d dpi; the PDF is the record." % dpi
        dest = os.path.join(DOCX, folder_of(stem)); os.makedirs(dest, exist_ok=True)
        doc.save(os.path.join(dest, stem + ".docx")); n += 1
    print("docx written: %d  -> %s" % (n, os.path.relpath(DOCX)))


def main(argv):
    ap = argparse.ArgumentParser(); ap.add_argument("--dpi", type=int, default=300)
    a = ap.parse_args(argv[1:]); build(a.dpi); return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
