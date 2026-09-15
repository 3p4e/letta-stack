#!/usr/bin/env python3
"""Truth check of 15.09.2026 — every parameter result and every cited eCoA code, date
and laboratory, from the primary records forward to what the desk prints.

An independent code path: nothing here imports the schedule, the export, the tracker
builder or the verifier. It reads the same primary records they read — the release
register with the intakes applied, the intake transcriptions (both reads where two
exist), the two-read corpus of the eCoA database — and compares, with its own key and
value normalisation, what the derived layers state:

  T1  intake transcriptions → the release register rows written from them
  T2  intake transcriptions → the tracker instances written from them
  T3  export (the certificates' rows) → a primary record: code exists, date, lab, result
  T4  tracker sheet cells (recalculated workbook) → a primary record: code, date, lab, result
  T5  export rows ↔ tracker cells: the certificate cites what the tracker cites, same value
  T6  compiled draft certificates → export: header code/date/supersedes, every result,
      every cited (code, date)
  T7  references table → export: per certificate and determination, code, date, lab
  T8  CoQ Register → export/primary: Total THC and its certificate, the latest eCoA cited

A finding is a sentence with the evidence beside it. Counts are printed per check so a
zero is a zero over N comparisons, never over none.

    python3 tracker/truth_check_2026-09-15.py --root <qc_gap_analysis with intakes applied> \
        --workbook <recalculated CoQ_Analysis_Master_v31.xlsx> --out tracker/TRUTH_CHECK_2026-09-15.md
"""
import argparse
import csv
import glob
import html
import json
import os
import re
import sys
from collections import Counter, defaultdict

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))

# ------------------------------------------------------------------ normalisation (own)
CYR = str.maketrans("АВЕКМНОРСТУХЈЅІавекмнорстухјѕі", "ABEKMHOPCTYXJSIabekmhopctyxjsi")
SUP = str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹", "0123456789")


def nkey(code):
    """A certificate code as a comparison key: Latin letters for the Cyrillic homoglyphs,
    one separator, no trailing gloss. 220-1-К/26 == 220-1-K-26 == 220-1-K/26."""
    c = str(code or "").strip()
    c = re.sub(r"\s*\([^)]*\)\s*$", "", c)
    c = re.sub(r"\s+(EN|MK)$", "", c)
    c = re.sub(r"[-/](GS|ГС|LOD|LoD)[-/]", "/LOD/", c)
    c = c.translate(CYR)
    return re.sub(r"[\s\-/_.]+", "/", c).strip("/").upper()


def dkey(d):
    """dd.mm.yyyy, dd.mm.yy or a datetime → yyyy-mm-dd; '' when not a date."""
    if hasattr(d, "strftime"):
        return d.strftime("%Y-%m-%d")
    m = re.match(r"^\s*(\d{1,2})\.(\d{1,2})\.(\d{2}|\d{4})\s*$", str(d or ""))
    if not m:
        return ""
    y = int(m.group(3))
    y = y + 2000 if y < 100 else y
    return "%04d-%02d-%02d" % (y, int(m.group(2)), int(m.group(1)))


def labkey(lab):
    """Any spelling of a laboratory to one of CNP, IJZ, FHM, PP, DFL, NGP, ''."""
    s = str(lab or "").strip().lower()
    if not s or s in ("n/a", "none", "—", "-"):
        return ""
    if "purely" in s or s in ("pp", "pp*", "ngp") or "in-house" in s:
        return "PP"
    if "фармахем" in s or "farmahem" in s or s.startswith("fhm"):
        return "FHM"
    if "природни" in s or "natural products" in s or "cnp" in s or "faculty of pharmacy" in s or "фармацевтски" in s:
        return "CNP"
    if "јавно здравје" in s or "public health" in s or s.startswith("iph") or s.startswith("ijz") or "ијз" in s:
        return "IJZ"
    if "фитосанитар" in s or "phytosanitary" in s or s == "dfl":
        return "DFL"
    if "ngp" in s or "нгп" in s:
        return "NGP"
    return s


def norm(v):
    """A result as a comparable value.

    ('n', x)        a number, units stripped, exponents resolved
    ('lt', x)       '< x' / '≤ x'
    ('gt', x)       '> x'
    ('nd',)         the not-detected family
    ('loq',)        below the limit of quantification (BLQ, < LOQ, < LOD)
    ('absent',)     absence (Salmonella, E. coli)
    ('conforms',)   a conformity verdict
    ('empty',)      no result — including the page's bracketed blank and 'not tested'
    ('t', text)     anything else, folded: superscripts to ^n, 'и' to 'and' and dropped,
                    the tracker's ᴿ/ᴰ markers and trailing glosses removed
    """
    raw = str(v if v is not None else "")
    s = raw.strip()
    s = s.split("|")[0].strip() if "|" in s else s          # 'Absent | Отсутна'
    s = re.sub(r"\s+—\s.*$", "", s)                           # '2.06 — DETECTED, >LOQ'
    s = re.sub(r"\s*±.*$", "", s)                             # '10,3 ± 0,2'
    glossed_absent = bool(re.search(r"\(.*absent.*\)", s, re.I))
    s = re.sub(r"\s*\([^)]*\)\s*$", "", s)                  # '(ImB spec.)', '(below limit of quantification)', '(absent)'
    s = re.sub(r"[ᴿᴰ*]+\s*$", "", s).strip()                  # 'ND ᴿ', '0.02 ᴰ', '<LOQ**'
    s = re.sub(r"([⁰¹²³⁴⁵⁶⁷⁸⁹]+)", lambda m: "^" + m.group(1).translate(SUP), s)
    low = s.lower().replace("\u00a0", " ").strip()
    if low in ("", "—", "-", "[—]", "n/a", "/", "none", "n.r.", "not reported", "not tested", "•", "✓"):
        return ("empty",)
    if re.match(r"^(nd|n\.d\.?|н\.д\.?|нд|not detected|не е детектиран[оа]?|не детектиран[оа]?)(?=$|[\s,;(])", low):
        return ("nd",)
    if re.match(r"^<\s*loq|^blq|^below (the )?loq|^< ?lq\b|^<\s*lod", low):
        return ("loq",)
    if glossed_absent or re.match(r"^(absent|отсут|отсуств|not found|negative|негатив)", low):
        return ("absent",)
    if re.match(r"^(conforms|одговара|odgovara|complies|соодветствува|passes)(?=$|[\s,;(])", low):
        return ("conforms",)
    if re.match(r"^не одговара|^does not conform|^fails", low):
        return ("t", "nonconforming")
    t = low
    t = re.sub(r"(cfu/g|[µμ]g/kg|ug/kg|mg/kg|%\s*w/w|% ?w/w|%|w/w)", "", t).strip()
    t = t.replace(",", ".").replace(" и ", " and ")
    t = re.sub(r"\s*[×x·]\s*10\s*\^?\s*(-?\d+)", lambda m: "e" + m.group(1), t)
    t = re.sub(r"\s*10\s*\^\s*(-?\d+)", lambda m: "1e" + m.group(1), t)
    t = t.replace(" and ", " ")
    m = re.match(r"^([<≤>≥])\s*([0-9.]+(?:e-?\d+)?)$", t)
    if m:
        try:
            return ("lt" if m.group(1) in "<≤" else "gt", float(m.group(2)))
        except ValueError:
            pass
    m = re.match(r"^([0-9]+(?:\.[0-9]+)?(?:e-?\d+)?)$", t)
    if m:
        try:
            return ("n", float(m.group(1)))
        except ValueError:
            pass
    return ("t", re.sub(r"\s+", "", t))


def normd(v, det):
    """norm(), knowing the determination: in the absence columns a verdict is an absence
    (the register writes 'Одговара (absent)' where the page prints the verdict)."""
    x = norm(v)
    if det in ("9.4", "9.5") and x[0] == "conforms":
        return ("absent",)
    return x


def same(a, b):
    """Two normalised results agree: same category, numbers within 0.005."""
    if a[0] != b[0]:
        return False
    if a[0] in ("n", "lt", "gt"):
        return abs(a[1] - b[1]) <= 0.005 or (b[1] and abs(a[1] - b[1]) / abs(b[1]) <= 1e-6)
    if a[0] == "t":
        return a[1] == b[1]
    return True


def show(x):
    return {"n": lambda: "%g" % x[1], "lt": lambda: "< %g" % x[1], "gt": lambda: "> %g" % x[1],
            "t": lambda: repr(x[1])}.get(x[0], lambda: x[0])()


# register column → determination, from the register's own header row (checked against
# the export's map at run time, so the two definitions are compared, not assumed)
HDR2DET = [("THC %", "4"), ("CBD", "5"), ("CBN", "6"), ("Loss on drying", "8"), ("TAMC", "9.1"),
           ("TYMC", "9.2"), ("Bile-tolerant", "9.3"), ("Salmonella", "9.4"), ("E. coli", "9.5"),
           ("Aflatoxins Σ", "10.2"), ("Aflatoxin B1", "10.1"), ("Ochratoxin", "10.3"),
           ("Pb ", "11.1"), ("Cd ", "11.2"), ("As ", "11.3"), ("Hg ", "11.4"), ("Pesticides", "12")]
GROUPS = {"9": ["9.1", "9.2", "9.3", "9.4", "9.5"], "10": ["10.1", "10.2", "10.3"],
          "11": ["11.1", "11.2", "11.3", "11.4"]}
CORPUS_KEY = {"1": "identification_a_macroscopic", "2": "identification_b_microscopic", "4": "total_thc",
              "5": "total_cbd", "6": "total_cbn", "8": "loss_on_drying", "9.1": "tamc", "9.2": "tymc",
              "9.3": "bile_tolerant_gram_negative", "9.4": "salmonella", "9.5": "escherichia_coli",
              "10.1": "aflatoxin_b1", "10.2": "aflatoxins_total", "10.3": "ochratoxin_a",
              "11.1": "lead", "11.2": "cadmium", "11.3": "arsenic", "11.4": "mercury"}

FIND = []
COUNT = Counter()


def bad(check, what, detail):
    FIND.append((check, what, detail))


# ------------------------------------------------------------------ primary sources
def read_register(path):
    """Every certificate row of the release register: code, date, lab, block batch and
    P lot, and the value in every determination column, keyed by the header text."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["Batch Release QC"]
    hdr = {c: str(ws.cell(4, c).value or "").strip() for c in range(1, 26)}
    col2det = {}
    for c, h in hdr.items():
        for needle, det in HDR2DET:
            if h.startswith(needle.strip()) or needle in h:
                col2det[c] = det
                break
    rows, cur = [], None
    for r in range(6, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if str(a or "").strip().upper() == "LEGEND":
            break
        b = str(ws.cell(r, 2).value or "").strip()
        if b:
            cur = {"batch": b, "pn": str(ws.cell(r, 3).value or "").strip(), "strain": str(ws.cell(r, 4).value or "").strip()}
        code = str(ws.cell(r, 23).value or "").strip()
        if not code or cur is None:
            continue
        vals = {}
        for c, det in col2det.items():
            v = ws.cell(r, c).value
            if v not in (None, "") and str(v).strip() not in ("/", "—", "-"):
                vals[det] = str(v).strip()
        rows.append({"row": r, "code": code, "key": nkey(code), "date": dkey(ws.cell(r, 24).value),
                     "date_raw": str(ws.cell(r, 24).value or "").strip()[:10], "lab": labkey(ws.cell(r, 25).value),
                     "lab_raw": str(ws.cell(r, 25).value or "").strip(), "vals": vals, **cur})
    return rows, col2det, hdr


def read_intakes(root):
    """The intake transcriptions as one list of records {code, key, date, received, lab,
    batch, p, vals{det: value}, second{...}} — the second read attached where one exists."""
    out = []
    # Tranche 2 mycotoxins: 220-n-М/26 — read A (Claude), read B (Gemini)
    p = os.path.join(root, "intake_220M_2026-09-14", "reads_claude.json")
    if os.path.exists(p):
        a = json.load(open(p, encoding="utf-8"))
        a = a if isinstance(a, list) else list(a.values())
        bp = os.path.join(root, "intake_220M_2026-09-14", "reads_gemini.json")
        b = json.load(open(bp, encoding="utf-8")) if os.path.exists(bp) else {}
        b = b if isinstance(b, list) else list(b.values())
        bmap = {}
        for x in b:
            px = x.get("parsed") or x
            bmap[nkey(px.get("cert_code"))] = px
        for x in a:
            res = x.get("results") or {}
            afl = [res.get(k) for k in ("Aflatoxin B1", "Aflatoxin B2", "Aflatoxin G1", "Aflatoxin G2")]
            vals = {"10.1": res.get("Aflatoxin B1"), "10.3": res.get("Ochratoxin A")}
            if afl and all(norm(v)[0] == "nd" for v in afl):
                vals["10.2"] = "ND"
            rec = {"series": "220-M", "code": x.get("cert_code"), "key": nkey(x.get("cert_code")),
                   "date": dkey(x.get("date_of_issue")), "received": dkey(x.get("date_received")),
                   "lab": "FHM", "batch": x.get("batch_printed") or "", "p": x.get("p_number") or "",
                   "vals": {k: v for k, v in vals.items() if v is not None}, "second": None}
            sb = bmap.get(rec["key"])
            if sb:
                sres = {}
                for pm in sb.get("parameters") or []:
                    sres[pm.get("parameter_printed") or pm.get("parameter")] = pm.get("result_printed")
                if not sres and isinstance(sb.get("results"), dict):
                    sres = sb["results"]
                rec["second"] = {"date": dkey(sb.get("date_of_issue")), "p": sb.get("p_number") or "",
                                 "label": sb.get("batch_printed") or "", "res": sres}
            out.append(rec)
    # Tranche 3 potency: 227-n-К/26 — read A (page read), read B (checkpoint table, 25)
    p = os.path.join(root, "intake_227K_2026-09-15", "reads_227K.json")
    if os.path.exists(p):
        a = json.load(open(p, encoding="utf-8"))
        a = a if isinstance(a, list) else list(a.values())
        bp = os.path.join(root, "intake_227K_2026-09-15", "checkpoint_master_coa_table.json")
        b = json.load(open(bp, encoding="utf-8")) if os.path.exists(bp) else {}
        b = b if isinstance(b, list) else list(b.values())
        bmap = {}
        for x in b:
            for k in (x.get("p"), x.get("cu")):
                if k:
                    bmap[str(k).replace("/", "-")] = x
        for x in a:
            res = x.get("results") or {}
            rec = {"series": "227-K", "code": x.get("cert_code"), "key": nkey(x.get("cert_code")),
                   "date": dkey(x.get("date_of_issue")), "received": dkey(x.get("date_received")),
                   "lab": "FHM", "batch": x.get("batch_printed") or "", "p": x.get("p_number") or "",
                   "vals": {"4": res.get("Total THC"), "5": res.get("Total CBD"), "6": res.get("Total CBN")},
                   "second": None}
            sb = bmap.get(str(rec["p"]).replace("/", "-")) or bmap.get(str(rec["batch"]).replace("/", "-"))
            if sb:
                rec["second"] = {"date": dkey(str(sb.get("issue", "")).split(" ")[0]), "p": sb.get("p") or "",
                                 "res": {"4": sb.get("Total THC"), "5": sb.get("Total CBD"), "6": sb.get("Total CBN")}}
            out.append(rec)
    # Tranche 2 potency: 220-n-К/26 — read A (page read), read B (independent transcription)
    p = os.path.join(root, "intake_220K_2026-09-15", "reads_220K.json")
    if os.path.exists(p):
        a = json.load(open(p, encoding="utf-8"))
        a = a if isinstance(a, list) else list(a.values())
        bp = os.path.join(root, "intake_220K_2026-09-15", "readB2_220K.json")
        b = json.load(open(bp, encoding="utf-8")) if os.path.exists(bp) else {}
        b = b if isinstance(b, list) else list(b.values())
        bmap = {nkey(x.get("report_number") or x.get("cert_code")): x for x in b}
        for x in a:
            res = x.get("results") or {}

            def rv(k):
                v = res.get(k)
                return v.get("result") if isinstance(v, dict) else v
            s = x.get("sample") or {}
            rec = {"series": "220-K", "code": x.get("report_number") or x.get("cert_code"),
                   "key": nkey(x.get("report_number") or x.get("cert_code")),
                   "date": dkey(x.get("issued") or x.get("date_of_issue")),
                   "received": dkey(x.get("received") or x.get("date_received")),
                   "lab": "FHM", "batch": x.get("batch_printed") or s.get("batch") or "", "p": x.get("p_number") or "",
                   "vals": {"4": rv("Total Δ9-THC") or rv("Total THC"), "5": rv("Total CBD"), "6": rv("Total CBN")},
                   "second": None}
            sb = bmap.get(rec["key"])
            if sb:
                sres = sb.get("results") or {}

                def sv(k):
                    v = sres.get(k)
                    return v.get("result") if isinstance(v, dict) else v
                lbl = sb.get("client_sample_label") or sb.get("sample") or {}
                rec["second"] = {"date": dkey(sb.get("issued") or sb.get("date_of_issue")),
                                 "p": lbl.get("batch") or sb.get("batch_printed") or "",
                                 "res": {"4": sv("Total Δ9-THC") or sv("Total THC"), "5": sv("Total CBD"), "6": sv("Total CBN")}}
            out.append(rec)
    return out


def read_corpus(path):
    by = {}
    if not os.path.exists(path):
        return by
    for r in json.load(open(path, encoding="utf-8")):
        k = nkey(r.get("cert_code"))
        if not k:
            # a document the reads could not code is still on file: its code and date
            # are in the file name, <batch>_<code>, <date>_<lab>.pdf
            m = re.match(r"^[^_]+_(.+?), (\d{2}\.\d{2}\.\d{4})_", str(r.get("document") or ""))
            if not m:
                continue
            k = nkey(m.group(1))
            r = dict(r, date_of_issue=m.group(2), _from_filename=True)
        params = {}
        for pm in r.get("parameters") or []:
            if pm.get("parameter") and pm.get("result_printed") not in (None, ""):
                params.setdefault(pm["parameter"], pm["result_printed"])
        by.setdefault(k, []).append({"date": dkey(r.get("date_of_issue")), "lab": labkey(r.get("lab")),
                                     "params": params, "batch": r.get("batch_printed") or ""})
    return by


def read_resolution(path):
    """The owner's 09.09.2026 pass over eCoA_DATABASE — cell_resolution_2026-09-09.tsv:
    per batch and determination the document cited, its date and laboratory and what it
    prints. A single page read; the certificates rest on it where the register is silent."""
    by = {}
    if not os.path.exists(path):
        return by
    for r in csv.DictReader(open(path, encoding="utf-8"), delimiter="\t"):
        code = (r.get("Document to cite") or "").strip()
        if not code or code.startswith(("—", "iCoA", "NO-DOC")):
            continue
        det = (r.get("Determination") or "").split(" ")[0]
        e = by.setdefault(nkey(code), {"date": dkey(r.get("Issued")), "lab": labkey(r.get("Laboratory")), "vals": {}})
        v = (r.get("What the document prints") or "").strip()
        if v and det:
            e["vals"].setdefault(det, v)
    return by


def read_instances(path):
    by = {}
    for i in json.load(open(path, encoding="utf-8")):
        by.setdefault(nkey(i["code"]), []).append({"date": dkey(i.get("date")), "lab": labkey(i.get("lab")),
                                                   "vals": i.get("vals") or {}, "p": i.get("p") or "", "cu": i.get("cu") or ""})
    return by


class Primary:
    """The union of the primary records, looked up by certificate key."""

    def __init__(self, reg_rows, corpus, instances, intakes, index, resolution):
        self.reg = defaultdict(list)
        for r in reg_rows:
            self.reg[r["key"]].append(r)
        self.corpus, self.inst = corpus, instances
        self.intake = {r["key"]: r for r in intakes}
        self.index = index
        self.res = resolution
        self.alias = {}

    def has0(self, key):
        return key in self.reg or key in self.corpus or key in self.inst or key in self.intake or key in self.index or key in self.res

    def resolve(self, key):
        """The key as the record holds it: itself; its last component for a composite
        in-house string ('PP CoA #027 / ППК25370' → the CNP certificate it rests on); or
        the register's long name that contains it ('NGP-QCG-SOP-024 F3' inside
        'In-house GC cross-check NGP/QCG/SOP-024')."""
        if key in self.alias:
            return self.alias[key]
        out = key
        if not self.has0(key) and "/" in key and re.match(r"^(PP|COA)", key):
            parts = key.split("/")
            for n in range(1, len(parts)):
                k2 = "/".join(parts[n:])
                if self.has0(k2):
                    out = k2
                    COUNT["documents cited through a composite in-house string"] += 1
                    break
        if out == key and not self.has0(key):
            short = key.replace("/F3", "")
            hit = next((rk for rk in self.reg if short and short in rk and rk != key), None)
            if hit:
                out = hit
                COUNT["documents cited by a short alias of the register's name"] += 1
        self.alias[key] = out
        return out

    def has(self, key):
        return self.has0(self.resolve(key))

    def dates(self, key):
        key = self.resolve(key)
        d = set()
        if key in self.res:
            d.add(self.res[key]["date"])
        for r in self.reg.get(key, []):
            d.add(r["date"])
        for r in self.corpus.get(key, []):
            d.add(r["date"])
        for r in self.inst.get(key, []):
            d.add(r["date"])
        if key in self.intake:
            d.add(self.intake[key]["date"])
        if key in self.index:
            d.add(self.index[key]["date"])
        return {x for x in d if x}

    def labs(self, key):
        key = self.resolve(key)
        d = set()
        if key in self.res:
            d.add(self.res[key]["lab"])
        for r in self.reg.get(key, []):
            d.add(r["lab"])
        for r in self.corpus.get(key, []):
            d.add(r["lab"])
        for r in self.inst.get(key, []):
            d.add(r["lab"])
        if key in self.intake:
            d.add(self.intake[key]["lab"])
        if key in self.index:
            d.add(self.index[key]["lab"])
        return {x for x in d if x}

    def values(self, key, det):
        """Every primary value on record for (certificate, determination): [(source, value)]."""
        key = self.resolve(key)
        out = []
        for r in self.reg.get(key, []):
            if det in r["vals"]:
                out.append(("register row %d" % r["row"], r["vals"][det]))
        if key in self.res and det in self.res[key]["vals"] and ";" not in self.res[key]["vals"][det]:
            # the pass records a single value per cell; a multi-value cell (free acid;
            # acid; total) is the page's whole line and is not compared
            out.append(("09.09 pass (one read)", self.res[key]["vals"][det]))
        for r in self.inst.get(key, []):
            if det in r["vals"]:
                out.append(("instance", r["vals"][det]))
        if key in self.intake and det in self.intake[key]["vals"]:
            out.append(("intake read", self.intake[key]["vals"][det]))
        ck = CORPUS_KEY.get(det)
        for r in self.corpus.get(key, []):
            if ck and ck in r["params"]:
                out.append(("corpus", r["params"][ck]))
        return out


# ------------------------------------------------------------------ the checks
def t1_t2(intakes, reg_rows, instances, cu_of_p):
    """Intake transcriptions against the register rows and the tracker instances."""
    reg_by = defaultdict(list)
    for r in reg_rows:
        reg_by[r["key"]].append(r)
    for rec in intakes:
        k = rec["key"]
        # the second read, where one exists, still agrees with the first
        if rec["second"]:
            COUNT["T0 two-read pairs"] += 1
            s = rec["second"]
            if s["date"] and s["date"] != rec["date"]:
                bad("T0", "the two reads disagree on the date of issue", "%s: A %s, B %s" % (rec["code"], rec["date"], s["date"]))
            sp = str(s["p"]).replace("/", "-")
            if sp and rec["p"] and re.match(r"^P\d{6}$", sp) and sp != str(rec["p"]).replace("/", "-"):
                bad("T0", "the two reads disagree on the lot", "%s: A %s, B %s" % (rec["code"], rec["p"], s["p"]))
            elif sp and rec["p"] and not re.match(r"^P\d{6}$", sp):
                COUNT["T0 second read's lot field holds no lot (label checked instead)"] += 1
                if str(rec["p"]) not in str(s.get("label") or ""):
                    bad("T0", "the second read's label does not name the first read's lot", "%s: A %s, B label %r" % (rec["code"], rec["p"], s.get("label")))
            for det, v in rec["vals"].items():
                sv = None
                for kk, vv in (s["res"] or {}).items():
                    kl = str(kk).lower()
                    if (det == "4" and "thc" in kl) or (det == "5" and "cbd" in kl) or (det == "6" and "cbn" in kl) \
                            or (det == "10.1" and "b1" in kl) or (det == "10.3" and "ochra" in kl):
                        sv = vv
                if sv is not None:
                    COUNT["T0 two-read results"] += 1
                    if not same(norm(v), norm(sv)):
                        bad("T0", "the two reads disagree on a result (#%s)" % det, "%s: A %r, B %r" % (rec["code"], v, sv))
        # T1 — the register row
        rows = reg_by.get(k, [])
        COUNT["T1 intake certificates"] += 1
        if not rows:
            bad("T1", "intake certificate has no register row", "%s (%s)" % (rec["code"], rec["p"] or rec["batch"]))
        else:
            if len(rows) > 1:
                bad("T1", "intake certificate on more than one register row", "%s: rows %s" % (rec["code"], [r["row"] for r in rows]))
            for r in rows:
                if r["date"] != rec["date"]:
                    bad("T1", "register date differs from the certificate", "%s: register %s, read %s" % (rec["code"], r["date_raw"], rec["date"]))
                if r["lab"] != rec["lab"]:
                    bad("T1", "register laboratory differs from the certificate", "%s: register %r" % (rec["code"], r["lab_raw"]))
                want = (rec["p"] or rec["batch"]).replace("/", "-")
                have = (r["pn"] or r["batch"]).replace("/", "-")
                listed = cu_of_p.get(rec["p"], "").replace("/", "-")
                if want and have and want != have and want not in (r["batch"].replace("/", "-"), r["pn"]) \
                        and listed != r["batch"].replace("/", "-"):
                    bad("T1", "register block is not the lot the certificate prints",
                        "%s: block %s / %s, certificate %s / %s" % (rec["code"], r["batch"], r["pn"], rec["batch"], rec["p"]))
                for det, v in rec["vals"].items():
                    COUNT["T1 results"] += 1
                    if det not in r["vals"]:
                        bad("T1", "register row lacks the result the certificate prints (#%s)" % det, "%s: read %r" % (rec["code"], v))
                    elif not same(norm(r["vals"][det]), norm(v)):
                        bad("T1", "register result differs from the certificate (#%s)" % det,
                            "%s: register %r, read %r" % (rec["code"], r["vals"][det], v))
        # T2 — the instance
        ins = instances.get(k, [])
        if not ins:
            bad("T2", "intake certificate has no tracker instance", rec["code"])
        for i in ins:
            COUNT["T2 instances"] += 1
            if i["date"] != rec["date"]:
                bad("T2", "instance date differs from the certificate", "%s: instance %s, read %s" % (rec["code"], i["date"], rec["date"]))
            if i["lab"] != rec["lab"]:
                bad("T2", "instance laboratory differs from the certificate", "%s: %r" % (rec["code"], i["lab"]))
            lot = (i["p"] or i["cu"]).replace("/", "-")
            want = (rec["p"] or rec["batch"]).replace("/", "-")
            if lot and want and lot != want and rec["batch"].replace("/", "-") not in (i["cu"].replace("/", "-"),):
                bad("T2", "instance lot is not the lot the certificate prints", "%s: instance %s/%s, certificate %s/%s" % (rec["code"], i["cu"], i["p"], rec["batch"], rec["p"]))
            for det, v in rec["vals"].items():
                if det in i["vals"]:
                    COUNT["T2 results"] += 1
                    if not same(norm(i["vals"][det]), norm(v)):
                        bad("T2", "instance result differs from the certificate (#%s)" % det, "%s: instance %r, read %r" % (rec["code"], i["vals"][det], v))
                elif det in ("4", "5", "6", "10.1", "10.3"):
                    bad("T2", "instance lacks a result the certificate prints (#%s)" % det, rec["code"])


def t3_export(export, prim, col2det_reg):
    """Every row of every certificate against the primary record of the document it cites."""
    # the register's own column map against the export's
    exp_map = {d["no"]: d["col"] for d in export["dets"] if d.get("col")}
    mine = {}
    for c, det in col2det_reg.items():
        mine[det] = openpyxl.utils.get_column_letter(c)
    for det, col in exp_map.items():
        COUNT["T3 column map entries"] += 1
        if mine.get(det) != col:
            bad("T3", "the export's register column for a determination is not the header's", "#%s: export %s, header %s" % (det, col, mine.get(det)))
    for c in export["coqs"]:
        lot = c.get("pp") or c.get("cb")
        for r in c["rows"]:
            doc = (r.get("doc") or "").strip()
            if not doc or doc.startswith(("—", "iCoA", "NO-DOC")):
                continue
            k = nkey(doc)
            det = str(r.get("no"))
            COUNT["T3 cited rows"] += 1
            if not prim.has(k):
                bad("T3", "certificate cites a document no primary record holds", "%s #%s: %s" % (lot, det, doc))
                continue
            dd = dkey(r.get("dd"))
            if dd and prim.dates(k) and dd not in prim.dates(k):
                bad("T3", "certificate dates a document differently from the record", "%s #%s: %s printed %s, record %s" % (lot, det, doc, r.get("dd"), sorted(prim.dates(k))))
            lb = labkey(r.get("lab"))
            if lb and prim.labs(k) and lb not in prim.labs(k):
                bad("T3", "certificate names a laboratory the record does not", "%s #%s: %s printed %s, record %s" % (lot, det, doc, r.get("lab"), sorted(prim.labs(k))))
            res = r.get("res")
            nv = normd(res, det)
            if nv[0] in ("empty", "conforms") or det in ("1", "2", "3", "7", "9.6", "9.7", "12"):
                continue
            vals = prim.values(k, det)
            if not vals:
                COUNT["T3 results without a primary value"] += 1
                continue
            COUNT["T3 results compared"] += 1
            agree = [s for s, v in vals if same(normd(v, det), nv)]
            if not agree:
                bad("T3", "certificate result differs from the record (#%s)" % det,
                    "%s: %s prints %r; record %s" % (lot, doc, res, ["%s %r" % (s, v) for s, v in vals]))
            elif len(agree) < len(vals):
                dis = ["%s %r" % (s, v) for s, v in vals if not same(normd(v, det), nv)]
                bad("T3", "the primary records disagree among themselves (#%s)" % det, "%s: %s prints %r; agrees with %s; not with %s" % (lot, doc, res, agree, dis))


def read_tracker(wb_values):
    """The tracker sheet's cells: per lot, per determination, the cited reference and the
    printed value, block by block (a round every two rows)."""
    name = next(n for n in wb_values.sheetnames if n.startswith("CoQ Parameter Tracker"))
    ws = wb_values[name]
    starts = [c for c in range(4, ws.max_column + 1) if str(ws.cell(2, c).value or "").startswith("#")]
    pcol = {}
    for i, s0 in enumerate(starts):
        e = (starts[i + 1] - 1) if i + 1 < len(starts) else ws.max_column
        n = re.match(r"#(\d+)", str(ws.cell(2, s0).value)).group(1)
        pcol[n] = (s0, e)
    anchors = [r for r in range(5, ws.max_row + 1) if ws.cell(r, 1).value not in (None, "")
               and not str(ws.cell(r, 1).value).startswith("KEY")]
    lots = []
    for a, nxt in zip(anchors, anchors[1:] + [ws.max_row + 1]):
        cu, p = str(ws.cell(a, 1).value), str(ws.cell(a, 2).value or "")
        cells = []                      # (det, code, date, lab, value, row, credited)
        for n, (s0, e) in pcol.items():
            single = (e - s0) <= 2
            subs = GROUPS.get(n, [n])
            for r in range(a, nxt, 2):
                ref = str(ws.cell(r if single else r + 1, s0 + 1 if single else s0).value or "")
                m = re.match(r"^(.+?), \((\d{2}\.\d{2}\.\d{4})\) \[([^\]]+)\]", ref)
                if not m:
                    continue
                credited = "not credited" not in ref
                for j, sub in enumerate(subs):
                    v = ws.cell(r, s0 + (0 if single else j)).value
                    cells.append((sub, m.group(1).strip(), m.group(2), m.group(3), str(v or ""), r, credited))
        lots.append({"cu": cu, "p": p, "row": a, "cells": cells})
    return lots


def t4_tracker(lots, prim):
    for lot in lots:
        for det, code, date, lab, val, row, credited in lot["cells"]:
            if code.startswith(("iCoA", "NO-DOC")) or labkey(lab) == "PP":
                continue
            k = nkey(code)
            COUNT["T4 tracker citations"] += 1
            if not prim.has(k):
                bad("T4", "tracker cites a document no primary record holds", "%s/%s #%s row %d: %s" % (lot["cu"], lot["p"], det, row, code))
                continue
            if prim.dates(k) and dkey(date) not in prim.dates(k):
                bad("T4", "tracker dates a document differently from the record", "%s/%s #%s: %s (%s), record %s" % (lot["cu"], lot["p"], det, code, date, sorted(prim.dates(k))))
            if prim.labs(k) and labkey(lab) not in prim.labs(k):
                bad("T4", "tracker names a laboratory the record does not", "%s/%s #%s: %s [%s], record %s" % (lot["cu"], lot["p"], det, code, lab, sorted(prim.labs(k))))
            nv = normd(val, det)
            if not credited or nv[0] in ("empty", "conforms") or det in ("1", "2", "3", "7", "12"):
                continue
            vals = prim.values(k, det)
            if not vals:
                continue
            COUNT["T4 results compared"] += 1
            if not any(same(normd(v, det), nv) for s, v in vals):
                bad("T4", "tracker result differs from the record (#%s)" % det,
                    "%s/%s: %s prints %r; record %s" % (lot["cu"], lot["p"], code, val, ["%s %r" % (s, v) for s, v in vals]))


def t5_export_vs_tracker(export, lots):
    """The certificate cites what the tracker cites for the lot, with the same value."""
    by_p, by_cu = {}, {}
    for lot in lots:
        if lot["p"] and not lot["p"].startswith("N/A"):
            by_p[nkey(lot["p"])] = lot
        by_cu[nkey(re.sub(r"[＊*]", "", lot["cu"]))] = lot
    for c in export["coqs"]:
        lot = (by_p.get(nkey(c["pp"])) if c.get("pp") else None) or by_cu.get(nkey(c.get("cb", "")))
        if not lot:
            COUNT["T5 certificates with no tracker lot"] += 1
            continue
        cited = defaultdict(list)
        for det, code, date, lab, val, row, credited in lot["cells"]:
            cited[(det, nkey(code))].append((val, date, credited))

        def tail(k):
            return k.split("/")[-1] if re.match(r"^(PP|COA)", k) and "/" in k else k
        for r in c["rows"]:
            doc = (r.get("doc") or "").strip()
            det = str(r.get("no"))
            if not doc or doc.startswith(("—", "iCoA", "NO-DOC")) or det in ("1", "2", "3", "7", "9.6", "9.7", "12"):
                continue
            COUNT["T5 rows"] += 1
            hits = cited.get((det, nkey(doc))) or cited.get((det, tail(nkey(doc))))
            if not hits:
                bad("T5", "certificate cites a document the tracker does not cite for the lot (#%s)" % det,
                    "%s: %s; tracker cites %s" % (c.get("pp") or c.get("cb"), doc, sorted({k[1] for k in cited if k[0] == det})))
                continue
            nv = normd(r.get("res"), det)
            if nv[0] in ("empty", "conforms"):
                continue
            if not any(same(normd(v, det), nv) for v, d, cr in hits):
                bad("T5", "certificate result differs from the tracker's for the same document (#%s)" % det,
                    "%s: %s certificate %r, tracker %r" % (c.get("pp") or c.get("cb"), doc, r.get("res"), [v for v, d, cr in hits]))


def t6_drafts(root, export):
    """Every compiled draft against the export record it was compiled from."""
    by_lot = defaultdict(dict)
    for c in export["coqs"]:
        series = "reissue" if c["t"].startswith("additional") else "initial"
        for nm in (c.get("pp"), c.get("cb")):
            if nm:
                by_lot[nkey(nm)].setdefault(series, c)
    files = sorted(glob.glob(os.path.join(root, "drafts", "DRAFT_CoQ_*.html")))
    for f in files:
        base = os.path.basename(f)[len("DRAFT_CoQ_"):-len(".html")]
        series = "reissue" if base.endswith("_reissue") else "initial"
        lot = base[:-len("_reissue")] if series == "reissue" else base
        c = by_lot.get(nkey(lot), {}).get(series)
        COUNT["T6 drafts"] += 1
        if not c:
            bad("T6", "draft has no export record", base)
            continue
        h = open(f, encoding="utf-8").read()
        m = re.search(r'class="hb-code">([^<]*)<', h)
        code = html.unescape(m.group(1)).strip() if m else ""
        want = c.get("regcode") or ""
        if want.startswith("CoQ-PP_26-") and code != want:
            bad("T6", "draft prints a code other than the register's", "%s: draft %r, register %r" % (base, code, want))
        m = re.search(r"Issued · Издаден <b>([^<]*)</b>", h)
        if m and dkey(m.group(1)) != dkey(c.get("issue")):
            bad("T6", "draft prints a date other than the export's", "%s: draft %r, export %r" % (base, m.group(1), c.get("issue")))
        m = re.search(r"\(supersedes ([^ ]+) of ([0-9.]+)\)", h)
        sup = c.get("supersedes") or {}
        if series == "reissue" and sup.get("code", "").startswith("CoQ-PP_26-"):
            if not m or m.group(1) != sup["code"] or dkey(m.group(2)) != dkey(sup.get("date")):
                bad("T6", "draft's supersedes line differs from the export", "%s: draft %r, export %r" % (base, m.groups() if m else None, sup))
        # the results table
        body = h[h.find('class="results"'):]
        body = body[:body.find("</table>")]
        printed = {}
        grp = None
        SUBLBL = [("tamc", "9.1"), ("tymc", "9.2"), ("bile", "9.3"), ("salmonella", "9.4"), ("escherichia", "9.5"),
                  ("aflatoxin b", "10.1"), ("aflatoxins", "10.2"), ("ochratoxin", "10.3"),
                  ("lead", "11.1"), ("cadmium", "11.2"), ("arsenic", "11.3"), ("mercury", "11.4")]
        for tr in re.finditer(r"<tr([^>]*)>(.*?)</tr>", body, re.S):
            attrs, inner = tr.group(1), tr.group(2)
            if "row-group" in attrs:
                mg = re.match(r"\s*<td>(\d+)</td>", inner)
                grp = mg.group(1) if mg else None
                continue
            mv = re.search(r'class="r-val[^"]*">(.*?)</span>\s*</td>', inner, re.S)
            if not mv:
                continue
            val = mv.group(1)
            val = val.split('<i class="bisep">')[0]
            val = html.unescape(re.sub(r"<[^>]+>", "", val)).strip()
            if "sub-row" in attrs and grp:
                ml = re.search(r'class="p-sub">(.*?)<', inner, re.S)
                lbl = html.unescape(ml.group(1)).lower() if ml else ""
                det = next((d for needle, d in SUBLBL if needle in lbl), None)
            else:
                md = re.match(r"\s*<td>(\d+(?:\.\d+)?)</td>", inner)
                det = md.group(1) if md else None
            if det:
                printed[det] = val
        for r in c["rows"]:
            det = str(r.get("no"))
            if det not in printed:
                continue
            res = r.get("res") or ""
            COUNT["T6 results compared"] += 1
            a, b = normd(printed[det], det), normd(res, det)
            if a[0] == "empty" and b[0] == "empty":
                continue
            if not same(a, b):
                # a blank line on the page for a result the export marks as not printed
                if a[0] == "empty" and str(r.get("st", "")).startswith(("to be performed", "upon request", "not required")):
                    continue
                bad("T6", "draft prints a result other than the export's (#%s)" % det, "%s: draft %r, export %r" % (base, printed[det], res))
        # the cited documents
        refs = set()

        def base(k):
            return k.split("/")[-1] if re.match(r"^(PP|COA)", k) and "/" in k else k
        for m2 in re.finditer(r'class="lr-mono">(.*?)</td>', h, re.S):
            txt = html.unescape(re.sub(r"<[^>]+>", "", m2.group(1)))
            for part in txt.split(" · "):
                cm = re.match(r"^\s*(.+?), (\d{2}\.\d{2}\.\d{4})\s*$", part)
                if cm:
                    refs.add((base(nkey(cm.group(1))), dkey(cm.group(2))))
        want_refs = set()
        for r in c["rows"]:
            doc = (r.get("doc") or "").strip()
            if doc and not doc.startswith(("—", "NO-DOC")) and dkey(r.get("dd")):
                want_refs.add((base(nkey(doc)), dkey(r.get("dd"))))
        COUNT["T6 cited documents"] += len(want_refs)
        if refs:
            for k, d in sorted(want_refs - refs):
                if any(k2 == k for k2, d2 in refs):
                    bad("T6", "draft dates a cited document differently from the export", "%s: %s export %s, draft %s" % (base, k, d, [d2 for k2, d2 in refs if k2 == k]))
                else:
                    bad("T6", "draft does not list a document the export cites", "%s: %s (%s)" % (base, k, d))
            for k, d in sorted(refs - want_refs):
                if not any(k2 == k for k2, d2 in want_refs):
                    bad("T6", "draft lists a document the export does not cite", "%s: %s (%s)" % (base, k, d))


def t7_references(path, export):
    if not os.path.exists(path):
        bad("T7", "references table not found", path)
        return
    by_code = {c["regcode"]: c for c in export["coqs"] if str(c.get("regcode", "")).startswith("CoQ-PP_26-")}
    for row in csv.DictReader(open(path, encoding="utf-8")):
        c = by_code.get(row.get("CoQ"))
        if not c:
            continue
        COUNT["T7 certificates"] += 1
        rows = {str(r["no"]): r for r in c["rows"]}
        for g in range(1, 13):
            cell = (row.get("#%d" % g) or "").strip()
            first = cell.split(";")[0].strip()
            m = re.match(r"^(\S+) (\d{2}\.\d{2}\.\d{2,4}|—) (\S+)", first)
            if not m:
                continue
            code, date, lab = m.groups()
            subs = GROUPS.get(str(g), [str(g)])
            er = next((rows[s] for s in subs if s in rows and (rows[s].get("doc") or "").strip() not in ("", "—")), None)
            if not er:
                continue
            COUNT["T7 cells"] += 1
            if nkey(code) != nkey(er["doc"]):
                bad("T7", "references table cites a document other than the certificate's (#%d)" % g, "%s: table %s, certificate %s" % (row["CoQ"], code, er["doc"]))
                continue
            if date != "—" and dkey(er.get("dd")) and dkey(date) != dkey(er["dd"]):
                bad("T7", "references table dates a document differently (#%d)" % g, "%s: %s table %s, certificate %s" % (row["CoQ"], code, date, er["dd"]))
            if labkey(lab.rstrip("*")) != labkey(er.get("lab")) and not (labkey(lab.rstrip("*")) == "IJZ" and labkey(er.get("lab")) == "IJZ"):
                bad("T7", "references table names a laboratory differently (#%d)" % g, "%s: %s table %s, certificate %s" % (row["CoQ"], code, lab, er.get("lab")))


def t8_register(wb_values, export, prim):
    ws = wb_values["CoQ Register"]
    rows = list(ws.iter_rows(values_only=True))
    hi = next(i for i, r in enumerate(rows) if r and any(str(x or "").strip() == "No." for x in r))
    hdr = [str(x or "").strip() for x in rows[hi]]
    by_code = {c["regcode"]: c for c in export["coqs"] if str(c.get("regcode", "")).startswith("CoQ-PP_26-")}
    for r in rows[hi + 1:]:
        d = dict(zip(hdr, r))
        code = str(d.get("CoQ code") or "")
        if not code.startswith("CoQ-PP_26-"):
            continue
        c = by_code.get(code)
        if not c:
            bad("T8", "a numbered register row has no export record", "%s %s" % (code, d.get("Key")))
            continue
        COUNT["T8 numbered rows"] += 1
        r4 = next((x for x in c["rows"] if str(x["no"]) == "4"), None)
        thc = d.get("Total THC (%)")
        if thc not in (None, "", "—") and r4 and norm(r4.get("res"))[0] == "n" and not same(norm(thc), norm(r4["res"])):
            bad("T8", "register Total THC differs from the certificate's row 4", "%s: register %r, certificate %r" % (code, thc, r4["res"]))
        cert = str(d.get("THC certificate") or "")
        if cert not in ("", "—") and r4 and nkey(cert) != nkey(r4.get("doc")):
            bad("T8", "register THC certificate differs from the certificate's row 4", "%s: register %r, certificate %r" % (code, cert, r4.get("doc")))
        lc, ld = str(d.get("Latest eCoA cited (code)") or ""), d.get("Latest eCoA cited (date)")
        if lc and not lc.startswith("—"):
            k = nkey(lc)
            if not prim.has(k):
                bad("T8", "register's latest eCoA is on no primary record", "%s: %s" % (code, lc))
            elif dkey(ld) and prim.dates(k) and dkey(ld) not in prim.dates(k):
                bad("T8", "register dates its latest eCoA differently from the record", "%s: %s %s, record %s" % (code, lc, dkey(ld), sorted(prim.dates(k))))
        ic = str(d.get("Ident C — eCoA (Total THC)") or "")
        if ic and not ic.startswith("—"):
            k = nkey(ic.split(",")[0])
            if not prim.has(k):
                bad("T8", "register's identification C certificate is on no primary record", "%s: %s" % (code, ic[:40]))


# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", default=os.path.dirname(HERE), help="qc_gap_analysis with the intakes applied to its register")
    ap.add_argument("--workbook", required=True, help="the recalculated CoQ_Analysis_Master workbook (values)")
    ap.add_argument("--corpus", default=None)
    ap.add_argument("--out", default=None)
    args = ap.parse_args()
    root = os.path.abspath(args.root)
    corpus_path = args.corpus or os.path.join(os.path.dirname(os.path.dirname(root)), "ingestion", "ecoa_runner", "records_corpus.json")

    reg_rows, col2det, hdr = read_register(os.path.join(root, "PP_Batch_Release_QC_Register_SUBLOT_2026-09-01.xlsx"))
    intakes = read_intakes(root)
    corpus = read_corpus(corpus_path)
    instances = read_instances(os.path.join(root, "tracker", "new_instances.json"))
    export = json.load(open(os.path.join(root, "coq_artifact_data.json"), encoding="utf-8"))
    index = {nkey(e["code"]): {"date": dkey(e["date"]), "lab": labkey(e["lab"])} for e in export.get("ecoa", [])}
    resolution = read_resolution(os.path.join(root, "cell_resolution_2026-09-09.tsv"))
    prim = Primary(reg_rows, corpus, instances, intakes, index, resolution)
    cu_of_p = {}
    bl = os.path.join(root, "tracker", "batch_dates.csv")
    if os.path.exists(bl):
        for r in csv.DictReader(open(bl, encoding="utf-8")):
            if r.get("p_batch"):
                cu_of_p[r["p_batch"]] = r["cu_batch"]
    wbv = openpyxl.load_workbook(args.workbook, data_only=True)

    print("primary records: %d register rows (%d codes), %d intake certificates (%d with a second read), "
          "%d corpus codes, %d instance codes, %d indexed documents, %d documents of the 09.09 pass"
          % (len(reg_rows), len(prim.reg), len(intakes), sum(1 for x in intakes if x["second"]), len(corpus), len(instances), len(index), len(resolution)))
    t1_t2(intakes, reg_rows, instances, cu_of_p)
    t3_export(export, prim, col2det)
    lots = read_tracker(wbv)
    print("tracker: %d lots, %d cited cells" % (len(lots), sum(len(l["cells"]) for l in lots)))
    t4_tracker(lots, prim)
    t5_export_vs_tracker(export, lots)
    t6_drafts(root, export)
    refs = sorted(glob.glob(os.path.join(root, "tracker", "CoQ_references_v*.csv")))
    t7_references(refs[-1] if refs else "", export)
    t8_register(wbv, export, prim)

    print()
    for k in sorted(COUNT):
        print("  %-40s %6d" % (k, COUNT[k]))
    print()
    by = Counter((f[0], f[1]) for f in FIND)
    print("%d finding(s)" % len(FIND))
    for (chk, what), n in sorted(by.items()):
        print("  %s  %-80s %d" % (chk, what, n))
    if args.out:
        with open(args.out, "w", encoding="utf-8") as fh:
            fh.write("# Truth check of 15.09.2026 — results and cited documents\n\n")
            fh.write("Built by `tracker/truth_check_2026-09-15.py` over the register with the intakes applied, the intake\n"
                     "transcriptions and their second reads, the two-read corpus, the tracker instances, the export,\n"
                     "the recalculated workbook (tracker and CoQ Register), the references table and every compiled draft.\n\n")
            fh.write("## Comparisons\n\n| check | count |\n| --- | ---: |\n")
            for k in sorted(COUNT):
                fh.write("| %s | %d |\n" % (k, COUNT[k]))
            fh.write("\n## Findings — %d\n\n" % len(FIND))
            if not FIND:
                fh.write("None.\n")
            for (chk, what), n in sorted(by.items()):
                fh.write("### %s · %s — %d\n\n" % (chk, what, n))
                for f in FIND:
                    if (f[0], f[1]) == (chk, what):
                        fh.write("* %s\n" % f[2])
                fh.write("\n")
    return 0


if __name__ == "__main__":
    sys.exit(main())
