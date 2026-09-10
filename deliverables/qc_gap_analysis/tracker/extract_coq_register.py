#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Lift the CoQ issue register out of the analysis workbook.

    python3 deliverables/qc_gap_analysis/tracker/extract_coq_register.py \
        CoQ_Analysis_Master_v21.xlsx

The `CoQ Register` sheet is where the desk states, per lot, the code a
certificate of quality will carry and **the date it is issued** — the owner's
ruling of 10.09.2026 is that the certificate prints that date and no other. Both
columns are formulas, and openpyxl stores no computed value for a formula it did
not evaluate, so the sheet is recalculated through LibreOffice first and the
computed values are read from the copy. That is the same step
`extract_artifact_data.py` takes, for the same reason.

The result is a flat source beside the schedule, `coq_register_<date>.csv`, keyed
the way the sheet keys itself: the cultivation batch and `I` for an initial
release or `R` for the reissue. `export_coq_artifact_data.py` reads it, so the
desk, the page and the certificates all take the issue date from one place.
"""
import csv
import os
import shutil
import subprocess
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
GAP = os.path.dirname(HERE)
COLS = [("Key", "key"), ("CU Batch", "cu_batch"), ("P Batch", "p_batch"),
        ("Series", "series"), ("CoQ code", "coq_code"),
        ("Issue date (planned)", "issue_date"), ("Rule date", "rule_date"),
        ("iCoA (register)", "icoa_ref"), ("iCoA issue date", "icoa_issue_date"),
        ("Issuable", "issuable")]


def recalculated(src):
    """A copy of the workbook with every formula evaluated."""
    tmp = tempfile.mkdtemp(prefix="coqreg_")
    shutil.copy(src, os.path.join(tmp, "in.xlsx"))
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx",
                    "--outdir", os.path.join(tmp, "out"),
                    os.path.join(tmp, "in.xlsx")],
                   check=True, capture_output=True)
    out = os.path.join(tmp, "out", "in.xlsx")
    if not os.path.exists(out):
        raise SystemExit("LibreOffice produced no output — is libreoffice-calc installed?")
    return out, tmp


def main(argv):
    src = argv[1] if len(argv) > 1 else os.path.join(HERE, "CoQ_Analysis_Master_v21.xlsx")
    out = argv[2] if len(argv) > 2 else os.path.join(GAP, "coq_register_2026-09-10.csv")
    import openpyxl
    path, tmp = recalculated(src)
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        if "CoQ Register" not in wb.sheetnames:
            raise SystemExit("no CoQ Register sheet in " + src)
        rows = list(wb["CoQ Register"].iter_rows(values_only=True))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    hdr = ["" if v is None else str(v).strip() for v in rows[0]]
    ix = {}
    for want, _ in COLS:
        if want not in hdr:
            raise SystemExit("column missing from the register: " + want)
        ix[want] = hdr.index(want)

    def cell(r, want):
        v = r[ix[want]] if ix[want] < len(r) else None
        if v is None:
            return ""
        # LibreOffice returns the date columns as real dates; the document writes
        # every date DD.MM.YYYY, so they are normalised here rather than at each
        # of the three places a certificate prints one.
        if hasattr(v, "strftime"):
            return v.strftime("%d.%m.%Y")
        return str(v).strip()

    body = []
    for r in rows[1:]:
        if not any(v is not None and str(v).strip() for v in r):
            continue
        if not cell(r, "Key"):
            continue
        body.append([cell(r, w) for w, _ in COLS])
    with open(out, "w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh)
        wr.writerow([n for _, n in COLS])
        wr.writerows(body)
    dated = sum(1 for b in body if b[5] and b[5] != "—")
    print("%s: %d row(s), %d with an issue date" % (os.path.basename(out), len(body), dated))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
