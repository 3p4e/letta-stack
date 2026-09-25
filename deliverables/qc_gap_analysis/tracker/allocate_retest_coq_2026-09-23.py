#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Allocate a certificate-of-quality number to the 35 retests that had none.

## KNOWN DEFECT — this script does not yet produce a correct workbook. Do not run it for real.

`retarget()` moves only the references to the pattern row, so `No.`, whose formula reads
`=IF(OR(C173="yes",...), COUNT(A$1:A172)+1, "")`, keeps pointing at `A172` in every row it
writes. All 35 new rows therefore count the same range and come out as `CoQ-PP_26-172`.
Verified on the workbook this produced: the first 172 code/key pairs were byte-identical to
v56 and the key join was one-to-one, but the new codes were not 173-207.

The fix is to shift **every** row reference that is not `$`-anchored by `to - frm`, not just
the ones naming the pattern row, and to check first whether the `iCoA (register)` lookup
hard-codes its key as a string — if it does, a copied row would look up the pattern row's lot,
which is the same class of fault and would be silent.

The `--check` path is sound and was correct: it reports the 35 and the numbers they would take
without writing anything.

    python3 tracker/allocate_retest_coq_2026-09-23.py            # writes v57 beside v56
    python3 tracker/allocate_retest_coq_2026-09-23.py --check    # report, write nothing

Head of QC, 23.09.2026, asked for "the missing CoQs that miss document code" to be given
"the next available according to their chronological order".

## What was actually missing

Nothing on the CoQ Register: all 172 rows carry a code, because the code is a formula that
fires for `yes`, `allocated` and `ruled` alike. What has no code is **35 retest instances that
have an iCoA row and no CoQ row at all** — iCoA Register rows 174-208, `— at issue —`,
`Issuable = no`, status *"not yet issuable — sampled 19-22.08.2026"*. Those are the 35 this
script gives a number to, by creating the CoQ row that was never there.

## Why `allocated` and not `yes`

`allocated` already exists on this register — thirty rows use it — and it produces a code just
as `yes` does. So the number is reserved and the register does not claim that a certificate was
issued for a retest whose results are not in. The owner authorised numbering; nothing here
authorises issuance, and moving a row to `yes` remains the owner's own step.

## Why the rows are appended and never inserted

`No.` is `=IF(OR(Issuable="yes","allocated","ruled"), COUNT(A$1:A_prev)+1, "")` and the code is
that count formatted — a **running count**, not a stored value. A row inserted in the middle
therefore renumbers every certificate after it. That is not hypothetical: 23 already-issued
certificates currently disagree with v56 about which code belongs to which lot, in two
contiguous runs, which is the signature of exactly this. So the new rows go after the last data
row, they take 173-207, and the first 172 codes are asserted unchanged afterwards.

Chronology is honoured **within the new block**: the 35 are sorted by the sampling date their
status records, so 173 is the earliest and 207 the latest. Chronological order across the whole
register cannot be had without renumbering what is already issued, and is not worth that.

The iCoA Register is left alone. Its 35 counterparts keep `— at issue —`, which is true: the
CoQ number is allocated, the internal certificate is not yet issued. That register has no
`allocated` state and introducing one could break the readers that test `== "yes"`.
"""
import argparse
import datetime
import os
import re
import shutil
import sys

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, "CoQ_Analysis_Master_v56.xlsx")
DST = os.path.join(HERE, "CoQ_Analysis_Master_v57.xlsx")

SAMPLED = re.compile(r"sampled\s+(\d{2})\.(\d{2})\.(\d{4})")
# the columns the new rows carry over from the iCoA row, by CoQ Register column number
CARRY = {10: "Group", 11: "Series", 12: "CU Batch", 13: "P Batch", 14: "Strain", 19: "Key"}
FORMULA_COLS = (1, 2, 4, 5, 8, 9)      # No., CoQ code, Issue date, Rule date, iCoA x2
BLANK = "—"


def awaiting(ws):
    """The iCoA rows that carry no code yet, with the sampling date their status records."""
    head = {ws.cell(1, c).value: c for c in range(1, ws.max_column + 1)}
    out = []
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, head["iCoA code"]).value or "").strip() != "— at issue —":
            continue
        status = str(ws.cell(r, head["Status"]).value or "")
        m = SAMPLED.search(status)
        row = {k: ws.cell(r, head[k]).value for k in
               ("Group", "Series", "CU Batch", "P Batch", "Strain", "Key", "Status")}
        row["_row"] = r
        row["_sampled"] = (datetime.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
                           if m else datetime.date(9999, 12, 31))
        out.append(row)
    # chronological, and Key breaks a tie so two runs of this script agree
    out.sort(key=lambda d: (d["_sampled"], str(d["Key"])))
    return out


def formulas_of(ws, row):
    """The formula pattern of one data row, as a template keyed by column."""
    out = {}
    for c in FORMULA_COLS:
        v = ws.cell(row, c).value
        if not (isinstance(v, str) and v.startswith("=")):
            raise SystemExit("column %d of row %d is not a formula — the sheet has changed" % (c, row))
        out[c] = v
    return out


def retarget(formula, frm, to):
    """The same formula pointed at another row: A173 -> A208, C173 -> C208.

    Only bare column-letter + row references to the pattern row are moved, and `$`-anchored
    ones are left alone, because those are the whole-column lookups into the other register.
    """
    return re.sub(r"(?<![$A-Z0-9])([A-Z]{1,3})%d(?![0-9])" % frm,
                  lambda m: "%s%d" % (m.group(1), to), formula)


def last_data_row(ws):
    """The last row of the table — the merged footnote and any blank line sit below it."""
    for r in range(ws.max_row, 1, -1):
        if str(ws.cell(r, 3).value or "").strip():        # Issuable is always present
            return r
    raise SystemExit("no data rows found on CoQ Register")


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--check", action="store_true", help="report and write nothing")
    ap.add_argument("--src", default=SRC)
    ap.add_argument("--dst", default=DST)
    a = ap.parse_args(argv)

    wb = openpyxl.load_workbook(a.src, data_only=False)
    cq, ic = wb["CoQ Register"], wb["iCoA Register"]
    head = {cq.cell(1, c).value: c for c in range(1, cq.max_column + 1)}

    todo = awaiting(ic)
    last = last_data_row(cq)
    before = [str(cq.cell(r, head["Issuable"]).value or "").strip() for r in range(2, last + 1)]
    numbered = sum(1 for v in before if v in ("yes", "allocated", "ruled"))
    print("CoQ Register: data rows 2-%d, %d of them numbered" % (last, numbered))
    print("iCoA rows awaiting a code: %d" % len(todo))
    print("they would take CoQ-PP_26-%03d .. CoQ-PP_26-%03d"
          % (numbered + 1, numbered + len(todo)))
    print()
    print("%-4s %-16s %-14s %-24s %s" % ("code", "key", "sampled", "strain", "series"))
    for i, d in enumerate(todo):
        print("%-4d %-16s %-14s %-24s %s"
              % (numbered + 1 + i, str(d["Key"]), d["_sampled"].strftime("%d.%m.%Y"),
                 str(d["Strain"])[:24], str(d["Series"])[:34]))
    if a.check:
        return 0

    # the footnote moves down; it is a merged row and must be unmerged before it is cleared
    foot_row, foot_text, foot_span = None, None, None
    for rng in list(cq.merged_cells.ranges):
        if rng.min_row > last:
            foot_row, foot_span = rng.min_row, (rng.min_col, rng.max_col)
            foot_text = cq.cell(rng.min_row, rng.min_col).value
            cq.unmerge_cells(str(rng))
            for c in range(rng.min_col, rng.max_col + 1):
                cq.cell(rng.min_row, c).value = None
            break

    tmpl = formulas_of(cq, last)
    src_style = last
    for i, d in enumerate(todo):
        r = last + 1 + i
        for c in FORMULA_COLS:
            cq.cell(r, c).value = retarget(tmpl[c], src_style, r)
        cq.cell(r, head["Issuable"]).value = "allocated"
        for c, key in CARRY.items():
            cq.cell(r, c).value = d[key]
        cq.cell(r, head["Status"]).value = (
            "allocated — retest sampled %s; number reserved on the Head of QC's instruction "
            "23.09.2026, results not yet recorded, not issued"
            % d["_sampled"].strftime("%d.%m.%Y"))
        cq.cell(r, head["Plan reference (31.08.2026)"]).value = "CoQ retest (assigned on issue)"
        for c in range(1, cq.max_column + 1):
            if c in FORMULA_COLS or c in CARRY:
                continue
            if cq.cell(r, c).value in (None, ""):
                cq.cell(r, c).value = BLANK

    if foot_row is not None:
        new = last + len(todo) + 2
        cq.cell(new, foot_span[0]).value = foot_text
        cq.merge_cells(start_row=new, start_column=foot_span[0],
                       end_row=new, end_column=foot_span[1])

    wb.save(a.dst)
    print()
    print("wrote %s" % a.dst)
    return 0


if __name__ == "__main__":
    sys.exit(main())
