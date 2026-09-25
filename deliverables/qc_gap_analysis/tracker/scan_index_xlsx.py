#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""The scan index as a spreadsheet — Head of QC, 25.09.2026, so it can be read without a terminal.

    python3 tracker/scan_index_xlsx.py

Reads `SCAN_INDEX_2026-09-25.tsv` and writes `SCAN_INDEX_2026-09-25.xlsx`. One sheet, one row per
scan, header frozen and filterable. The `register_agrees` column is the finding of 25.09 and is
coloured where it reads NO, so the disagreement is visible on opening rather than needing a query.
"""
import csv
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, 'SCAN_INDEX_2026-09-25.tsv')
DEST = os.path.join(HERE, 'SCAN_INDEX_2026-09-25.xlsx')

HEAD = PatternFill('solid', fgColor='1F3864')
FLAG = PatternFill('solid', fgColor='FCE4D6')
IDCOL = PatternFill('solid', fgColor='F2F2F2')


def main():
    with open(SRC, encoding='utf-8') as fh:
        rows = list(csv.DictReader(fh, delimiter='\t'))
    cols = list(rows[0].keys())

    wb = Workbook()
    ws = wb.active
    ws.title = 'Approved scans'

    for j, c in enumerate(cols, 1):
        cell = ws.cell(1, j, c)
        cell.font = Font(bold=True, color='FFFFFF', size=9)
        cell.fill = HEAD
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    for i, r in enumerate(rows, 2):
        for j, c in enumerate(cols, 1):
            cell = ws.cell(i, j, r[c])
            cell.font = Font(size=9)
            cell.alignment = Alignment(vertical='top',
                                       wrap_text=c == 'parameters_referenced_named')
            if c == 'register_agrees' and r[c] == 'NO':
                cell.fill = FLAG
                cell.font = Font(size=9, bold=True, color='9C2500')
            if c in ('drive_file_id', 'drive_title', 'drive_bytes'):
                cell.fill = IDCOL

    width = {'batch': 11, 'coq_code': 15, 'issued': 12, 'supersedes': 15,
             'supersedes_issued': 12, 'icoa_cited': 16, 'icoa_cited_issued': 12,
             'icoa_scope': 11, 'parameters_referenced': 12,
             'parameters_referenced_named': 46, 'icoa_per_register': 17,
             'register_agrees': 9, 'pheno': 17, 'pheno_split': 11,
             'drive_file_id': 36, 'drive_title': 20, 'drive_bytes': 11, 'note': 34}
    for j, c in enumerate(cols, 1):
        ws.column_dimensions[get_column_letter(j)].width = width.get(c, 9)
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(cols)), len(rows) + 1)

    wb.save(DEST)
    no = sum(1 for r in rows if r['register_agrees'] == 'NO')
    print('spreadsheet written: %s' % os.path.relpath(DEST, os.path.dirname(HERE)))
    print('  %d scans x %d columns, %d flagged register_agrees=NO, %.0f KiB'
          % (len(rows), len(cols), no, os.path.getsize(DEST) / 1024.0))
    return 0


if __name__ == '__main__':
    raise SystemExit(main())
