#!/usr/bin/env python3
"""Every value v11 prints, keyed to the certificate that must show it on its page."""
import json, re, sys, importlib.util, collections, subprocess, shutil, tempfile, os
import openpyxl
HERE='/home/user/letta-stack/deliverables/qc_gap_analysis/tracker'
S='/tmp/claude-0/-home-user-letta-stack/4877ce6e-ae82-551e-bf35-5698c379c3be/scratchpad'
spec=importlib.util.spec_from_file_location('T', HERE+'/tracker_data.py'); T=importlib.util.module_from_spec(spec); spec.loader.exec_module(T)
tmp=tempfile.mkdtemp(prefix='cl_'); shutil.copy(HERE+'/CoQ_Analysis_Master_v11.xlsx', tmp+'/in.xlsx')
subprocess.run(['soffice','--headless','--calc','--convert-to','xlsx','--outdir',tmp+'/out',tmp+'/in.xlsx'],check=True,capture_output=True,timeout=900)
WV=openpyxl.load_workbook(tmp+'/out/in.xlsx', data_only=True); shutil.rmtree(tmp, ignore_errors=True)
WB=openpyxl.load_workbook(HERE+'/CoQ_Analysis_Master_v11.xlsx')
TR=next(n for n in WB.sheetnames if n.startswith('CoQ Parameter Tracker'))
sh=WB[TR]
starts=[c for c in range(4, sh.max_column+1) if str(sh.cell(2,c).value or '').startswith('#')]
PC={}
for i,s0 in enumerate(starts):
    e=(starts[i+1]-1) if i+1<len(starts) else sh.max_column
    PC[int(re.match(r'#(\d+)', str(sh.cell(2,s0).value)).group(1))]=(s0,e)
anchors=[r for r in range(5, sh.max_row+1) if sh.cell(r,1).value not in (None,'') and not str(sh.cell(r,1).value).startswith('KEY')]
def fill(c):
    try: return c.fill.fgColor.rgb[-6:] if c.fill and c.fill.fill_type=='solid' and isinstance(c.fill.fgColor.rgb,str) else None
    except Exception: return None
out=[]
for a,nx in zip(anchors, anchors[1:]+[sh.max_row+1]):
    cu, p = str(sh.cell(a,1).value), str(sh.cell(a,2).value or '')
    end=a
    for mr in sh.merged_cells.ranges:
        if mr.min_col==1 and mr.min_row==a: end=mr.max_row
    for n,(s0,e) in PC.items():
        single=(e-s0)<=2
        subs=T.GROUPS.get(n) or [str(n)]
        for r in range(a, end+1, 2):
            ref=str(WV[TR].cell(r if single else r+1, s0+1 if single else s0).value or '')
            if not ref or ref.startswith('—') or 'no certificate' in ref: continue
            code=ref.split(',')[0].strip()
            if code.startswith('iCoA') or 'at issue' in code: continue   # in-house, no eCoA page
            m=re.search(r'\((\d{2}\.\d{2}\.\d{4})\)', ref); date=m.group(1) if m else ''
            m=re.search(r'\[([^\]]+)\]', ref); lab=m.group(1) if m else ''
            for j,sub in enumerate(subs):
                c=sh.cell(r, s0+(0 if single else j))
                v=str(WV[TR].cell(c.row, c.column).value or '').strip()
                if not v or v.startswith(('—','not ','n.r.','no result','held','on file')): continue
                out.append({'cu':cu,'p':p,'det':sub,'value':v,'code':code,'date':date,'lab':lab,
                            'state':{'C6EFCE':'release','FCE5CD':'stability','EDEDED':'uncredited'}.get(fill(c),'other'),
                            'row':c.row,'col':c.column})
json.dump(out, open(S+'/checklist.json','w'), ensure_ascii=False)
print('values to verify:', len(out), '| distinct certificates:', len({x['code'] for x in out}))
print('by lab:', collections.Counter(x['lab'] for x in out).most_common())
print('by state:', collections.Counter(x['state'] for x in out).most_common())
