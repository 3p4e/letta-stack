#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""QA's proposed specifications against ours: potency grades, nominal values and ranges. Read only.

    python3 tracker/compare_qa_specs_2026-10-07.py

Head of QC, 07.10.2026: *"check the specifications documents … that you created above and their potency grades and
ranges and compare them with the following folder contents, which contains specifications that my colleague from QA
tried to help us and rectify some inconsistencies and presumably … a wrong potency grade and range. Do not change any
ranges and grades from our just created PDFs, just compare the suggested potency grades, their nominal values and
ranges that QA proposes."*

Inputs:

* **Ours**:
  * `potency_grades_2026-09-15.csv`, the grade table behind the 58 sheets (KVM4 builder, 07.10.2026);
  * `specs/QCSP_001_ImB/INDEX.json`, each sheet's product code;
  * `DELIVER_2026-10-07_T3_Final/CONTENTS.tsv`, the 37 sheets the Tranche 3 PDF holds (the Drive file
    `T3_Specifications_2026-10-07.pdf` is byte for byte the committed one) and each T3 CoQ's printed Total THC.
* **QA's**: `QA_SPEC_PROPOSALS_2026-10-07.json`, what the QA colleague's documents print (Drive folder
  `1TY-W5G2G8l7I6dLS1ITXH0e5WHcfruQW`), read by one agent per strain group and re-read by a second.
* **Results**: `POTENCY_RESULTS_KVM4_2026-10-07.tsv`, every Total THC result on file per strain, the basis of
  "no empty range".

Nothing is changed: not the grade table, not a sheet, not a certificate. Output, `tracker/`:

* `QA_SPEC_COMPARISON_2026-10-07.xlsx`;
* the same as TSVs (`…_qa.tsv`, `…_ours.tsv`, `…_t3.tsv`);
* `QA_SPEC_COMPARISON_2026-10-07.md`, the summary.
"""
import csv
import json
import os
import re
import sys
from collections import defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
GAP = os.path.dirname(HERE)
QA = os.path.join(HERE, 'QA_SPEC_PROPOSALS_2026-10-07.json')
STEM = os.path.join(HERE, 'QA_SPEC_COMPARISON_2026-10-07')
# the strain words in QA's file names
TITLE_ABBR = {'GORILLAGLUE': 'GG', 'GRAPEPIE': 'GP', 'OPM': 'OPM', 'CAPJUNKY': 'CJ', 'CASHCOW': 'CC', 'BSS': 'BSS',
              'BLUEGELATO': 'BG', 'APPLES&BANANA': 'AB', 'ACC': 'ACC', 'FATBASTARD': 'FB', 'J31': 'J31', 'JD': 'JD',
              'KC': 'KC', 'PUM': 'PUM', 'CF': 'CF', 'HPA': 'HPA', 'PM': 'PM', 'WC': 'WC', 'WED': 'WED', 'SJ': 'SJ',
              'SCR': 'SCR', 'MB': 'MB', 'GRC': 'GRC'}


def f2(x):
    return '' if x is None else ('%.2f' % float(x))


def ours():
    grades = list(csv.DictReader(open(os.path.join(GAP, 'potency_grades_2026-09-15.csv'), encoding='utf-8')))
    index = {s['code']: s for s in json.load(open(os.path.join(GAP, 'specs', 'QCSP_001_ImB', 'INDEX.json'),
                                                  encoding='utf-8'))['sheets']}
    t3 = list(csv.DictReader(open(os.path.join(GAP, 'DELIVER_2026-10-07_T3_Final', 'CONTENTS.tsv'), encoding='utf-8'),
                             delimiter='\t'))
    in_t3 = {r['spec'] for r in t3}
    out = []
    for g in grades:
        code = 'QCSP_001_%s-%s_v.01' % (g['abbr'], g['numeral'])
        out.append({'abbr': g['abbr'], 'strain': g['strain'], 'numeral': g['numeral'], 'code': code,
                    'nominal': float(g['nominal']), 'tol': float(g['tolerance']),
                    'low': float(g['window_low']), 'high': float(g['window_high']),
                    'pcode': (index.get(code) or {}).get('product_code', ''), 'in_t3_pdf': code in in_t3})
    return out, t3


def results():
    by = defaultdict(list)
    for r in csv.DictReader(open(os.path.join(HERE, 'POTENCY_RESULTS_KVM4_2026-10-07.tsv'), encoding='utf-8'),
                            delimiter='\t'):
        try:
            by[r['strain']].append((float(r['total_thc']), r['kind'], r['batch'] or r['p_lot']))
        except ValueError:
            pass
    return by


def abbr_of(e, title, known, names):
    a = (e.get('abbr') or '').strip().upper()
    if a in known:
        return a, 'printed'
    t = re.sub(r'^NE\s*-\s*', '', title).replace('ImB_Specification_', '').upper()
    for k in sorted(TITLE_ABBR, key=len, reverse=True):
        if t.startswith(k):
            return TITLE_ABBR[k], 'file name'
    s = (e.get('strain') or '').strip().lower()
    for ab, nm in names.items():
        if s and (s == nm.lower() or s.startswith(nm.lower()) or nm.lower().startswith(s)):
            return ab, 'strain name'
    return '', 'unknown'


def main():
    if not os.path.exists(QA):
        raise SystemExit('%s is missing — the extraction of QA\'s documents writes it' % os.path.basename(QA))
    qa = json.load(open(QA, encoding='utf-8'))
    our, t3 = ours()
    res = results()
    known = {g['abbr'] for g in our}
    names = {g['abbr']: g['strain'] for g in our}
    our_by = defaultdict(list)
    for g in our:
        our_by[g['abbr']].append(g)

    # every QA grade, one row
    rows = []
    for f in qa['files']:
        for e in f.get('entries') or []:
            ab, how = abbr_of(e, f['title'], known, names)
            n, t, lo, hi = e.get('nominal_thc'), e.get('tol_thc'), e.get('range_low'), e.get('range_high')
            m = re.search(r'THC\s*(\d+(?:[.,]\d+)?)', e.get('product_code') or '')
            rows.append({'file': f['title'], 'file_id': f['file_id'], 'ne': f.get('ne_prefix', False),
                         'kind': f['kind'], 'abbr': ab, 'abbr_from': how, 'strain': e.get('strain', ''),
                         'grade': e.get('grade_label', ''), 'spec_code': e.get('spec_code', ''),
                         'pcode': e.get('product_code', ''), 'n': n, 't': t, 'lo': lo, 'hi': hi,
                         'pcode_n': float(m.group(1).replace(',', '.')) if m else None,
                         'text': e.get('potency_text', ''), 'qa_flags': list(e.get('internal_inconsistencies') or [])})

    qa_by = defaultdict(list)
    for r in rows:
        if r['abbr'] and r['n'] is not None and r['kind'] == 'specification' and not r['ne']:
            qa_by[r['abbr']].append(r)

    # remarks QA's documents share, said once in the summary rather than on every grade
    COMMON = (r'footer', r'TEMPLATE', r'mixes precision', r'v03', r'v\.03')
    for r in rows:
        r['qa_flags'] = [f for f in r['qa_flags'] if not any(re.search(c, f) for c in COMMON)]

    def checks(r):
        """What is wrong with a QA grade on its own terms and against the results on file."""
        out = []
        n, t, lo, hi = r['n'], r['t'], r['lo'], r['hi']
        # the builder allows a tolerance up to 10 % of the nominal (solveTolerances, maxTol)
        if n is not None and t is not None and t > round(0.1 * n, 2) + 0.005:
            out.append('tolerance %.2f is more than 10 %% of %.2f (%.2f)' % (t, n, 0.1 * n))
        if None not in (n, t, lo) and abs(lo - (n - t)) > 0.006:
            out.append('low %.2f is not %.2f − %.2f' % (lo, n, t))
        if None not in (n, t, hi) and abs(hi - (n + t)) > 0.006 and abs(hi - (n + t - 0.01)) > 0.006:
            out.append('high %.2f is not %.2f + %.2f' % (hi, n, t))
        if r['pcode_n'] is not None and n is not None and abs(r['pcode_n'] - n) > 0.001:
            out.append('product code says THC%s, nominal %.2f' % (r['pcode_n'], n))
        if lo is not None and hi is not None:
            inside = [x for x in res.get(r['abbr'], []) if lo <= x[0] <= hi]
            r['results_in'] = ', '.join('%.2f %s' % (x[0], x[1]) for x in sorted(inside))
            if not inside:
                out.append('EMPTY: no Total THC result on file falls in %.2f – %.2f' % (lo, hi))
        return out + r['qa_flags']

    for r in rows:
        r['results_in'] = ''
        r['problems'] = checks(r) if r['kind'] == 'specification' else r['qa_flags']
        same = [g for g in our_by.get(r['abbr'], []) if r['n'] is not None and abs(g['nominal'] - r['n']) < 0.001]
        r['ours'] = same[0]['code'] if same else ''
        if same:
            g = same[0]
            diff = []
            if r['t'] is not None and abs(g['tol'] - r['t']) > 0.005:
                diff.append('tolerance %.2f vs ours %.2f' % (r['t'], g['tol']))
            if r['lo'] is not None and abs(g['low'] - r['lo']) > 0.005:
                diff.append('low %.2f vs ours %.2f' % (r['lo'], g['low']))
            # QA closes a range at nominal + tolerance, we at nominal + tolerance - 0.01: the same grade
            if r['hi'] is not None and abs(g['high'] - r['hi']) > 0.005 and abs(g['high'] + 0.01 - r['hi']) > 0.005:
                diff.append('high %.2f vs ours %.2f' % (r['hi'], g['high']))
            if r['pcode'] and re.sub(r'\s', '', r['pcode']) != re.sub(r'\s', '', g['pcode']):
                diff.append('product code %s vs ours %s' % (r['pcode'], g['pcode']))
            r['vs_ours'] = '; '.join(diff) or 'same nominal, tolerance and range'
        else:
            r['vs_ours'] = 'no grade of ours with this nominal' if r['abbr'] else 'strain not identified'

    # overlaps between QA's own grades of a strain, and results no QA range covers
    strain_notes = {}
    for ab, rs in qa_by.items():
        notes = []
        rs = sorted(rs, key=lambda r: r['n'])
        for a, b in zip(rs, rs[1:]):
            if None not in (a['hi'], b['lo']) and b['lo'] <= a['hi']:
                notes.append('QA %s (%s–%s) overlaps %s (%s–%s)' % (a['file'], f2(a['lo']), f2(a['hi']), b['file'],
                                                                    f2(b['lo']), f2(b['hi'])))
        uncovered = [x for x in res.get(ab, [])
                     if not any(r['lo'] is not None and r['hi'] is not None and r['lo'] <= x[0] <= r['hi'] for r in rs)]
        if uncovered:
            notes.append('results in no QA range: ' + ', '.join('%.2f %s (%s)' % x for x in sorted(uncovered)))
        strain_notes[ab] = notes

    # our grades: does QA propose the same?
    our_rows = []
    for g in our:
        q = [r for r in qa_by.get(g['abbr'], []) if abs(r['n'] - g['nominal']) < 0.001]
        our_rows.append(dict(g, qa=', '.join(r['file'] for r in q) or '—',
                             qa_agrees='yes' if q and all(r['vs_ours'] == 'same nominal, tolerance and range' for r in q)
                             else ('differs' if q else 'no QA document for this grade'),
                             results=', '.join('%.2f' % x[0] for x in sorted(res.get(g['abbr'], []))
                                               if g['low'] <= x[0] <= g['high'])))

    # the Tranche 3 certificates: which QA grade would each fall in
    t3_rows = []
    for c in t3:
        ab = re.match(r'QCSP_001_([A-Z0-9]+)-', c['spec']).group(1)
        thc = float(re.match(r'\s*(\d+[.,]\d+)', c['thc']).group(1).replace(',', '.'))
        hit = [r for r in qa_by.get(ab, []) if r['lo'] is not None and r['hi'] is not None and r['lo'] <= thc <= r['hi']]
        g = next(x for x in our if x['code'] == c['spec'])
        t3_rows.append({'coq': c['coq'], 'batch': c['batch'], 'strain': c['strain'], 'thc': thc, 'ours': c['spec'],
                        'ours_window': '%.2f – %.2f' % (g['low'], g['high']), 'ours_pcode': c['pcode'],
                        'qa': ' / '.join('%s (%s ± %s, %s – %s, %s)' % (r['file'], f2(r['n']), f2(r['t']), f2(r['lo']),
                                                                      f2(r['hi']), r['pcode']) for r in hit) or 'no QA range',
                        'same_nominal': 'yes' if any(abs(r['n'] - g['nominal']) < 0.001 for r in hit) else 'no'})

    write(rows, our_rows, t3_rows, strain_notes, qa)
    return 0


def write(rows, our_rows, t3_rows, strain_notes, qa):
    qa_cols = ['file', 'ne', 'kind', 'abbr', 'grade', 'spec_code', 'pcode', 'n', 't', 'lo', 'hi', 'ours', 'vs_ours',
               'problems', 'results_in', 'text', 'file_id']
    our_cols = ['code', 'strain', 'nominal', 'tol', 'low', 'high', 'pcode', 'in_t3_pdf', 'qa', 'qa_agrees', 'results']
    t3_cols = ['coq', 'batch', 'strain', 'thc', 'ours', 'ours_window', 'ours_pcode', 'qa', 'same_nominal']
    flat = lambda v: '; '.join(v) if isinstance(v, list) else ('' if v is None else str(v))
    for name, cols, data in (('qa', qa_cols, rows), ('ours', our_cols, our_rows), ('t3', t3_cols, t3_rows)):
        with open('%s_%s.tsv' % (STEM, name), 'w', encoding='utf-8', newline='') as fh:
            w = csv.writer(fh, delimiter='\t', lineterminator='\n')
            w.writerow(cols)
            for r in data:
                w.writerow([flat(r.get(c)) for c in cols])
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = openpyxl.Workbook()
    for i, (title, cols, data) in enumerate((('QA proposals', qa_cols, rows), ('Our grades', our_cols, our_rows),
                                             ('T3 certificates', t3_cols, t3_rows))):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.append(cols)
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF')
            c.fill = PatternFill('solid', fgColor='1B3A5C')
        for r in data:
            ws.append([flat(r.get(c)) for c in cols])
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = min(60, max(10, max(len(str(c.value or '')) for c in col) + 2))
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.freeze_panes = 'A2'
    ws = wb.create_sheet('Per strain')
    ws.append(['strain', 'notes'])
    for ab in sorted(strain_notes):
        ws.append([ab, '\n'.join(strain_notes[ab]) or '—'])
    wb.save(STEM + '.xlsx')

    live = [r for r in rows if r['kind'] == 'specification' and not r['ne'] and r['abbr']]
    ne = [r for r in rows if r['kind'] == 'specification' and r['ne']]
    same = [r for r in live if r['vs_ours'] == 'same nominal, tolerance and range']
    diff = [r for r in live if r['ours'] and r not in same]
    new = [r for r in live if not r['ours']]
    empty = [r for r in live if any(p.startswith('EMPTY') for p in r['problems'])]
    missing = [o for o in our_rows if o['qa'] == '—']
    t3_same = sum(r['same_nominal'] == 'yes' for r in t3_rows)
    t3_none = sum(r['qa'] == 'no QA range' for r in t3_rows)
    g = lambda r: '%s ± %s (%s – %s)' % (f2(r['n']), f2(r['t']), f2(r['lo']), f2(r['hi']))
    abbr = lambda code: code.split('_')[2].split('-')[0]
    lines = ['# QA\'s proposed specifications against ours — potency grades, nominals and ranges', '',
             'Read only: nothing of ours changed. QA: folder `1TY-W5G2G8l7I6dLS1ITXH0e5WHcfruQW`, %d documents, each read '
             'twice by independent agents. Ours: the 58 sheets on `potency_grades_2026-09-15.csv` (KVM4 builder); the '
             'Tranche 3 PDF holds 37 of them. Detail: `QA_SPEC_COMPARISON_2026-10-07.xlsx`.' % len(qa['files']), '',
             '## In short', '',
             '- QA proposes **%d grades** in %d documents, plus %d in the two documents marked "NE".' % (
                 len(live), len({r['file'] for r in live}), len(ne)),
             '- **%d** are our grade exactly. QA writes the upper limit as nominal + tolerance and we write it as nominal '
             '+ tolerance - 0.01; that counts as the same grade.' % len(same),
             '- **%d** share a nominal with ours but not its tolerance or range.' % len(diff),
             '- **%d** are nominals we have no grade for, and **%d** of our 58 grades have no QA document.' % (
                 len(new), len(missing)),
             '- **%d** QA grades would hold no Total THC result on file (an empty range).' % len(empty),
             '- QA\'s ranges are closed at both ends and set side by side, so neighbouring grades share a boundary value '
             '(17.00 is in both GG 16 and GG 18).',
             '- Every QA document prints the code `…_v.01` with the footer `QCSP 001v03`, and most say "TEMPLATE" in '
             'the header.',
             '- Of the 60 Tranche 3 certificates:',
             '  - **%d** fall in a QA grade of the same nominal as ours;' % t3_same,
             '  - **%d** fall in a QA grade of a different nominal;' % (len(t3_rows) - t3_same - t3_none),
             '  - **%d** fall in no QA range.' % t3_none,
             '', '## Per strain — ours | QA\'s', '']
    for ab in sorted({r['abbr'] for r in live} | {abbr(o['code']) for o in our_rows}):
        o = sorted([x for x in our_rows if abbr(x['code']) == ab], key=lambda x: -x['nominal'])
        q = sorted([r for r in live if r['abbr'] == ab], key=lambda r: r['n'] or 0)
        ours_txt = '; '.join('%s %.2f ± %.2f (%.2f – %.2f)%s' % (x['code'].split('_')[2], x['nominal'], x['tol'], x['low'],
                                                                x['high'], ' ·T3' if x['in_t3_pdf'] else '') for x in o)
        qa_txt = '; '.join(g(r) + (' — same' if r['vs_ours'] == 'same nominal, tolerance and range' and not r['problems']
                                   else ' — ' + '; '.join([r['vs_ours']] + r['problems'])) for r in q)
        lines.append('- **%s** — ours: %s | QA: %s' % (ab, ours_txt or '—', qa_txt or 'no document'))
        if strain_notes.get(ab):
            lines.append('  - ' + ' · '.join(strain_notes[ab]))
    lines += ['', '## The two documents marked "NE" (read, not counted above)', '']
    for r in sorted(ne, key=lambda r: (r['abbr'], r['n'] or 0)):
        lines.append('- %s: %s %s%s' % (r['file'], r['abbr'], g(r), (' — ' + '; '.join(r['problems'])) if r['problems'] else ''))
    if qa.get('disagreements'):
        lines += ['', '## Where the second reading corrected the first (%d fields; the second is used)' % len(qa['disagreements']), '']
        for d in qa['disagreements']:
            lines.append('- %s — %s: first %r, second %r' % (d.get('file_title') or d['file_id'], d['field'],
                                                             str(d['first_read'])[:80], str(d['second_read'])[:80]))
    open(STEM + '.md', 'w', encoding='utf-8').write('\n'.join(lines) + '\n')
    print('\n'.join(lines[:20]))


if __name__ == '__main__':
    sys.exit(main())
