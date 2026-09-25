#!/usr/bin/env python3
"""Third pass: the figures and sentences each sheet states about itself.

The two passes in verify_workbook.py test the tables against each other and against the
certificate records. This one tests what the sheets SAY — the counts in a status cell, the
figures inside a note, the finding beside an audited certificate, the dates against the list
as the Head of QC sent it — because a true table under a false sentence is still a false sheet.
"""
import csv, importlib.util, json, os, re, shutil, subprocess, sys, tempfile, collections
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)
from reference_sections import latest_master, sheet_or_section, has as _readable

SRC = sys.argv[1] if len(sys.argv) > 1 else latest_master(HERE)
if not SRC or not os.path.exists(SRC):
    raise SystemExit("no workbook to check in " + HERE)
spec = importlib.util.spec_from_file_location("tracker_data", os.path.join(HERE, "tracker_data.py"))
T = importlib.util.module_from_spec(spec); spec.loader.exec_module(T)
F = []
def bad(sheet, what, detail=""): F.append((sheet, what, detail))

tmp = tempfile.mkdtemp(prefix="prose_")
shutil.copy(SRC, os.path.join(tmp, "in.xlsx"))
subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir",
                os.path.join(tmp, "out"), os.path.join(tmp, "in.xlsx")], check=True, capture_output=True, timeout=900)
WV = openpyxl.load_workbook(os.path.join(tmp, "out", "in.xlsx"), data_only=True)
shutil.rmtree(tmp, ignore_errors=True)
WB = openpyxl.load_workbook(SRC)
TR = next(n for n in WB.sheetnames if n.startswith("CoQ Parameter Tracker"))
sh = WB[TR]
from openpyxl.utils import get_column_letter as L

# ---------------------------------------------------------------- the tracker's own blocks
starts = [c for c in range(4, sh.max_column + 1) if str(sh.cell(2, c).value or "").startswith("#")]
PCOL = {}
for i, s0 in enumerate(starts):
    e = (starts[i + 1] - 1) if i + 1 < len(starts) else sh.max_column
    PCOL[int(re.match(r"#(\d+)", str(sh.cell(2, s0).value)).group(1))] = (s0, e)
anchors = [r for r in range(5, sh.max_row + 1) if sh.cell(r, 1).value not in (None, "")
           and not str(sh.cell(r, 1).value).startswith("KEY")]
LOT = {}
for a, nxt in zip(anchors, anchors[1:] + [sh.max_row + 1]):
    end = a
    for mr in sh.merged_cells.ranges:
        if mr.min_col == 1 and mr.min_row == a:
            end = mr.max_row
    refs, credited, uncredited = collections.defaultdict(list), collections.defaultdict(list), set()
    for n, (s0, e) in PCOL.items():
        single = (e - s0) <= 2
        for r in range(a, end + 1, 2):
            t = str(WV[TR].cell(r if single else r + 1, s0 + 1 if single else s0).value or "")
            if not t or "no certificate" in t:
                continue
            refs[n].append(t)
            # Coverage is the green fill, which is how the tracker and
            # verify_workbook.py both define it. This used to skip any reference
            # opening with an em dash, which is the shape of every in-house
            # result ("— at issue —, (16.02.2026) [PP]"): identification A and B
            # and foreign matter are performed in house on every lot, so the
            # recount dropped three covered parameters per lot and read the
            # STATUS cell as undercounting when the cell was right.
            c = sh.cell(r, s0)
            fill = (c.fill.fgColor.rgb[-6:] if c.fill and c.fill.fill_type == "solid"
                    and isinstance(c.fill.fgColor.rgb, str) else None)
            if fill == "C6EFCE":
                credited[n].append(t)
            elif fill == "EDEDED":
                uncredited.add(t.split(",")[0].strip())
    LOT[(str(sh.cell(a, 1).value), str(sh.cell(a, 2).value or ""))] = {
        "row": a, "end": end, "blocks": (end - a + 1) // 2, "refs": refs, "credited": credited,
        "uncredited": uncredited,
        "status": str(WV[TR].cell(a, 3).value or "")}

# ---------------------------------------------------------------- 1. the tracker's STATUS cell
for (cu, p), d in LOT.items():
    st = d["status"]
    # the tracker counts a parameter as NO RESULT unless a credited certificate reports a release
    # result; Batch Coverage now says the same, so the check reads the coverage row
    miss = [n for n in range(1, 13) if not d["credited"].get(n)]
    m = re.search(r"(\d+) NO RESULT", st)
    parts = re.search(r"\((\d+) no cert / (\d+) cert w/o result(?: / (\d+) stability only)?\)", st)
    if m and parts:
        tot = sum(int(x) for x in parts.groups() if x)
        if int(m.group(1)) != tot:
            bad(TR, "STATUS's total does not equal the buckets it names", f"{cu}/{p}: {st[:60]!r}")
    if m and int(m.group(1)) < len(miss):
        bad(TR, "STATUS names fewer missing parameters than the block shows", f"{cu}/{p}: says {m.group(1)}, block shows {len(miss)}")
    if not m and not st.startswith("✓ COMPLETE"):
        bad(TR, "STATUS is neither COMPLETE nor a NO RESULT count", f"{cu}/{p}: {st[:40]!r}")
    m = re.search(r"(\d+) testing instances", st)
    if m and int(m.group(1)) != d["blocks"]:
        bad(TR, "STATUS names a different number of testing instances", f"{cu}/{p}: says {m.group(1)}, block has {d['blocks']}")
    m = re.search(r"• (\d+) document\(s\) on file, not credited", st)
    # an uncredited document is a grey cell, the tracker's own definition; the
    # suffix in the text follows the fill and is not what is counted
    extra = len(d["uncredited"])
    if m and int(m.group(1)) != extra:
        bad(TR, "STATUS names a different number of uncredited documents", f"{cu}/{p}: says {m.group(1)}, block has {extra}")

# ---------------------------------------------------------------- 2. Batch Coverage: the certificate count and the laboratories
cov = WB["Batch Coverage"]
hdr = [str(c.value or "") for c in cov[1]]
for r in range(2, cov.max_row + 1):
    cu, p = str(cov.cell(r, 1).value or ""), str(cov.cell(r, 2).value or "")
    if not cu:
        continue
    key = next((k for k in LOT if re.sub(r"[＊*]", "", k[0]) == re.sub(r"[＊*]", "", cu)
                and (k[1] == p or (p.startswith("—") and k[1].startswith(("N/A", "—"))))), None)
    if key is None:
        continue
    d = LOT[key]
    codes = {t.split(",")[0].strip() for n in d["credited"] for t in d["credited"][n]}
    codes = {c for c in codes if not c.startswith("iCoA") and "at issue" not in c}
    n_sheet = cov.cell(r, 19).value
    if n_sheet is not None and str(n_sheet).isdigit() and int(n_sheet) < len(codes):
        bad("Batch Coverage", "certificate count is lower than the tracker's credited certificates",
            f"{cu}/{p}: sheet {n_sheet}, tracker {len(codes)}")
    labs_sheet = str(cov.cell(r, 20).value or "")
    labs_track = {re.search(r"\[([^\]]+)\]", t).group(1) for n in d["credited"] for t in d["credited"][n]
                  if re.search(r"\[([^\]]+)\]", t)}
    for lab in labs_track:
        if lab not in ("PP",) and lab not in labs_sheet:
            bad("Batch Coverage", "a laboratory on the tracker is not in the Labs column", f"{cu}/{p}: {lab} missing from {labs_sheet!r}")

# ---------------------------------------------------------------- 3. Mikro CoQ Parameter against the tracker
if _readable(WB, "Mikro CoQ Parameter"):
    mk = sheet_or_section(WB, "Mikro CoQ Parameter")
    mkv = sheet_or_section(WV, "Mikro CoQ Parameter")
    mst = [c for c in range(4, mk.max_column + 1) if str(mk.cell(2, c).value or "").startswith("#")]
    MP = {}
    for i, s0 in enumerate(mst):
        e = (mst[i + 1] - 1) if i + 1 < len(mst) else mk.max_column
        MP[int(re.match(r"#(\d+)", str(mk.cell(2, s0).value)).group(1))] = (s0, e)
    if set(MP) - {7, 8, 9, 10, 11, 12}:
        bad("Mikro CoQ Parameter", "carries a parameter outside #7–#12", str(sorted(set(MP) - {7, 8, 9, 10, 11, 12})))
    m_an = [r for r in range(5, mk.max_row + 1) if mk.cell(r, 1).value not in (None, "") and not str(mk.cell(r, 1).value).startswith("KEY")]
    checked = 0
    for a in m_an:
        cu, p = str(mk.cell(a, 1).value), str(mk.cell(a, 2).value or "")
        key = next((k for k in LOT if re.sub(r"[＊*]", "", k[0]) == re.sub(r"[＊*]", "", cu) and (k[1] == p or p in k[1])), None)
        if key is None:
            bad("Mikro CoQ Parameter", "row has no lot on the tracker", f"{cu}/{p}")
            continue
        t0 = LOT[key]["row"]
        for n, (s0, e) in MP.items():
            single = (e - s0) <= 2
            ts0, te = PCOL[n]
            tsingle = (te - ts0) <= 2
            for j in range(0, (1 if single else e - s0)):
                a_v = str(mkv.cell(a, s0 + j).value or "")
                b_v = str(WV[TR].cell(t0, ts0 + j).value or "")
                if a_v or b_v:
                    checked += 1
                    if a_v != b_v:
                        bad("Mikro CoQ Parameter", f"#{n} first-instance value differs from the tracker",
                            f"{cu}/{p} col {j}: mikro {a_v[:22]!r}, tracker {b_v[:22]!r}")
    print(f"Mikro: {checked} cell(s) compared with the tracker")

# ---------------------------------------------------------------- 4. Credit Audit findings against the corpus
CORPUS = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(HERE))), "ingestion", "ecoa_runner", "records_corpus.json")
recs = json.load(open(CORPUS)) if os.path.exists(CORPUS) else []
in_corpus = {T.nkey(str(r.get("cert_code") or "")) for r in recs}
aud = sheet_or_section(WB, "Credit Audit")
for r in range(2, aud.max_row + 1):
    code = str(aud.cell(r, 3).value or "")
    why = str(aud.cell(r, 8).value or "")
    if not code or code.startswith("Head of QC"):
        continue
    if why == "not ingested" and T.nkey(code) in in_corpus:
        bad("Credit Audit", "says 'not ingested' but the corpus holds the certificate", f"row {r}: {code}")
    lab = str(aud.cell(r, 5).value or "")
    if why == "not on this certificate" and T.nkey(code) not in in_corpus and lab not in ("PP", "NGP"):
        # an in-house form is never in the eCoA corpus; the finding there is the desk's own reading
        bad("Credit Audit", "says 'not on this certificate' but the corpus does not hold it", f"row {r}: {code}")

# ---------------------------------------------------------------- 5. Work Order against the audit and the registers
wo = sheet_or_section(WB, "Work Order")
audit_codes = {str(aud.cell(r, 3).value or "") for r in range(2, aud.max_row + 1)}
for r in range(2, wo.max_row + 1):
    task = str(wo.cell(r, 1).value or "")
    code = str(wo.cell(r, 4).value or "")
    if not task or task.startswith("Head of QC"):
        continue
    # A Work Order row need not be about a certificate at all. The strain-name rows added
    # 07.09.2026 name no document — they ask a person to rule on a spelling — and an em dash
    # in the certificate column is how such a row says so.
    if code.strip() in ("", "—", "-"):
        continue
    if code and code not in audit_codes and "lot" not in task.lower():
        bad("Work Order", "names a certificate that is not on the Credit Audit", f"row {r}: {task[:30]} {code}")

# ---------------------------------------------------------------- 6. the notes' own figures
def note_of(name):
    """The footnote under a table: one long string in the first column, nothing beside it.

    It used to be recognised by the opening words "Head of QC", so rewriting a
    note's first sentence made this return "" — and every check that reads the
    note then passed on an empty string instead of failing.
    """
    s = sheet_or_section(WB, name)
    for row in s.iter_rows():
        if isinstance(row[0].value, str) and len(row[0].value) > 200 and not any(
                c.value not in (None, "") for c in row[1:]):
            return row[0].value
    return ""


def rows_of(name, key_row=1):
    s = WV[name]
    hdr = [str(c.value or "").strip() for c in s[key_row]]
    out = []
    for r in range(2, s.max_row + 1):
        vals = [s.cell(r, c).value for c in range(1, len(hdr) + 1)]
        if not any(v not in (None, "") for v in vals[:3]) or str(vals[0] or "").startswith("Head of QC"):
            continue
        out.append(dict(zip(hdr, vals)))
    return out


ic = rows_of("iCoA Register"); cq = rows_of("CoQ Register")
icn = [d for d in ic if d["No."] not in (None, "")]
cqn = [d for d in cq if d["No."] not in (None, "")]
if "iCoA-PP_26-nnn" not in note_of("iCoA Register"):
    bad("iCoA Register", "the note does not state the code series", "")


def fmt(v):
    return v.strftime("%d.%m.%Y") if hasattr(v, "strftime") else ("" if v is None else str(v))


# The legacy backlog is released on one day, the note explains that day, and the
# two have to agree. The day is read off the rows and looked for in the note; it
# used to be a literal here, so the rows were compared against a third copy of the
# date that no sheet held, and the check passed or failed on whether someone had
# updated the checker rather than on whether the sheet contradicted itself. Only
# the release round is held to the blanket day: a legacy lot's retest is sampled
# in its own campaign and issued then, which is the ruling of 05.09.2026.
for name, rows in (("iCoA Register", icn), ("CoQ Register", cqn)):
    days = {fmt(d["Issue date (planned)"]) for d in rows
            if str(d["Group"]) == "legacy" and str(d["Series"]) == "initial release"}
    txt = note_of(name)
    for day in sorted(d for d in days if d):
        if day not in txt:
            bad(name, "the note never names the day the legacy series is released on",
                "rows say %s; the note names %s" % (day, sorted(set(re.findall(r"\d{2}\.\d{2}\.\d{4}", txt)))))
    # A Status that says "issued <day>" is a sentence about its own row: the day it
    # names is the row's planned issue date or the sentence is false. The builder
    # wrote this day as a literal too, so a register could issue on 06.06.2026 and
    # say 27.05.2026 in the same row.
    #
    # A Status is a semicolon-joined list of clauses, and not every clause speaks
    # about this row's own certificate: a retest row ends "...; iCoA tested
    # 14.08.2026, issued 17.08.2026", which dates the INTERNAL certificate of
    # analysis, a different document the register dates in its own iCoA issue date
    # column. Reading the whole Status as one sentence made that clause answer for
    # the row, and it reported CoQ-PP_26-171 and -172 as false when both are true —
    # they are two of the seven the register withheld on 18.09.2026, so they
    # genuinely carry no planned issue day yet, while their iCoA issue date column
    # holds exactly the 17.08.2026 the clause names. So: read clause by clause, and
    # hold an iCoA clause against the iCoA column rather than skipping it, because
    # a clause nothing checks is how the wrong date gets in.
    for d in rows:
        for clause in str(d.get("Status") or "").split(";"):
            said = re.search(r"issued (?:with the legacy series on )?(\d{2}\.\d{2}\.\d{4})", clause)
            if not said:
                continue
            about_icoa = "iCoA" in clause and name != "iCoA Register"
            col = "iCoA issue date" if about_icoa else "Issue date (planned)"
            if col not in d:
                continue
            if said.group(1) != fmt(d[col]):
                bad(name,
                    "Status names an iCoA issue day that is not the row's" if about_icoa
                    else "Status names an issue day that is not the row's",
                    "%s: Status says %s, row's %s is %s"
                    % (d.get("iCoA code") or d.get("CoQ code"), said.group(1), col, fmt(d[col])))

# ---------------------------------------------------------------- 7. Batch Dates against the list as it was sent
raw = os.path.join(HERE, "batch_dates_raw_2026-09-04.tsv")
if os.path.exists(raw):
    sent = [r for r in csv.DictReader(open(raw, encoding="utf-8"), delimiter="\t")]
    bd = rows_of("Batch Dates")
    if len(bd) != len(sent):
        bad("Batch Dates", "row count differs from the list as sent", f"sheet {len(bd)}, sent {len(sent)}")
    by_seq = {str(d["Seq"]): d for d in bd}
    for s_ in sent:
        d = by_seq.get(str(int(s_["#"])))
        if not d:
            bad("Batch Dates", "a row of the sent list is missing", s_["#"])
            continue
        cu = re.sub(r"\s*-\s*R&D$", "", s_["Batch"].strip())
        if str(d["Batch (as listed)"]) != cu:
            bad("Batch Dates", "batch name differs from the list as sent", f"{s_['#']}: sheet {d['Batch (as listed)']!r}, sent {cu!r}")
        if str(d["P Batch"] or "") != (s_["Packaging batch"].strip() if s_["Packaging batch"].strip().startswith("P") else ""):
            bad("Batch Dates", "P batch differs from the list as sent", f"{s_['#']}: sheet {d['P Batch']!r}, sent {s_['Packaging batch']!r}")
        # the printed span must contain the normalised dates
        for col, printed in (("Packaging from", s_["Packaging date"]), ("Harvest from", s_["Date of harvest"])):
            v = fmt(d[col])
            if v.startswith("—") or printed.strip() in ("0", "]", ""):
                continue
            # the day and the month the sheet carries must be numbers the list prints for that batch
            nums = {int(x) for x in re.findall(r"\d+", printed)}
            if int(v[:2]) not in nums or int(v[3:5]) not in nums:
                bad("Batch Dates", f"{col} is not a date the list prints", f"{s_['#']} {cu}: sheet {v}, list {printed!r}")

print(f"\n{len(F)} finding(s)")
for s, w, d in F:
    print(f"  [{s}] {w}" + (f" — {d}" if d else ""))

# This pass had no exit code, so it printed its findings and returned 0 — and the
# workflow step that runs it passed no matter what it found. That is the same flaw
# the workflow's own comment records for verify_workbook.py ("a verifier that
# cannot fail is a report, not a check"), fixed there and never here: two findings
# on v48 rode a green CI, and both turned out to be the checker misreading a
# clause rather than the register lying, which is exactly the kind of thing that
# only gets looked at once something fails.
if F:
    print(f"\nFAILED: {len(F)} finding(s) on what the sheets say about themselves")
    sys.exit(1)
print("\nprose verified: no findings")
sys.exit(0)
