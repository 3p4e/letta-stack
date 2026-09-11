#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Lift the harvest and packaging dates out of the analysis workbook.

    python3 deliverables/qc_gap_analysis/tracker/extract_batch_dates.py

The owner's ruling of 10.09.2026 dates the internal certificate of analysis on
the batch's **packaging date**: "start and end date of testing that is the same
date of packaging of the production batch". The desk was taking that from the
release record, which carries it only for the 48 lots that have a P number — so
half the internal certificates had no testing date and fell back to whatever
laboratory date the round happened to end on.

`Batch Dates` in the workbook has it for every batch, and it has it as a **range**
— `Packaging from` and `Packaging to` — which is what "start and end date of
testing" actually wants. This lifts it to a flat source beside the schedule.
"""
import csv
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
GAP = os.path.dirname(HERE)
OUT = os.path.join(GAP, "batch_dates_2026-09-10.csv")
SHEET = "Batch Dates"
COLS = [("Batch (as listed)", "batch"), ("P Batch", "p_batch"),
        ("Harvest from", "harvest_from"), ("Harvest to", "harvest_to"),
        ("Packaging from", "packaging_from"), ("Packaging to", "packaging_to")]


def main(argv):
    src = argv[1] if len(argv) > 1 else os.path.join(HERE, "CoQ_Analysis_Master_v21.xlsx")
    import openpyxl
    wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
    if SHEET not in wb.sheetnames:
        raise SystemExit("no %s sheet in %s" % (SHEET, src))
    rows = list(wb[SHEET].iter_rows(values_only=True))
    hdr = ["" if v is None else str(v).strip() for v in rows[0]]
    ix = {}
    for want, _ in COLS:
        if want not in hdr:
            raise SystemExit("column missing from %s: %s" % (SHEET, want))
        ix[want] = hdr.index(want)

    def cell(r, want):
        v = r[ix[want]] if ix[want] < len(r) else None
        if v is None:
            return ""
        # the sheet holds real dates; every Purely Plant document writes DD.MM.YYYY
        if hasattr(v, "strftime"):
            return v.strftime("%d.%m.%Y")
        return str(v).strip()

    body = []
    for r in rows[1:]:
        if not cell(r, "Batch (as listed)"):
            continue
        body.append([cell(r, w) for w, _ in COLS])
    with open(OUT, "w", newline="", encoding="utf-8") as fh:
        wr = csv.writer(fh)
        wr.writerow([n for _, n in COLS])
        wr.writerows(body)
    packed = sum(1 for b in body if b[4])
    print("%s: %d batch(es), %d with a packaging date"
          % (os.path.basename(OUT), len(body), packed))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
