#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""The list of the 12-month retest certificates of quality, tranches 1, 2 and 3.

    python3 deliverables/qc_gap_analysis/retest_list.py
        -> tracker/CoQ_Retest_List_<today>.xlsx / .csv / .md

The Head of QC, 17.09.2026: "a list of all the certificates of quality for the retest of
tranche one, two and three, in ascending order, with batch number, strain, product type,
potency grade and range, and the actual THC result of the CoQ — in a logical column order."

Read from the certificates' own data (coq_artifact_data.json), so the list says what the
pages say: the certificate code and issue date, the release certificate it supersedes, the
production and cultivation batch, the strain, the product type as Section 01 prints it
(phenotype, chemotype, processing), the potency grade with its product code and the
specification window, and the Total Δ⁹-THC, CBD and CBN the certificate prints with the
document they come from. Ordered by certificate code; the two not yet numbered last.
"""
import csv, datetime as dt, io, json, os, re, sys
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(os.path.dirname(HERE))
SRC = os.path.join(HERE, "coq_artifact_data.json")
SCOPE = os.path.join(HERE, "tracker", "coq_reissue_scope_2026-09-15.csv")
COLS = ["Tranche", "CoQ code", "Issued", "Supersedes", "Superseded issued", "Production batch (P lot)", "Cultivation batch",
        "Strain", "Product type", "Grade", "Product code", "Potency window (%)", "Nominal ± tol (%)",
        "Total THC result (%)", "CBD (%)", "CBN (%)", "Potency certificate", "Certificate issued", "Laboratory",
        "Retest campaign", "Internal CoA", "Internal CoA issued"]


def tranche_of():
    out = {}
    if os.path.exists(SCOPE):
        for r in csv.DictReader(io.open(SCOPE, encoding="utf-8-sig")):
            out[(r.get("coq_code") or "").strip()] = (r.get("tranche") or "").strip()
            out[(r.get("p_lot") or "").strip()] = (r.get("tranche") or "").strip()
            out[(r.get("batch_as_delivered") or "").strip()] = (r.get("tranche") or "").strip()
    return out


def window(crit):
    m = re.search(r"([\d.]+)\s*[–-]\s*([\d.]+)\s*%", crit or "")
    n = re.search(r"nominal\s*([\d.]+)\s*±\s*([\d.]+)", crit or "")
    return ("%s – %s" % (m.group(1), m.group(2)) if m else ""), ("%s ± %s" % (n.group(1), n.group(2)) if n else "")


def rows():
    d = json.load(io.open(SRC, encoding="utf-8")); tr = tranche_of(); out = []
    for c in d["coqs"]:
        if "additional" not in (c.get("t") or ""):
            continue
        r4 = next((r for r in c["rows"] if r["no"] == "4"), {})
        r5 = next((r for r in c["rows"] if r["no"] == "5"), {}); r6 = next((r for r in c["rows"] if r["no"] == "6"), {})
        spc = c.get("spc") or {}
        ptype = " · ".join(x for x in (spc.get("pheno"), spc.get("chemo"), spc.get("proc")) if x)
        if spc.get("dominance"):
            ptype += " (%s)" % spc["dominance"]
        w, nt = window(r4.get("crit"))
        sup = c.get("supersedes") or {}
        code = c.get("regcode") or ""
        t = tr.get(code) or tr.get(c.get("pp") or "") or tr.get(c.get("cb") or "") or ("3" if "T3" in (c.get("icoa_campaign") or "") else "")
        out.append({
            "Tranche": t, "CoQ code": code, "Issued": c.get("issue") or "",
            "Supersedes": sup.get("code") or "", "Superseded issued": sup.get("date") or "",
            "Production batch (P lot)": c.get("pp") or "", "Cultivation batch": c.get("cb") or "",
            "Strain": c.get("strain") or "", "Product type": ptype, "Grade": ("Grade %s" % c["grade"]) if c.get("grade") else "",
            "Product code": c.get("pcode") or "", "Potency window (%)": w, "Nominal ± tol (%)": nt,
            "Total THC result (%)": (r4.get("res") or "").replace("%", "").strip(), "CBD (%)": (r5.get("res") or "").strip(), "CBN (%)": (r6.get("res") or "").strip(),
            "Potency certificate": r4.get("doc") or "", "Certificate issued": r4.get("dd") or "", "Laboratory": r4.get("lab") or "",
            "Retest campaign": ("Farmahem %s-series" % c["icoa_campaign"]) if c.get("icoa_campaign") else "",
            "Internal CoA": c.get("icoa_code") or "", "Internal CoA issued": c.get("icoa_issue") or ""})
    def key(r):
        m = re.search(r"_26-(\d+)$", r["CoQ code"])
        return (0, int(m.group(1))) if m else (1, r["Production batch (P lot)"] or r["Cultivation batch"])
    return sorted(out, key=key)


def write(R, base):
    with io.open(base + ".csv", "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=COLS); w.writeheader(); w.writerows(R)
    L = ["# Certificates of quality — 12-month retest, tranches 1, 2 and 3", "",
         "%d certificates, ordered by certificate code. Written %s from the certificates' own data by `retest_list.py`." % (len(R), dt.date.today().strftime("%d.%m.%Y")), "",
         "| " + " | ".join(COLS) + " |", "| " + " | ".join("---" for _ in COLS) + " |"]
    for r in R:
        L.append("| " + " | ".join(str(r[c]).replace("|", "·") for c in COLS) + " |")
    io.open(base + ".md", "w", encoding="utf-8").write("\n".join(L) + "\n")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    def sheet(ws, rows_):
        ws.append(COLS)
        for r in rows_:
            ws.append([r[c] for c in COLS])
        hdr = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="1B3A5C")
        for cell in ws[1]:
            cell.font = hdr; cell.fill = fill; cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "C2"; ws.auto_filter.ref = ws.dimensions; ws.row_dimensions[1].height = 30
        widths = {"Tranche": 8, "CoQ code": 15, "Issued": 11, "Supersedes": 15, "Superseded issued": 11, "Production batch (P lot)": 12, "Cultivation batch": 14,
                  "Strain": 22, "Product type": 34, "Grade": 9, "Product code": 18, "Potency window (%)": 15, "Nominal ± tol (%)": 14,
                  "Total THC result (%)": 12, "CBD (%)": 9, "CBN (%)": 9, "Potency certificate": 16, "Certificate issued": 11, "Laboratory": 22,
                  "Retest campaign": 20, "Internal CoA": 16, "Internal CoA issued": 11}
        for i, c in enumerate(COLS, 1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 14)
        thin = Side(style="thin", color="D9E0E8")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = Border(bottom=thin)
                if cell.column_letter in ("A", "J", "N", "O", "P"):
                    cell.alignment = Alignment(horizontal="center")
    ws = wb.active; ws.title = "Retest CoQs — all"; sheet(ws, R)
    for t in ("1", "2", "3"):
        sheet(wb.create_sheet("Tranche %s" % t), [r for r in R if r["Tranche"] == t])
    wb.save(base + ".xlsx")


def main(argv):
    R = rows(); base = os.path.join(HERE, "tracker", "CoQ_Retest_List_%s" % dt.date.today().isoformat())
    write(R, base)
    import collections
    print("retest certificates:", len(R), dict(collections.Counter(r["Tranche"] for r in R)))
    print("wrote", os.path.relpath(base, ROOT) + ".xlsx / .csv / .md"); return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
