#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""One register of every certificate of quality, by document code, with its lineage and its references.

    python3 deliverables/qc_gap_analysis/coq_master_register.py
        -> tracker/CoQ_Master_Register_<today>.xlsx / .csv / .md

The Head of QC, 17.09.2026: one table ordered by the certificate of quality's own document
code, carrying the supersession — which certificate a reissue supersedes, and which reissue
supersedes an initial certificate — and, for the data reference, the internal certificate of
analysis each one cites, with its code and its date of issue. For every certificate, initial
issue and reissue alike.

Read from the certificates' own data (`coq_artifact_data.json`) and the internal-certificate
register (`icoa_register_2026-09-10.csv`), so the register says what the pages say.

The columns run identity, lineage, batch, product, result, references:

* **identity** — the code, the testing round, the date of issue, what the CoQ Register says
  about issuing it, and the legacy number the document carried before the 26-series;
* **lineage** — for a reissue, the release certificate it supersedes and that certificate's
  date; for a release certificate, the reissue that supersedes it and its date. The join is
  the reissue's own supersedes line, and the lot where a pair is not yet numbered;
* **batch** — tranche, production batch, cultivation batch, strain;
* **product** — the type as Section 01 prints it, the grade, the product code, the
  specification window and the nominal with its tolerance;
* **result** — Total Δ⁹-THC, CBD and CBN as the certificate prints them, with the document
  they come from, its date and its laboratory;
* **references** — the internal certificate of analysis: its code, its date of issue, the
  day its determinations were made, and the parameters it covers; and the retest campaign.

Three sheets: every certificate in code order, the release round alone, the retest round
alone. The Markdown carries the same table for reading on GitHub.
"""
import csv, datetime as dt, io, json, os, re, sys
HERE = os.path.dirname(os.path.abspath(__file__)); ROOT = os.path.dirname(os.path.dirname(HERE))
SRC = os.path.join(HERE, "coq_artifact_data.json")
ICOA = os.path.join(HERE, "icoa_register_2026-09-10.csv")
SCOPE = os.path.join(HERE, "tracker", "coq_reissue_scope_2026-09-15.csv")

COLS = ["CoQ code", "Testing round", "Issued", "Register status", "Legacy code",
        "Supersedes", "Superseded of", "Superseded by", "Reissue issued",
        "Tranche", "Production batch (P lot)", "Cultivation batch", "Strain", "Product type",
        "Grade", "Product code", "Potency window (%)", "Nominal ± tol (%)",
        "Total THC (%)", "CBD (%)", "CBN (%)", "Potency certificate", "Certificate issued", "Laboratory",
        "Internal CoA (iCoA)", "iCoA issued", "iCoA tested", "iCoA covers parameters", "Retest campaign"]
STATUS = {"yes": "issued", "allocated": "code allocated, date provisional", "no": "not yet issuable"}


def nk(s):
    return re.sub(r"[\s\-/_.＊*]+", "", str(s or "")).upper()


def lotkey(c):
    return nk(c.get("pp") or c.get("cb") or "")


def order(code, lot):
    m = re.search(r"_26-(\d+)$", code or "")
    return (0, int(m.group(1))) if m else (1, lot or "")


def window(crit):
    m = re.search(r"([\d.]+)\s*[–-]\s*([\d.]+)\s*%", crit or "")
    n = re.search(r"nominal\s*([\d.]+)\s*±\s*([\d.]+)", crit or "")
    return ("%s – %s" % (m.group(1), m.group(2)) if m else ""), ("%s ± %s" % (n.group(1), n.group(2)) if n else "")


def tranche_of():
    out = {}
    if os.path.exists(SCOPE):
        for r in csv.DictReader(io.open(SCOPE, encoding="utf-8-sig")):
            t = (r.get("tranche") or "").strip()
            for k in ("coq_code", "p_lot", "batch_as_delivered"):
                v = (r.get(k) or "").strip()
                if v:
                    out[v] = t
    return out


def icoa_index():
    out = {}
    if os.path.exists(ICOA):
        for r in csv.DictReader(io.open(ICOA, encoding="utf-8")):
            out[(r.get("code") or "").strip()] = r
    return out


def rows():
    d = json.load(io.open(SRC, encoding="utf-8")); tr = tranche_of(); ico = icoa_index()
    retests = [c for c in d["coqs"] if "additional" in (c.get("t") or "")]
    by_super = {(c.get("supersedes") or {}).get("code"): c for c in retests
                if str((c.get("supersedes") or {}).get("code") or "").startswith("CoQ-PP")}
    by_lot = {lotkey(c): c for c in retests}
    out = []
    for c in d["coqs"]:
        ret = "additional" in (c.get("t") or "")
        r4 = next((r for r in c["rows"] if r["no"] == "4"), {})
        r5 = next((r for r in c["rows"] if r["no"] == "5"), {}); r6 = next((r for r in c["rows"] if r["no"] == "6"), {})
        spc = c.get("spc") or {}
        ptype = " · ".join(x for x in (spc.get("pheno"), spc.get("chemo"), spc.get("proc")) if x)
        if spc.get("dominance") and spc["dominance"] != "TO BE DETERMINED":
            ptype += " (%s)" % spc["dominance"]
        w, nt = window(r4.get("crit"))
        sup = c.get("supersedes") or {}
        rt = None if ret else (by_super.get(c.get("regcode")) or by_lot.get(lotkey(c)))
        ic = ico.get(str(c.get("icoa_code") or "").strip(), {})
        out.append({
            "CoQ code": c.get("regcode") or "",
            "Testing round": "12-month retest (reissue)" if ret else "Initial release",
            "Issued": c.get("issue") or "",
            "Register status": STATUS.get((c.get("reg_issuable") or "").strip(), c.get("reg_issuable") or ""),
            "Legacy code": c.get("n") or "",
            "Supersedes": sup.get("code") or "",
            "Superseded of": sup.get("date") or "",
            "Superseded by": (rt or {}).get("regcode") or "",
            "Reissue issued": (rt or {}).get("issue") or "",
            "Tranche": tr.get(c.get("regcode") or "") or tr.get((rt or {}).get("regcode") or "")
                       or tr.get(c.get("pp") or "") or tr.get(c.get("cb") or ""),
            "Production batch (P lot)": c.get("pp") or "",
            "Cultivation batch": c.get("cb") or "",
            "Strain": c.get("strain") or "",
            "Product type": ptype,
            "Grade": ("Grade %s" % c["grade"]) if c.get("grade") else "",
            "Product code": c.get("pcode") or "",
            "Potency window (%)": w,
            "Nominal ± tol (%)": nt,
            "Total THC (%)": (r4.get("res") or "").replace("%", "").strip(),
            "CBD (%)": (r5.get("res") or "").strip(),
            "CBN (%)": (r6.get("res") or "").strip(),
            "Potency certificate": r4.get("doc") or "",
            "Certificate issued": r4.get("dd") or "",
            "Laboratory": r4.get("lab") or "",
            "Internal CoA (iCoA)": c.get("icoa_code") or "",
            "iCoA issued": c.get("icoa_issue") or "",
            "iCoA tested": c.get("icoa_tested") or "",
            "iCoA covers parameters": (ic.get("parameters") or "").strip(),
            "Retest campaign": ("Farmahem %s-series" % c["icoa_campaign"]) if c.get("icoa_campaign") else "",
        })
    return sorted(out, key=lambda r: order(r["CoQ code"], r["Production batch (P lot)"] or r["Cultivation batch"]))


def write(R, base):
    with io.open(base + ".csv", "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=COLS); w.writeheader(); w.writerows(R)
    ini = [r for r in R if r["Testing round"] == "Initial release"]
    ret = [r for r in R if r["Testing round"] != "Initial release"]
    L = ["# Certificates of quality — the register, by document code", "",
         "%d certificates: %d of the release round and %d twelve-month retests, ordered by certificate "
         "code. Each row carries the certificate it supersedes or is superseded by, and the internal "
         "certificate of analysis it cites with its date of issue. Written %s from the certificates' own "
         "data by `coq_master_register.py`." % (len(R), len(ini), len(ret), dt.date.today().strftime("%d.%m.%Y")), "",
         "| " + " | ".join(COLS) + " |", "| " + " | ".join("---" for _ in COLS) + " |"]
    for r in R:
        L.append("| " + " | ".join(str(r[c]).replace("|", "·") for c in COLS) + " |")
    io.open(base + ".md", "w", encoding="utf-8").write("\n".join(L) + "\n")
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    widths = {"CoQ code": 15, "Testing round": 22, "Issued": 11, "Register status": 26, "Legacy code": 17,
              "Supersedes": 15, "Superseded of": 12, "Superseded by": 15, "Reissue issued": 12,
              "Tranche": 8, "Production batch (P lot)": 12, "Cultivation batch": 14, "Strain": 22,
              "Product type": 32, "Grade": 9, "Product code": 18, "Potency window (%)": 15,
              "Nominal ± tol (%)": 14, "Total THC (%)": 11, "CBD (%)": 9, "CBN (%)": 9,
              "Potency certificate": 16, "Certificate issued": 11, "Laboratory": 22,
              "Internal CoA (iCoA)": 16, "iCoA issued": 11, "iCoA tested": 11,
              "iCoA covers parameters": 30, "Retest campaign": 20}
    centre = {"Tranche", "Grade", "Total THC (%)", "CBD (%)", "CBN (%)", "Issued", "Superseded of", "Reissue issued",
              "Certificate issued", "iCoA issued", "iCoA tested"}
    wb = Workbook()

    def sheet(ws, rows_):
        ws.append(COLS)
        for r in rows_:
            ws.append([r[c] for c in COLS])
        hdr = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="1B3A5C")
        for cell in ws[1]:
            cell.font = hdr; cell.fill = fill; cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions; ws.row_dimensions[1].height = 32
        for i, c in enumerate(COLS, 1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 14)
        thin = Side(style="thin", color="D9E0E8")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = Border(bottom=thin)
                if COLS[cell.column - 1] in centre:
                    cell.alignment = Alignment(horizontal="center")
                if COLS[cell.column - 1] == "CoQ code":
                    cell.font = Font(bold=True)
    ws = wb.active; ws.title = "All certificates"; sheet(ws, R)
    sheet(wb.create_sheet("Initial release"), ini)
    sheet(wb.create_sheet("Retest reissue"), ret)
    wb.save(base + ".xlsx")


def main(argv):
    R = rows(); base = os.path.join(HERE, "tracker", "CoQ_Master_Register_%s" % dt.date.today().isoformat())
    write(R, base)
    import collections
    print("certificates: %d  %s" % (len(R), dict(collections.Counter(r["Testing round"] for r in R))))
    pairs = sum(1 for r in R if r["Supersedes"]) ; back = sum(1 for r in R if r["Superseded by"])
    print("reissues naming the certificate they supersede: %d; release certificates naming their reissue: %d" % (pairs, back))
    print("rows citing an internal certificate: %d" % sum(1 for r in R if r["Internal CoA (iCoA)"]))
    print("wrote", os.path.relpath(base, ROOT) + ".xlsx / .csv / .md")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv))
