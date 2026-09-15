#!/usr/bin/env python3
"""Write the Tranche 3 Farmahem cannabinoid certificates (227-1-К/26 … 227-30-К/26) into the
owner's release register, in the shape of the Tranche 1 rows (197-n-К/26).

    python3 deliverables/qc_gap_analysis/intake_227K_2026-09-15/apply_227K.py
        [--register PATH]   the register to extend  (default: …SUBLOT_2026-09-01.xlsx)
        [--out PATH]        where to write          (default: in place)
        [--no-new-blocks]   hold the four certificates whose batch has no block

Idempotent: a certificate whose code is already in the register is skipped. The desk
does not write the owner's register on its own; this script is the write, ready to run.

Two reads must agree before a certificate is written: reads_227K.json (the page read
of 12.09.2026) and checkpoint_master_coa_table.json (the transcription of the same day
checkpointed in master_coa_table.tsv, commit 7393bc4), compared on the batch, the P lot,
the issue date and the three results. The checkpoint holds 25 of the 30; the five it does
not hold (227-1, -4, -8, -16, -29) rest on the page read alone and are written with that
said in the intake README and in OI-32 — the owner ruled on 12.09.2026 that every one of
the thirty is a retest, and on 15.09.2026 gave the campaign its sampling dates.

Row by row:
  * 26 certificates go into the batch's existing block, after its last real
    certificate row and before any "(not numbered)" / "n/a" placeholder — the block by
    the cultivation batch as printed, else by the P lot (227-29-К/26 prints
    BSS1024_01/1 and P050122; the register's block is BSS1024_01 / P050122, the same
    packaged lot);
  * 4 batches have no block and get one, numbered on from the last block: BSS1024_01/2
    (P050142), WED102501 (P060102), SCR012601 (P060342) and GRC102501/1 (P060142) — all
    four on the Head of QC's batch list with these P lots, three of them on the Tranche 3
    delivery list. The label row carries the certificate, as every block's does; strain
    and P lot are the certificate's own;
  * every row: THC % as printed, "≥ 5.00 %" in THC spec as the 197 rows carry it, CBD %
    and CBN % as printed with "< LOQ" written "<LOQ" as the 197 rows write it, "/" in
    every column the certificate does not report, the code as printed (Cyrillic К), the
    date of issue 11.09.2026, "Farmahem", "Open".
"""
import argparse, copy, json, os, sys
import openpyxl
from openpyxl.utils import range_boundaries

HERE = os.path.dirname(os.path.abspath(__file__))
GAP = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(os.path.dirname(os.path.dirname(GAP)), 'ingestion', 'common'))
from batch_id import batch_key  # noqa: E402

PLACEHOLDER = {'', '(not numbered)', 'n/a'}
THC_SPEC = '≥ 5.00 %'            # what every 197-n-К/26 row carries in THC spec
RESULTS = ('Total THC', 'Total CBD', 'Total CBN')


def reg_value(v):
    """A printed result in the register's own spelling: '< LOQ' is written '<LOQ'."""
    return str(v).replace('< LOQ', '<LOQ').strip()


def reg_code(code):
    """The code in the register's own spelling: Cyrillic К, as every 197-n-К/26 row.

    The page reads transcribed the series letter as Latin K on 12 certificates and
    Cyrillic К on 18 — the two glyphs are identical on the page and the reader said
    so. One spelling in the register, the one its Tranche 1 rows already use.
    """
    return str(code).strip().replace('-K/', '-К/')


def disagreements(m, c):
    """Where the page read and the checkpoint differ; empty when they agree; None with no checkpoint."""
    if not c:
        return None
    out = []
    if batch_key(c['cu']) != batch_key(m['batch_printed']):
        out.append('batch %r vs %r' % (c['cu'], m['batch_printed']))
    if (c['p'] or '') != (m['p_number'] or ''):
        out.append('P lot %r vs %r' % (c['p'], m['p_number']))
    if not c['issue'].startswith(m['date_of_issue']):
        out.append('issue date %r vs %r' % (c['issue'], m['date_of_issue']))
    for name in RESULTS:
        a = reg_value(c.get(name, '')).replace('%w/w', '').strip()
        b = reg_value(m['results'][name])
        if a != b:
            out.append('%s %r vs %r' % (name, c.get(name), m['results'][name]))
    return out


def load_reads():
    R = json.load(open(os.path.join(HERE, 'reads_227K.json'), encoding='utf-8'))
    C = json.load(open(os.path.join(HERE, 'checkpoint_master_coa_table.json'), encoding='utf-8'))
    bad, single = {}, []
    for fn, m in R.items():
        d = disagreements(m, C.get(m['cert_code']))
        if d is None:
            single.append(m['cert_code'])
        elif d:
            bad[fn] = d
    if bad:
        for fn, d in bad.items():
            print('DISAGREE', fn, '—', '; '.join(d))
        sys.exit('%d certificate(s) where the two reads differ — nothing written' % len(bad))
    return R, single


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--register', default=os.path.join(GAP, 'PP_Batch_Release_QC_Register_SUBLOT_2026-09-01.xlsx'))
    ap.add_argument('--out', default=None)
    ap.add_argument('--no-new-blocks', action='store_true',
                    help='write only the certificates whose batch already has a block; hold the four that would open one')
    a = ap.parse_args()
    out = a.out or a.register
    R, single = load_reads()
    wb = openpyxl.load_workbook(a.register)
    ws = wb['Batch Release QC']
    present = {reg_code(ws.cell(r, 23).value or '') for r in range(6, ws.max_row + 1)}
    real = lambda r: str(ws.cell(r, 23).value or '').strip() not in PLACEHOLDER

    labels = []
    for r in range(6, ws.max_row + 1):
        v = ws.cell(r, 2).value
        if v:
            labels.append((r, str(v).strip(), str(ws.cell(r, 3).value or '').strip()))
        elif str(ws.cell(r, 1).value or '').upper() == 'LEGEND':
            labels.append((r, None, None)); break
    blocks = {}
    for i, (r, b, p) in enumerate(labels):
        if b:
            blk = {'first': r, 'last': labels[i + 1][0] - 1, 'label': b}
            blocks[batch_key(b)] = blk
            if p.startswith('P'):
                blocks.setdefault(batch_key(p), blk)
    tmpl_row = next(r for r in range(6, ws.max_row + 1) if str(ws.cell(r, 23).value or '') == '197-17-К/26')
    assert str(ws.cell(tmpl_row, 6).value).strip() == THC_SPEC, ws.cell(tmpl_row, 6).value
    tmpl_label = next(r for r in range(6, ws.max_row + 1) if str(ws.cell(r, 2).value or '') == 'OPM1024')
    last_data = max(r for r in range(6, ws.max_row + 1) if ws.cell(r, 23).value not in (None, ''))
    last_no = max(int(ws.cell(r, 1).value) for r in range(6, ws.max_row + 1) if isinstance(ws.cell(r, 1).value, (int, float)))

    def values(m):
        v = [None] * 26
        r = m['results']
        v[4], v[5], v[6], v[7] = reg_value(r['Total THC']), THC_SPEC, reg_value(r['Total CBD']), reg_value(r['Total CBN'])
        for c in range(9, 23): v[c - 1] = '/'
        v[22], v[23], v[24], v[25] = reg_code(m['cert_code']), m['date_of_issue'], 'Farmahem', 'Open'
        return v

    def style_from(src, dst):
        for c in range(1, 27):
            s, d = ws.cell(src, c), ws.cell(dst, c)
            d.font, d.fill, d.border, d.alignment, d.number_format = copy.copy(s.font), copy.copy(s.fill), copy.copy(s.border), copy.copy(s.alignment), s.number_format
        ws.row_dimensions[dst].height = ws.row_dimensions[src].height

    todo = sorted(R.values(), key=lambda m: m['n'])
    todo = [m for m in todo if reg_code(m['cert_code']) not in present]
    into, new = [], []
    for m in todo:
        blk = next((blocks[batch_key(c)] for c in (m['batch_printed'], m['p_number']) if c and batch_key(c) in blocks), None)
        (into if blk else new).append((m, blk))
    if a.no_new_blocks:
        print('held (no block, --no-new-blocks): %s' % ', '.join(m['batch_printed'] for m, _ in new))
        new = []
    inserts = sorted(((max(r for r in range(blk['first'], blk['last'] + 1) if real(r)) + 1, m) for m, blk in into), key=lambda t: -t[0])
    merged_below = [str(mr) for mr in ws.merged_cells.ranges if mr.min_row > last_data]
    for mr in merged_below: ws.unmerge_cells(mr)
    for at, m in inserts:
        ws.insert_rows(at); style_from(tmpl_row + (1 if tmpl_row >= at else 0), at)
        for c, v in enumerate(values(m), 1): ws.cell(at, c).value = v
    last_data += len(inserts); tmpl_label += sum(1 for at, _ in inserts if at <= tmpl_label)
    for i, (m, _) in enumerate(new):
        at = last_data + 1 + i; ws.insert_rows(at); style_from(tmpl_label, at)
        v = values(m); v[0], v[1], v[2], v[3] = last_no + 1 + i, m['batch_printed'], m['p_number'], m['strain_printed']
        for c, val in enumerate(v, 1): ws.cell(at, c).value = val
    shift = len(inserts) + len(new)
    for mr in merged_below:
        c1, r1, c2, r2 = range_boundaries(mr); ws.merge_cells(start_row=r1 + shift, start_column=c1, end_row=r2 + shift, end_column=c2)
    ws.auto_filter.ref = 'A4:Z%d' % (last_data + len(new))
    wb.save(out)
    print('%s: %d row(s) into existing blocks, %d new block(s) No. %d-%d, %d already present, legend shifted by %d'
          % (os.path.basename(out), len(inserts), len(new), last_no + 1, last_no + len(new), 30 - len(todo), shift))
    if single:
        print('on the page read alone (no checkpoint transcription): ' + ', '.join(reg_code(c) for c in single))


if __name__ == '__main__':
    main()
