#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Bind the print fleet's record to the register's current numbering.

    python3 icoa_handoff/v3/sync_recs_to_register.py [--dry-run]

`build_fleet.js` reads `icoa_v44_recs.json` for each certificate's content and then
keeps only the records whose CODE a certificate of quality cites. That works while
the codes mean what they meant when the file was frozen at v44, and silently stops
working the moment the series is renumbered: the owner's ruling of 20.09.2026 ("one
iCOA per COQ, retest and initial") cut the register from 215 rows to 172 and
renumbered it contiguously, so code 007 now names Grape Pie's release round where
the frozen file still has it naming OMP1024_01's first retest. Filtering the old
records by the new codes would have printed 172 documents with the wrong content
under each code — and it very nearly did.

So the record is re-bound here, by IDENTITY rather than by number. Every row carries
a `key` — `CJ1024|I`, `P050202|R`, `P050202|R3` — which names the lot and the round
and does not move when the series is renumbered; all 172 of the register's rows are
found in the frozen file by it, exactly once each.

What this rewrites and what it deliberately does not:

  REWRITTEN   `code` and `no`, from the register, matched on `key`.
  UNTOUCHED   everything else — `packaging`, `harvest`, `identC`, `tested`, `scope`,
              `status` and the rest are copied across verbatim.

That distinction is the whole point. The register sheet holds these fields in a
different shape from the record: a display span ("24.04.2025 – 25.04.2025") where the
record holds one date, an annotation ("03.03.2025 (the issuance plan's basis:
27.02.2025)"), "— not stated —" where the record holds a date, and an empty identC
because the generator sources identification C from the certificate data instead.
Mapping the sheet's columns onto the record's fields would therefore change what
prints on 172 already-signed certificates while appearing to be a mere re-sync. Only
the numbering is stale, so only the numbering is rewritten.

The 43 rows the ruling withdrew are written to `icoa_withdrawn_2026-09-20.json`
rather than dropped, so the record of what they were survives the change — and they
are kept out of the fleet file itself, because leaving them in would give two records
the same code (a withdrawn row's old number is now some other row's new number) and
`build_fleet.js` would fail its own count check.
"""
import argparse
import json
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
GAP = os.path.dirname(os.path.dirname(HERE))
RECS = os.path.join(HERE, "icoa_v44_recs.json")
WITHDRAWN = os.path.join(HERE, "icoa_withdrawn_2026-09-20.json")
sys.path.insert(0, os.path.join(GAP, "tracker"))
from reference_sections import latest_master  # noqa: E402


def register_rows():
    """(key -> {code, no}) from the newest master's iCoA Register.

    Read through a spreadsheet engine because the sheet's dates are formulas and
    openpyxl stores no computed values — the same reason verify_workbook.py does it.
    """
    src = latest_master(os.path.join(GAP, "tracker"))
    tmp = tempfile.mkdtemp(prefix="sync_recs_")
    try:
        shutil.copy(src, os.path.join(tmp, "in.xlsx"))
        subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx",
                        "--outdir", os.path.join(tmp, "out"), os.path.join(tmp, "in.xlsx")],
                       check=True, capture_output=True, timeout=900)
        wb = openpyxl.load_workbook(os.path.join(tmp, "out", "in.xlsx"), data_only=True)
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
    ws = wb["iCoA Register"]
    hdr = [c.value for c in ws[1]]
    out = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(hdr, row))
        code = str(d.get("iCoA code") or "")
        if not code.startswith("iCoA-PP"):
            continue
        key = str(d.get("Key") or "").strip()
        if not key:
            raise SystemExit("an iCoA Register row has no Key, so it cannot be matched: " + code)
        if key in out:
            raise SystemExit("two iCoA Register rows share the Key %r — identity is not unique" % key)
        out[key] = {"code": code, "no": int(d["No."])}
    return os.path.basename(src), out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()

    master, reg = register_rows()
    recs = json.load(open(RECS, encoding="utf-8"))
    by_key = {}
    for r in recs:
        if r["key"] in by_key:
            raise SystemExit("the record holds two rows with key %r" % r["key"])
        by_key[r["key"]] = r
    print("%s: %d register row(s)" % (master, len(reg)))
    print("%s: %d record(s)" % (os.path.basename(RECS), len(recs)))

    missing = [k for k in reg if k not in by_key]
    if missing:
        raise SystemExit("%d register row(s) have no record to print from, so the fleet would be "
                         "short: %s" % (len(missing), missing[:6]))

    kept, withdrawn = [], []
    moved = 0
    for key, r in by_key.items():
        if key in reg:
            new = dict(r)
            if new["code"] != reg[key]["code"]:
                moved += 1
            new["code"] = reg[key]["code"]
            new["no"] = reg[key]["no"]
            kept.append(new)
        else:
            withdrawn.append(r)
    kept.sort(key=lambda r: r["no"])

    codes = [r["code"] for r in kept]
    nos = [r["no"] for r in kept]
    if len(set(codes)) != len(codes):
        raise SystemExit("the rebound record has duplicate codes, which build_fleet.js would print twice")
    if nos != list(range(1, len(nos) + 1)):
        raise SystemExit("the register's numbers are not contiguous from 1: %s" % nos[:8])

    print("kept %d (renumbered %d, unchanged %d) | withdrawn %d"
          % (len(kept), moved, len(kept) - moved, len(withdrawn)))
    print("codes %s .. %s, contiguous" % (codes[0], codes[-1]))
    if a.dry_run:
        print("--dry-run: nothing written")
        return 0
    json.dump(kept, open(RECS, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    json.dump(withdrawn, open(WITHDRAWN, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("wrote", RECS)
    print("wrote", WITHDRAWN, "(%d row(s) the ruling withdrew, kept as a record)" % len(withdrawn))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
