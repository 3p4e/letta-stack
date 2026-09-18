#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""The tranche retest sets of internal certificates, each as one Word document.

    python3 export_docx_icoa.py                 # Tranche 1 and Tranche 2
    python3 export_docx_icoa.py --dpi 200       # smaller files

The Head of QC, 18.09.2026: "give me Word documents for the internal certificates of
analysis for Tranche 1 and Tranche 2 — two documents with merged certificates."

The same rule as `export_docx_v40.py` for the certificates of quality, and for the same
reason: the certificate is an intricate fixed layout — three stacked style sheets, faded
bands, a 19-of-55 tick sheet, 14 px of clearance to the footer — that no HTML-to-Word
conversion reproduces faithfully, and a Word rebuild of the template would be a second
template to keep in step with the first. So each document carries the certificates as the
pages **exactly as they print**: every page of the merged tranche PDF rendered at 300 dpi
and placed edge to edge on an A4 page with zero margins, one certificate per page, in the
order the merge put them. It cannot drift from the PDF, because it is made from it.

What it is not: the text is not editable in Word. The editable source is the HTML under
ISSUE_iCOA/, which travels in the package; the controlled record is the PDF.

Output: pdf/iCoA_Retest_Tranche_<n>.docx, beside the PDF it is made from.
"""
import argparse
import io
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
PDF = os.path.join(HERE, "pdf")
SETS = [("iCoA_Retest_Tranche_1", "Tranche 1"), ("iCoA_Retest_Tranche_2", "Tranche 2")]


def build(dpi):
    import pymupdf
    from docx import Document
    from docx.enum.text import WD_BREAK
    from docx.shared import Mm

    made = 0
    for stem, label in SETS:
        src = os.path.join(PDF, stem + ".pdf")
        if not os.path.exists(src):
            print("  missing, skipped: %s" % os.path.relpath(src, HERE))
            continue
        book = pymupdf.open(src)
        doc = Document()
        s = doc.sections[0]
        s.page_width, s.page_height = Mm(210), Mm(297)
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Mm(0)
        s.header_distance = s.footer_distance = Mm(0)
        for i, page in enumerate(book):
            pix = page.get_pixmap(dpi=dpi, colorspace=pymupdf.csRGB, alpha=False)
            para = doc.add_paragraph()
            para.paragraph_format.space_before = para.paragraph_format.space_after = Mm(0)
            run = para.add_run()
            # a page break BEFORE every page but the first, so one certificate = one A4 sheet
            if i:
                run.add_break(WD_BREAK.PAGE)
            run.add_picture(io.BytesIO(pix.tobytes("png")), width=Mm(210), height=Mm(297))
        n = book.page_count
        book.close()
        core = doc.core_properties
        core.title = "%s — internal certificates of analysis, retest round" % label
        core.subject = "Internal Certificate of Analysis — Purely Plant GmbH"
        core.comments = ("%d certificates, each page rendered from the controlled vector "
                         "PDF at %d dpi; the PDF is the record." % (n, dpi))
        out = os.path.join(PDF, stem + ".docx")
        doc.save(out)
        made += 1
        print("  %-12s %3d certificate(s)  %s (%.1f MiB)"
              % (label, n, os.path.basename(out), os.path.getsize(out) / 1048576))
    return made


def main(argv):
    ap = argparse.ArgumentParser()
    ap.add_argument("--dpi", type=int, default=300)
    a = ap.parse_args(argv[1:])
    made = build(a.dpi)
    print("Word documents written: %d" % made)
    return 0 if made == len(SETS) else 1


if __name__ == "__main__":
    raise SystemExit(main(sys.argv))
