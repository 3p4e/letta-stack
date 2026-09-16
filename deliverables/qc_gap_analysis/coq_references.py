#!/usr/bin/env python3
"""One row per certificate of quality, one column per determination #1 to #12: the document
the CoQ cites for it — code, date of issue, laboratory abbreviation — read from
coq_artifact_data.json (the CoQ rows the drafts print from) with the desk's own laboratory
abbreviations from the tracker index. Writes <out>.csv, <out>.md, <out>.xlsx and <out>_nt.md,
and, from v29, fills the workbook's own `CoQ References` tab and its `Not Tested Review` section.

    python3 deliverables/qc_gap_analysis/coq_references.py [--src coq_artifact_data.json] [--out PREFIX]

Cell notation: `code date LAB`; `*` an internal CoA to be issued at the certificate's issue;
`also X` a later document on file for the same determination (the retest); `n/t` not tested,
with the sub-determinations when only some are; `(initial)` a reissue's row carried from the
initial certificate (owner, 15.09.2026); `DAB` the cited CNP certificate used the DAB monograph
(cnp_methods.py); `u/r` upon request; `OOS`, `undetermined`, `BLOCKED` the record's own verdict
flags. Rows: numbered CoQs in code order, then the unnumbered by batch.

Two columns beside the determinations (owner, 15.09.2026): `Sampled` — the retest campaign's
sampling day for a reissue, the packaging day for a release certificate — and `Received` — the
date each cited external certificate says the laboratory admitted the sample
(receipt_dates.py), one entry per certificate. In the spreadsheet every `n/t` cell is painted
red for the owner's check, and `<out>_nt.md` lists them.

The CoQ code is the register's: on the command line it is the code the export carries from
the extracted CoQ Register (coq_register_<date>.csv); inside the workbook build it is the
code the CoQ Register sheet of the same workbook prints, so the two tabs never disagree.
"""
import json, re, sys, os, collections, csv
G = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(G, 'tracker'))
sys.path.insert(0, G)
import tracker_data as T
import receipt_dates as RD

LONG = {'UKIM Faculty of Pharmacy — Center for Natural Products': 'CNP', 'CNP': 'CNP', 'IPH — Institute of Public Health': 'IJZ', 'IJZ': 'IJZ',
        'Purely Plant GmbH (in-house)': 'PP', 'State Phytosanitary Laboratory': 'DFL'}
COLS = ['CoQ', 'Batch', 'Series', 'Issue', 'Sampled', 'Received'] + [f'#{n}' for n in range(1, 13)]
WIDTHS = [16, 22, 18, 10, 11, 40] + [30] * 12
RED = 'F4B6B6'
NT_RX = re.compile(r'(^|; )n/t\b')


def _lab_map(instances):
    LAB = {}
    _, index_rows = T.load_owner()
    for r in index_rows: LAB[T.nkey(r['code'])] = r['lab']
    if instances and os.path.exists(instances):
        for i in json.load(open(instances, encoding='utf-8')): LAB[T.nkey(i['code'])] = i['lab']
    return LAB


def sd(dd):
    m = re.match(r'(\d\d)\.(\d\d)\.(\d{4})$', dd or '')
    return f"{m.group(1)}.{m.group(2)}.{m.group(3)[2:]}" if m else (dd or '')


def series(t):
    return {'initial release': 'release', 'initial release — predicted': 'release (predicted)',
            'additional testing (12-month)': 'retest', 'additional testing (12-month) — predicted': 'retest (predicted)'}[t]


def build_rows(src=None, instances=None, codes=None):
    """The table. `codes` — {(batch key, 'I'|'R'): CoQ code} — overrides the export's code
    (the workbook build passes its own CoQ Register)."""
    src = src or os.path.join(G, 'coq_artifact_data.json')
    instances = instances or os.path.join(G, 'tracker', 'new_instances.json')
    d = json.load(open(src, encoding='utf-8'))
    LAB = _lab_map(instances)

    def abbr(doc, lab):
        k = T.nkey(doc)
        if k in LAB: return LAB[k]
        if lab.startswith('Farmahem') or lab == 'FHM':
            return 'FHM-M' if re.search(r'[МM]/\d\d$', doc) else 'FHM-K' if re.search(r'[КK]/\d\d$', doc) else 'FHM'
        if doc.startswith('iCoA'): return 'PP'
        return LONG.get(lab, lab or '?')

    def cell(rows):
        """rows = the CoQ's rows for one determination (sub-rows folded)."""
        docs, notes = [], collections.OrderedDict()
        multi = len(rows) > 1
        for r in rows:
            doc, st = (r['doc'] or '').strip(), r['st']
            if doc and doc != '—':
                if doc.startswith('n/a — Purely Plant in-house CoA'): doc = 'in-house CoA, no number'
                tag = ''
                if st.startswith('to be performed'): tag = '*'
                elif 'OUT OF SPEC' in st: tag = ' OOS'
                elif 'UNDETERMINED' in st: tag = ' undetermined'
                elif 'BLOCKED' in st: tag = ' BLOCKED'
                if st.startswith('carried from the initial testing'): tag += ' (initial)'
                if 'DAB' in (r.get('mth') or ''): tag += ' DAB'
                item = f"{doc} {sd(r['dd'])} {abbr(doc, r['lab'])}{tag}"
                if item not in docs: docs.append(item)
                m = re.search(r'\(([^()]+)\)\s*$', r.get('also') or '')
                if m:
                    a = f"also {m.group(1)}"
                    if a not in docs: docs.append(a)
            else:
                if st.startswith('outside the retest scope'): n = '→ release'
                elif st.startswith('carried from the initial testing'): n = 'n/t (initial)'
                elif st.startswith('upon request'): n = 'u/r'
                elif st.startswith('not tested'): n = 'n/t'
                elif st.startswith('to be performed'): n = 'to be performed'
                elif st.startswith('awaiting the cannabinoid'): n = 'awaiting FHM-K'
                elif st.startswith('awaiting the mycotoxin'): n = 'awaiting FHM-M'
                elif st.startswith('in-house CoA only'): n = 'in-house record only'
                else: n = st[:40]
                notes.setdefault(n, []).append(str(r['no']))
        parts = list(docs)
        for n, subs in notes.items():
            if docs and n == 'u/r': continue
            parts.append(f"{n} {','.join(subs)}" if (multi and docs) else n)
        return '; '.join(parts)

    out = []
    for c in d['coqs']:
        code = c.get('regcode') or c.get('n') or '—'
        if codes:
            sfx = 'R' if c['t'].startswith('additional') else 'I'
            for name in (c.get('pp'), c.get('cb')):
                if name and (T.batch_key(name), sfx) in codes:
                    code = codes[(T.batch_key(name), sfx)]; break
        byno = collections.defaultdict(list)
        for r in c['rows']:
            n = int(str(r['no']).split('.')[0])
            if n == 9 and r['no'] in ('9.6', '9.7'): continue   # upon-request organisms, not release determinations
            byno[n].append(r)
        b = c['pp'] or c['cb']
        row = {'CoQ': code, 'Batch': b if (not c['cb'] or c['cb'] == b) else f"{b} ({c['cb']})", 'Series': series(c['t']), 'Issue': sd(c['issue']) or '—',
               'Sampled': sd(c.get('icoa_tested') or '') or '—'}
        rec, seen = [], set()
        for r in c['rows']:
            doc = (r['doc'] or '').strip()
            if not doc or doc == '—' or doc.startswith('iCoA') or doc.startswith('n/a') or T.nkey(doc) in seen: continue
            seen.add(T.nkey(doc))
            rec.append(f"{doc} {sd(RD.received(doc)) or '—'}")
        row['Received'] = '; '.join(rec) or '—'
        for n in range(1, 13): row[f'#{n}'] = cell(byno.get(n, []))
        out.append(row)

    def sk(r):
        m = re.match(r'CoQ-PP_26-(\d+)', r['CoQ'])
        return (0, int(m.group(1)), r['Batch']) if m else (1, 0, r['CoQ'], r['Batch'])
    out.sort(key=sk)
    return out


def nt_rows(rows):
    """Every n/t cell — (CoQ, Batch, Series, Parameter, Cell) — for the owner's check."""
    return [(r['CoQ'], r['Batch'], r['Series'], c, r[c]) for r in rows for c in COLS[6:] if NT_RX.search(str(r[c] or ''))]


def fill_sheet(ws, rows):
    """The table on a worksheet: header, widths, wrap, every n/t cell red, filter, freeze."""
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    ws.append(COLS)
    for r in rows: ws.append([r[c] for c in COLS])
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F3864'); c.alignment = Alignment(wrap_text=True, vertical='center')
    for i, w in enumerate(WIDTHS, 1): ws.column_dimensions[get_column_letter(i)].width = w
    red = PatternFill('solid', fgColor=RED)
    n = 0
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top'); c.font = Font(size=9)
            if c.column > 6 and NT_RX.search(str(c.value or '')):
                c.fill = red; n += 1
    ws.freeze_panes = 'E2'; ws.auto_filter.ref = ws.dimensions
    return n


def fill_nt_sheet(ws, nt):
    from openpyxl.styles import Font, Alignment, PatternFill
    from openpyxl.utils import get_column_letter
    ws.append(['CoQ', 'Batch', 'Series', 'Parameter', 'Cell'])
    for r in nt: ws.append(list(r))
    for c in ws[1]: c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='1F3864')
    for i, w in enumerate([16, 22, 18, 10, 60], 1): ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = Font(size=9); c.alignment = Alignment(wrap_text=True, vertical='top')
        row[4].fill = PatternFill('solid', fgColor=RED)


def write_files(rows, out):
    with open(out + '.csv', 'w', newline='', encoding='utf-8') as f:
        w = csv.DictWriter(f, fieldnames=COLS); w.writeheader(); w.writerows(rows)
    md = ['| ' + ' | '.join(COLS) + ' |', '|' + '---|' * len(COLS)]
    for r in rows: md.append('| ' + ' | '.join(str(r[c]).replace('|', '¦') for c in COLS) + ' |')
    open(out + '.md', 'w', encoding='utf-8').write('\n'.join(md) + '\n')
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'CoQ references'
    n = fill_sheet(ws, rows)
    nt = nt_rows(rows)
    fill_nt_sheet(wb.create_sheet('Not Tested Review'), nt)
    wb.save(out + '.xlsx')
    with open(out + '_nt.md', 'w', encoding='utf-8') as f:
        f.write("# Cells marked n/t — for the owner's check (%d)\n\n" % len(nt))
        f.write('| CoQ | Batch | Series | Parameter | Cell |\n|---|---|---|---|---|\n')
        for r in nt: f.write('| ' + ' | '.join(str(x).replace('|', '¦') for x in r) + ' |\n')
    return n


def main():
    import argparse
    ap = argparse.ArgumentParser()
    ap.add_argument('--src', default=os.path.join(G, 'coq_artifact_data.json'))
    ap.add_argument('--out', default=os.path.join(G, 'tracker', 'CoQ_references'))
    ap.add_argument('--instances', default=os.path.join(G, 'tracker', 'new_instances.json'))
    a = ap.parse_args()
    rows = build_rows(a.src, a.instances)
    n = write_files(rows, a.out)
    print(a.out + '.{csv,md,xlsx}:', len(rows), 'rows;', n, 'n/t cell(s) painted red, listed in', a.out + '_nt.md')
    print(collections.Counter(r['Series'] for r in rows))
    print(collections.Counter(r['CoQ'][:10] for r in rows))


if __name__ == '__main__':
    main()
